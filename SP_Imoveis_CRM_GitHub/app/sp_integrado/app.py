import io
import base64
import hashlib
import html
import importlib.util
import json
import math
import os
import random
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import unicodedata
from datetime import datetime, timedelta
from pathlib import Path
from urllib.parse import quote
from urllib import request as urlrequest

import pandas as pd
from flask import Flask, jsonify, render_template, request, send_file, send_from_directory
from werkzeug.utils import secure_filename

try:
    from selenium import webdriver
    from selenium.common.exceptions import TimeoutException
    from selenium.webdriver.chrome.options import Options
    from selenium.webdriver.chrome.webdriver import WebDriver as ChromeWebDriver
    from selenium.webdriver.common.by import By
    from selenium.webdriver.support import expected_conditions as EC
    from selenium.webdriver.support.ui import WebDriverWait

    SELENIUM_IMPORT_ERROR = ""
    SELENIUM_AVAILABLE = True
except Exception as exc:
    SELENIUM_IMPORT_ERROR = str(exc)
    SELENIUM_AVAILABLE = False

BASE_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = BASE_DIR.parent.parent
LEADS_SYNC_SCRIPT_PATH = PROJECT_ROOT / "AUTOMAÇÔES" / "Relatório_Leads" / "atualizar_crm_v0.py"
PROPERTY_SYNC_SCRIPT_PATH = PROJECT_ROOT / "AUTOMAÇÔES" / "Imoveis" / "Atualizar_imoveis_do_imobizi.py"
EXTERNAL_SELECTION_DIR = PROJECT_ROOT / "Atualização externa"
EXTERNAL_SCRIPT_PATH = EXTERNAL_SELECTION_DIR / "VIVAREAL - CODMAX - V1.py"


def load_local_env():
    env_path = BASE_DIR / ".env"
    if not env_path.exists():
        return
    for raw_line in env_path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


load_local_env()

configured_data_dir = os.environ.get("CRM_DATA_DIR", "").strip()
DATA_DIR = (
    Path(os.path.expandvars(configured_data_dir)).expanduser()
    if configured_data_dir
    else BASE_DIR / "data"
)
DATA_DIR = DATA_DIR.resolve()
UPLOAD_DIR = DATA_DIR / "uploads"
DB_PATH = DATA_DIR / "crm.sqlite3"

APP_BUILD = "20260505-dados-persistentes-v1"
INTEGRATION_TOKEN = os.environ.get("CRM_INTEGRATION_TOKEN", "").strip()
N8N_CRM_WEBHOOK_URL = os.environ.get("N8N_CRM_WEBHOOK_URL", "").strip()
N8N_CRM_WEBHOOK_SECRET = os.environ.get("N8N_CRM_WEBHOOK_SECRET", "").strip()
N8N_NOTIFY_TIMEOUT = float(os.environ.get("N8N_NOTIFY_TIMEOUT", "3") or 3)
DEFAULT_PROPERTY_FILE = os.environ.get("CRM_PROPERTY_FILE", "").strip()

configured_property_dir = os.environ.get("CRM_PROPERTY_DIR", "").strip()
PROPERTY_DIR = (
    Path(os.path.expandvars(configured_property_dir)).expanduser()
    if configured_property_dir
    else None
)

configured_export_columns_dir = os.environ.get("CRM_EXPORT_COLUMNS_DIR", "").strip()
EXPORT_COLUMNS_DIR = (
    Path(os.path.expandvars(configured_export_columns_dir)).expanduser()
    if configured_export_columns_dir
    else None
)

for folder in (DATA_DIR, UPLOAD_DIR):
    folder.mkdir(parents=True, exist_ok=True)
if EXPORT_COLUMNS_DIR:
    EXPORT_COLUMNS_DIR.mkdir(parents=True, exist_ok=True)

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 64 * 1024 * 1024

DEFAULT_STAGES = [
    {"id": "ativo", "label": "Ativo"},
    {"id": "negociando", "label": "Negociando"},
    {"id": "fechado", "label": "Fechado"},
    {"id": "corretores", "label": "Corretores"},
    {"id": "proprietarios", "label": "Proprietarios"},
]


def now_iso():
    return datetime.now().isoformat(timespec="seconds")


def parse_contact_datetime(value, label="data", fallback=None):
    try:
        if value is None or pd.isna(value):
            return fallback
    except Exception:
        if value is None:
            return fallback

    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime().replace(tzinfo=None).isoformat(timespec="seconds")
    if isinstance(value, datetime):
        return value.replace(tzinfo=None).isoformat(timespec="seconds")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
        if 1 <= number <= 60000:
            date_value = datetime(1899, 12, 30) + timedelta(days=number)
            return date_value.isoformat(timespec="seconds")

    text = clean_text(value)
    if not text:
        return fallback
    number = parse_brazilian_number(text)
    if pd.notna(number) and 1 <= number <= 60000:
        date_value = datetime(1899, 12, 30) + timedelta(days=float(number))
        return date_value.isoformat(timespec="seconds")

    parsed = pd.to_datetime(text.replace("T", " "), errors="coerce", dayfirst=True)
    if pd.isna(parsed):
        raise ValueError(f"{label} invalida: {text}")
    return parsed.to_pydatetime().replace(tzinfo=None).isoformat(timespec="seconds")


def normalize_key(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def slugify(value, fallback="coluna"):
    slug = normalize_key(value).replace(" ", "_").strip("_")
    return slug or fallback


def clean_text(value):
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    return str(value).strip()


def clean_phone(value):
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if value.is_integer():
            value = int(value)
    text = clean_text(value)
    text = re.sub(r"\.0+$", "", text)
    digits = re.sub(r"\D", "", text)
    return digits


def format_preview_cell(value):
    if isinstance(value, float) and not math.isnan(value) and value.is_integer():
        value = int(value)
    return str(value)


def parse_brazilian_number(value):
    if value is None:
        return math.nan
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = clean_text(value)
    if not text:
        return math.nan
    text = text.replace("R$", "").replace("r$", "").replace("\xa0", " ")
    text = re.sub(r"[^0-9,.-]", "", text)
    if not text:
        return math.nan
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", text):
        text = text.replace(".", "")
    try:
        return float(text)
    except ValueError:
        return math.nan


def parse_optional_number(value, label):
    text = clean_text(value)
    if not text:
        return None
    number = parse_brazilian_number(text)
    if pd.isna(number):
        raise ValueError(f"Valor invalido em {label}: {text}")
    return number


def parse_positive_int(value, label):
    text = clean_text(value)
    if not text:
        return None
    if not text.isdigit():
        raise ValueError(f"Valor invalido em {label}: {text}")
    number = int(text)
    if number <= 0:
        raise ValueError(f"Valor invalido em {label}: {text}")
    return number


def format_currency(value):
    number = parse_brazilian_number(value)
    if pd.isna(number):
        return "Não informado"
    return f"R$ {number:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def format_area(value):
    number = parse_brazilian_number(value)
    if pd.isna(number):
        return "Não informado"
    if abs(number - round(number)) < 1e-9:
        return f"{int(round(number)):,}".replace(",", ".") + " m²"
    return f"{number:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".") + " m²"


def cell_display(value):
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if hasattr(value, "strftime"):
        try:
            has_time = getattr(value, "hour", 0) or getattr(value, "minute", 0)
            return value.strftime("%d/%m/%Y %H:%M" if has_time else "%d/%m/%Y")
        except Exception:
            pass
    return str(value).strip()


def read_tabular_file(path):
    path = Path(path)
    if path.suffix.lower() in {".xlsx", ".xls"}:
        return pd.read_excel(path)
    last_error = None
    for encoding in ("utf-8", "utf-8-sig", "latin1", "cp1252"):
        try:
            return pd.read_csv(path, encoding=encoding)
        except Exception as exc:
            last_error = exc
    raise last_error


def save_upload(file_storage, prefix):
    original = secure_filename(file_storage.filename or "upload")
    suffix = Path(original).suffix.lower() or ".xlsx"
    filename = f"{prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{suffix}"
    path = UPLOAD_DIR / filename
    file_storage.save(path)
    return path


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with get_db() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS stages (
                id TEXT PRIMARY KEY,
                label TEXT NOT NULL,
                position INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS conversations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL DEFAULT '',
                phone TEXT NOT NULL DEFAULT '',
                stage TEXT NOT NULL DEFAULT 'ativo',
                kind TEXT NOT NULL DEFAULT 'Cliente',
                source TEXT NOT NULL DEFAULT '',
                owner TEXT NOT NULL DEFAULT '',
                last_message TEXT NOT NULL DEFAULT '',
                notes TEXT NOT NULL DEFAULT '',
                property_codes TEXT NOT NULL DEFAULT '',
                color TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            )
            """
        )
        existing_columns = {
            row["name"] for row in conn.execute("PRAGMA table_info(conversations)").fetchall()
        }
        if "color" not in existing_columns:
            conn.execute("ALTER TABLE conversations ADD COLUMN color TEXT NOT NULL DEFAULT ''")
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                conversation_id INTEGER NOT NULL,
                direction TEXT NOT NULL DEFAULT 'note',
                body TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(conversation_id) REFERENCES conversations(id)
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS automation_runs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                conversation_id INTEGER NOT NULL,
                demand_hash TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'recebido',
                demand_json TEXT NOT NULL DEFAULT '{}',
                selector_json TEXT NOT NULL DEFAULT '{}',
                dispatch_json TEXT NOT NULL DEFAULT '{}',
                message TEXT NOT NULL DEFAULT '',
                error TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                sent_at TEXT NOT NULL DEFAULT '',
                UNIQUE(conversation_id, demand_hash),
                FOREIGN KEY(conversation_id) REFERENCES conversations(id)
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS automation_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                conversation_id INTEGER,
                demand_hash TEXT NOT NULL DEFAULT '',
                event TEXT NOT NULL,
                message TEXT NOT NULL DEFAULT '',
                payload_json TEXT NOT NULL DEFAULT '{}',
                created_at TEXT NOT NULL
            )
            """
        )
        existing = conn.execute("SELECT COUNT(*) AS total FROM stages").fetchone()["total"]
        if existing == 0:
            now = now_iso()
            for position, stage in enumerate(DEFAULT_STAGES):
                conn.execute(
                    "INSERT INTO stages (id, label, position, created_at) VALUES (?, ?, ?, ?)",
                    (stage["id"], stage["label"], position, now),
                )


def get_stages(conn=None):
    owns_conn = conn is None
    conn = conn or get_db()
    try:
        rows = conn.execute(
            "SELECT id, label FROM stages ORDER BY position ASC, label ASC"
        ).fetchall()
        return [{"id": row["id"], "label": row["label"]} for row in rows]
    finally:
        if owns_conn:
            conn.close()


def get_stage_ids(conn=None):
    return {stage["id"] for stage in get_stages(conn)}


def fallback_stage(conn=None, avoid=None):
    stages = [stage["id"] for stage in get_stages(conn) if stage["id"] != avoid]
    if "ativo" in stages:
        return "ativo"
    return stages[0] if stages else "ativo"


def resolve_stage(value, conn=None, default="ativo"):
    raw = clean_text(value)
    stages = get_stages(conn)
    by_id = {stage["id"]: stage["id"] for stage in stages}
    by_label = {normalize_key(stage["label"]): stage["id"] for stage in stages}
    by_label.update({normalize_key(stage["id"]): stage["id"] for stage in stages})
    if raw in by_id:
        return raw
    return by_label.get(normalize_key(raw), default if default in {stage["id"] for stage in stages} else fallback_stage(conn))


CARD_COLORS = {"verde", "amarelo", "azul", "vermelho"}
CARD_COLOR_ALIASES = {
    "verde": "verde",
    "green": "verde",
    "amarelo": "amarelo",
    "amarela": "amarelo",
    "yellow": "amarelo",
    "azul": "azul",
    "blue": "azul",
    "vermelho": "vermelho",
    "vermelha": "vermelho",
    "red": "vermelho",
}


def resolve_card_color(value):
    key = normalize_key(value)
    if not key:
        return ""
    return CARD_COLOR_ALIASES.get(key, "")


def make_unique_stage_id(label, conn):
    base = slugify(label)
    stage_id = base
    existing = get_stage_ids(conn)
    index = 2
    while stage_id in existing:
        stage_id = f"{base}_{index}"
        index += 1
    return stage_id


def conversation_row_to_dict(row):
    stage_ids = get_stage_ids()
    return {
        "id": row["id"],
        "name": row["name"],
        "phone": row["phone"],
        "stage": row["stage"] if row["stage"] in stage_ids else fallback_stage(),
        "kind": row["kind"],
        "source": row["source"],
        "owner": row["owner"],
        "last_message": row["last_message"],
        "notes": row["notes"],
        "property_codes": row["property_codes"],
        "color": row["color"] if row["color"] in CARD_COLORS else "",
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
    }


def integration_is_authorized():
    if not INTEGRATION_TOKEN:
        return True
    token = (
        request.headers.get("X-Integration-Token")
        or request.headers.get("Authorization", "").replace("Bearer ", "", 1)
        or request.args.get("token")
        or ""
    ).strip()
    return token == INTEGRATION_TOKEN


def integration_auth_error():
    return jsonify({"ok": False, "error": "Token de integracao invalido"}), 401


def strip_html_to_text(value):
    text = html.unescape(clean_text(value))
    if not text:
        return ""
    text = re.sub(r"<\s*br\s*/?\s*>", "\n", text, flags=re.IGNORECASE)
    text = re.sub(r"</\s*(p|div|tr|li|th|td|h[1-6]|section|article)\s*>", "\n", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)
    text = text.replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    return "\n".join([line for line in lines if line]).strip()


OUTBOUND_METADATA_LINES = {
    "acao",
    "id",
    "nome",
    "telefone",
    "numero wpp",
    "whatsapp",
    "tipo",
    "etapa",
    "stage",
    "coluna",
    "status",
    "origem",
    "responsavel",
    "lead id",
    "demanda hash",
    "nome etapa",
}


def sanitize_outbound_message(message, name=""):
    text = strip_html_to_text(message)
    if not text:
        return ""
    name_key = normalize_key(name)
    cleaned = []
    for line in text.split("\n"):
        line = line.strip()
        key = normalize_key(line)
        if not line or key in OUTBOUND_METADATA_LINES:
            continue
        cleaned.append(line)
    while cleaned and name_key and normalize_key(cleaned[-1]) == name_key:
        cleaned.pop()
    return "\n".join(cleaned).strip()


def json_dumps(data):
    return json.dumps(data or {}, ensure_ascii=False, sort_keys=True, default=str)


def automation_status_is_duplicate(status):
    return status in {
        "disparo_preparado",
        "disparo_iniciado",
        "disparo_feito",
        "sent",
        "done",
    }


def add_automation_log(conversation_id, event, message="", payload=None, demand_hash="", conn=None):
    owns_conn = conn is None
    if owns_conn:
        conn = get_db()
    try:
        conn.execute(
            """
            INSERT INTO automation_logs
            (conversation_id, demand_hash, event, message, payload_json, created_at)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                conversation_id,
                clean_text(demand_hash),
                clean_text(event),
                clean_text(message),
                json_dumps(payload),
                now_iso(),
            ),
        )
        if owns_conn:
            conn.commit()
    finally:
        if owns_conn:
            conn.close()


def upsert_automation_run(
    conn,
    conversation_id,
    demand_hash,
    status,
    demand=None,
    selector=None,
    dispatch=None,
    message="",
    error="",
    sent=False,
):
    now = now_iso()
    existing = conn.execute(
        "SELECT * FROM automation_runs WHERE conversation_id = ? AND demand_hash = ?",
        (conversation_id, demand_hash),
    ).fetchone()
    demand_json = json_dumps(demand) if demand is not None else None
    selector_json = json_dumps(selector) if selector is not None else None
    dispatch_json = json_dumps(dispatch) if dispatch is not None else None
    if existing is None:
        conn.execute(
            """
            INSERT INTO automation_runs
            (conversation_id, demand_hash, status, demand_json, selector_json, dispatch_json,
             message, error, created_at, updated_at, sent_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                conversation_id,
                demand_hash,
                status,
                demand_json or "{}",
                selector_json or "{}",
                dispatch_json or "{}",
                clean_text(message),
                clean_text(error),
                now,
                now,
                now if sent else "",
            ),
        )
        return
    updates = ["status = ?", "updated_at = ?"]
    values = [status, now]
    if demand_json is not None:
        updates.append("demand_json = ?")
        values.append(demand_json)
    if selector_json is not None:
        updates.append("selector_json = ?")
        values.append(selector_json)
    if dispatch_json is not None:
        updates.append("dispatch_json = ?")
        values.append(dispatch_json)
    if message:
        updates.append("message = ?")
        values.append(clean_text(message))
    if error or status.startswith("erro"):
        updates.append("error = ?")
        values.append(clean_text(error))
    if sent:
        updates.append("sent_at = ?")
        values.append(now)
    values.extend([conversation_id, demand_hash])
    conn.execute(
        f"UPDATE automation_runs SET {', '.join(updates)} WHERE conversation_id = ? AND demand_hash = ?",
        values,
    )


def insert_crm_message(conn, conversation_id, direction, body):
    body = sanitize_outbound_message(body) if direction == "out" else strip_html_to_text(body)
    if not body:
        return None
    now = now_iso()
    cur = conn.execute(
        "INSERT INTO messages (conversation_id, direction, body, created_at) VALUES (?, ?, ?, ?)",
        (conversation_id, direction if direction in {"note", "in", "out"} else "note", body, now),
    )
    conn.execute(
        "UPDATE conversations SET last_message = ?, updated_at = ? WHERE id = ?",
        (body[:240], now, conversation_id),
    )
    return cur.lastrowid


init_db()


def post_json_async(url, payload, headers=None):
    def worker():
        try:
            body = json.dumps(payload, ensure_ascii=False, default=str).encode("utf-8")
            req = urlrequest.Request(
                url,
                data=body,
                headers={"Content-Type": "application/json", **(headers or {})},
                method="POST",
            )
            with urlrequest.urlopen(req, timeout=N8N_NOTIFY_TIMEOUT):
                pass
        except Exception as exc:
            add_blaster_log(f"Falha ao notificar N8N: {exc}", "warn")

    threading.Thread(target=worker, daemon=True).start()


def notify_n8n_lead_event(event, conversation_id):
    if not N8N_CRM_WEBHOOK_URL:
        return
    payload = {
        "event": event,
        "lead_id": conversation_id,
        "conversation_id": conversation_id,
        "source": "crm_app",
        "created_at": now_iso(),
    }
    headers = {}
    if N8N_CRM_WEBHOOK_SECRET:
        headers["X-CRM-Webhook-Secret"] = N8N_CRM_WEBHOOK_SECRET
    post_json_async(N8N_CRM_WEBHOOK_URL, payload, headers=headers)


# CRM endpoints
@app.get("/api/crm/board")
def api_crm_board():
    with get_db() as conn:
        stages = get_stages(conn)
        rows = conn.execute(
            "SELECT * FROM conversations ORDER BY datetime(updated_at) DESC, id DESC"
        ).fetchall()
    return jsonify(
        {"ok": True, "stages": stages, "contacts": [conversation_row_to_dict(row) for row in rows]}
    )


@app.get("/api/crm/template")
def api_crm_template():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    stages = get_stages()
    wb = Workbook()
    ws = wb.active
    ws.title = "Contatos"
    headers = [
        "Acao",
        "ID",
        "Nome",
        "Telefone",
        "Tipo",
        "Etapa",
        "Observacoes",
        "Mensagem / nota",
        "Direcao da mensagem",
        "Origem",
        "Responsavel",
        "Codigos imoveis",
        "Data de criacao",
        "Ultima mensagem",
    ]
    ws.append(headers)
    ws.append(
        [
            "I",
            "",
            "Joao da Silva",
            "11999999999",
            "Cliente",
            "Ativo",
            "Cliente interessado em galpao na Mooca.",
            "Primeiro contato feito pelo WhatsApp.",
            "Nota",
            "Importacao",
            "Equipe SP",
            "A1; B2",
            datetime.now().strftime("%d/%m/%Y %H:%M"),
            "Aguardando retorno.",
        ]
    )
    header_fill = PatternFill("solid", fgColor="0F766E")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    widths = [12, 10, 30, 18, 18, 20, 36, 42, 22, 20, 22, 24, 22, 34]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    ws.freeze_panes = "A2"

    opts = wb.create_sheet("Opcoes")
    opts.append(["Acoes", "Tipos", "Etapas", "Direcoes"])
    option_rows = max(3, len(stages))
    for row_num in range(2, option_rows + 2):
        action_value = ["I", "A", "E"][row_num - 2] if row_num <= 4 else ""
        type_value = ["Cliente", "Corretor", "Proprietario"][row_num - 2] if row_num <= 4 else ""
        stage_value = stages[row_num - 2]["label"] if row_num - 2 < len(stages) else ""
        direction_value = ["Nota", "Recebida", "Enviada"][row_num - 2] if row_num <= 4 else ""
        opts.append([action_value, type_value, stage_value, direction_value])
    for cell in opts[1]:
        cell.font = Font(bold=True)
    for col in ("A", "B", "C", "D"):
        opts.column_dimensions[col].width = 24

    max_stage_row = max(2, len(stages) + 1)
    action_validation = DataValidation(type="list", formula1="=Opcoes!$A$2:$A$4", allow_blank=True)
    type_validation = DataValidation(type="list", formula1="=Opcoes!$B$2:$B$4", allow_blank=True)
    stage_validation = DataValidation(type="list", formula1=f"=Opcoes!$C$2:$C${max_stage_row}", allow_blank=True)
    direction_validation = DataValidation(type="list", formula1="=Opcoes!$D$2:$D$4", allow_blank=True)
    ws.add_data_validation(action_validation)
    ws.add_data_validation(type_validation)
    ws.add_data_validation(stage_validation)
    ws.add_data_validation(direction_validation)
    action_validation.add("A2:A500")
    type_validation.add("E2:E500")
    stage_validation.add("F2:F500")
    direction_validation.add("I2:I500")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="modelo_importacao_crm.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/api/crm/export_columns", methods=["GET", "POST"])
def api_crm_export_columns():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    headers = [
        "Acao",
        "ID",
        "Nome",
        "Telefone",
        "Tipo",
        "Etapa",
        "Observacoes",
        "Origem",
        "Responsavel",
        "Codigos imoveis",
        "Ultima mensagem",
        "Data de criacao",
        "Atualizado em",
    ]

    with get_db() as conn:
        stages = get_stages(conn)
        rows = conn.execute(
            "SELECT * FROM conversations ORDER BY datetime(updated_at) DESC, id DESC"
        ).fetchall()

    contacts = [conversation_row_to_dict(row) for row in rows]
    stage_labels = {stage["id"]: stage["label"] for stage in stages}

    def sheet_name(label):
        cleaned = re.sub(r"[\[\]\*:/\\?]", " ", clean_text(label))[:31].strip()
        return cleaned or "Coluna"

    def write_sheet(ws, sheet_contacts):
        ws.append(headers)
        fill = PatternFill("solid", fgColor="0F766E")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
        for contact in sheet_contacts:
            ws.append(
                [
                    "",
                    contact["id"],
                    contact["name"],
                    contact["phone"],
                    contact["kind"],
                    stage_labels.get(contact["stage"], contact["stage"]),
                    contact["notes"],
                    contact["source"],
                    contact["owner"],
                    contact["property_codes"],
                    contact["last_message"],
                    contact["created_at"],
                    contact["updated_at"],
                ]
            )
        widths = [12, 10, 30, 18, 18, 22, 38, 18, 22, 24, 40, 20, 20]
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width
        ws.freeze_panes = "A2"

    wb = Workbook()
    first = wb.active
    first.title = "Todas as colunas"
    write_sheet(first, contacts)

    used_names = {first.title}
    for stage in stages:
        base = sheet_name(stage["label"])
        name = base
        suffix = 2
        while name in used_names:
            name = f"{base[:28]} {suffix}"[:31]
            suffix += 1
        used_names.add(name)
        ws = wb.create_sheet(name)
        write_sheet(ws, [contact for contact in contacts if contact["stage"] == stage["id"]])

    if EXPORT_COLUMNS_DIR:
        filepath = EXPORT_COLUMNS_DIR / "crm_colunas_sp_imoveis.xlsx"
        wb.save(filepath)
        return jsonify(ok=True, path=str(filepath))

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="crm_colunas_sp_imoveis.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.post("/api/crm/stages")
def api_crm_create_stage():
    payload = request.get_json(silent=True) or {}
    label = clean_text(payload.get("label"))
    if not label:
        return jsonify({"ok": False, "error": "Informe o nome da coluna"}), 400
    with get_db() as conn:
        stage_id = make_unique_stage_id(label, conn)
        position = conn.execute("SELECT COALESCE(MAX(position), -1) + 1 AS next_pos FROM stages").fetchone()["next_pos"]
        conn.execute(
            "INSERT INTO stages (id, label, position, created_at) VALUES (?, ?, ?, ?)",
            (stage_id, label, position, now_iso()),
        )
    return jsonify({"ok": True, "stage": {"id": stage_id, "label": label}})


@app.post("/api/crm/stages/reorder")
def api_crm_reorder_stages():
    payload = request.get_json(silent=True) or {}
    order = payload.get("order")
    if not isinstance(order, list) or not order:
        return jsonify({"ok": False, "error": "Informe a nova ordem das colunas"}), 400
    with get_db() as conn:
        existing = [stage["id"] for stage in get_stages(conn)]
        existing_set = set(existing)
        cleaned = []
        for item in order:
            stage_id = str(item)
            if stage_id in existing_set and stage_id not in cleaned:
                cleaned.append(stage_id)
        # qualquer coluna que nao veio na lista (nunca deveria acontecer) fica
        # no fim, na ordem em que ja estava, pra nao sumir do board
        missing = [stage_id for stage_id in existing if stage_id not in cleaned]
        final_order = cleaned + missing
        for position, stage_id in enumerate(final_order):
            conn.execute("UPDATE stages SET position = ? WHERE id = ?", (position, stage_id))
    return jsonify({"ok": True, "stages": get_stages()})


@app.delete("/api/crm/stages/<stage_id>")
def api_crm_delete_stage(stage_id):
    with get_db() as conn:
        stages = get_stages(conn)
        if stage_id not in {stage["id"] for stage in stages}:
            return jsonify({"ok": False, "error": "Coluna nao encontrada"}), 404
        if len(stages) <= 1:
            return jsonify({"ok": False, "error": "Mantenha pelo menos uma coluna"}), 400
        target = resolve_stage((request.get_json(silent=True) or {}).get("target_stage"), conn, fallback_stage(conn, stage_id))
        if target == stage_id:
            target = fallback_stage(conn, stage_id)
        conn.execute("UPDATE conversations SET stage = ?, updated_at = ? WHERE stage = ?", (target, now_iso(), stage_id))
        conn.execute("DELETE FROM stages WHERE id = ?", (stage_id,))
    return jsonify({"ok": True, "target_stage": target})


@app.post("/api/crm/conversations")
def api_crm_create_conversation():
    payload = request.get_json(silent=True) or request.form
    name = clean_text(payload.get("name")) or "Sem nome"
    phone = clean_phone(payload.get("phone"))
    now = now_iso()
    with get_db() as conn:
        stage = resolve_stage(payload.get("stage"), conn, fallback_stage(conn))
        cur = conn.execute(
            """
            INSERT INTO conversations
            (name, phone, stage, kind, source, owner, last_message, notes, property_codes, color, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                name,
                phone,
                stage,
                clean_text(payload.get("kind")) or "Cliente",
                clean_text(payload.get("source")),
                clean_text(payload.get("owner")),
                clean_text(payload.get("last_message")),
                clean_text(payload.get("notes")),
                clean_text(payload.get("property_codes")),
                resolve_card_color(payload.get("color")),
                now,
                now,
            ),
        )
        row = conn.execute("SELECT * FROM conversations WHERE id = ?", (cur.lastrowid,)).fetchone()
    notify_n8n_lead_event("lead_created", row["id"])
    return jsonify({"ok": True, "contact": conversation_row_to_dict(row)})


@app.patch("/api/crm/conversations/<int:conversation_id>")
def api_crm_update_conversation(conversation_id):
    payload = request.get_json(silent=True) or {}
    allowed = {
        "name",
        "phone",
        "stage",
        "kind",
        "source",
        "owner",
        "last_message",
        "notes",
        "property_codes",
        "color",
        "created_at",
    }
    updates = []
    values = []
    for field in allowed:
        if field in payload:
            value = clean_phone(payload[field]) if field == "phone" else clean_text(payload[field])
            if field == "stage":
                with get_db() as conn:
                    value = resolve_stage(value, conn, fallback_stage(conn))
            if field == "color":
                value = resolve_card_color(value)
            if field == "created_at":
                if not value:
                    continue
                try:
                    value = parse_contact_datetime(value, "Data de criacao")
                except ValueError as exc:
                    return jsonify({"ok": False, "error": str(exc)}), 400
            updates.append(f"{field} = ?")
            values.append(value)
    if not updates:
        return jsonify({"ok": False, "error": "Nenhum campo para atualizar"}), 400
    updates.append("updated_at = ?")
    values.append(now_iso())
    values.append(conversation_id)
    with get_db() as conn:
        conn.execute(f"UPDATE conversations SET {', '.join(updates)} WHERE id = ?", values)
        row = conn.execute("SELECT * FROM conversations WHERE id = ?", (conversation_id,)).fetchone()
    if row is None:
        return jsonify({"ok": False, "error": "Contato nao encontrado"}), 404
    notify_n8n_lead_event("lead_updated", conversation_id)
    return jsonify({"ok": True, "contact": conversation_row_to_dict(row)})


@app.get("/api/crm/conversations/<int:conversation_id>/messages")
def api_crm_messages(conversation_id):
    with get_db() as conn:
        rows = conn.execute(
            "SELECT * FROM messages WHERE conversation_id = ? ORDER BY datetime(created_at) ASC, id ASC",
            (conversation_id,),
        ).fetchall()
    messages = [
        {
            "id": row["id"],
            "conversation_id": row["conversation_id"],
            "direction": row["direction"],
            "body": row["body"],
            "created_at": row["created_at"],
        }
        for row in rows
    ]
    return jsonify({"ok": True, "messages": messages})


@app.post("/api/crm/conversations/<int:conversation_id>/messages")
def api_crm_add_message(conversation_id):
    payload = request.get_json(silent=True) or {}
    body = clean_text(payload.get("body"))
    if not body:
        return jsonify({"ok": False, "error": "Mensagem vazia"}), 400
    direction = clean_text(payload.get("direction")) or "note"
    if direction not in {"note", "in", "out"}:
        direction = "note"
    now = now_iso()
    with get_db() as conn:
        exists = conn.execute(
            "SELECT id FROM conversations WHERE id = ?", (conversation_id,)
        ).fetchone()
        if not exists:
            return jsonify({"ok": False, "error": "Contato nao encontrado"}), 404
        cur = conn.execute(
            "INSERT INTO messages (conversation_id, direction, body, created_at) VALUES (?, ?, ?, ?)",
            (conversation_id, direction, body, now),
        )
        conn.execute(
            "UPDATE conversations SET last_message = ?, updated_at = ? WHERE id = ?",
            (body[:240], now, conversation_id),
        )
        row = conn.execute("SELECT * FROM messages WHERE id = ?", (cur.lastrowid,)).fetchone()
    notify_n8n_lead_event("message_added", conversation_id)
    return jsonify(
        {
            "ok": True,
            "message": {
                "id": row["id"],
                "conversation_id": row["conversation_id"],
                "direction": row["direction"],
                "body": row["body"],
                "created_at": row["created_at"],
            },
        }
    )


@app.delete("/api/crm/conversations/<int:conversation_id>/messages/<int:message_id>")
def api_crm_delete_message(conversation_id, message_id):
    with get_db() as conn:
        row = conn.execute(
            "SELECT id FROM messages WHERE id = ? AND conversation_id = ?",
            (message_id, conversation_id),
        ).fetchone()
        if row is None:
            return jsonify({"ok": False, "error": "Nota nao encontrada"}), 404
        conn.execute(
            "DELETE FROM messages WHERE id = ? AND conversation_id = ?",
            (message_id, conversation_id),
        )
        latest = conn.execute(
            "SELECT body FROM messages WHERE conversation_id = ? ORDER BY datetime(created_at) DESC, id DESC LIMIT 1",
            (conversation_id,),
        ).fetchone()
        conn.execute(
            "UPDATE conversations SET last_message = ?, updated_at = ? WHERE id = ?",
            ((latest["body"][:240] if latest else ""), now_iso(), conversation_id),
        )
    return jsonify({"ok": True})


def get_import_col(row, normalized_cols, *names):
    col = find_import_col(normalized_cols, *names)
    if col is not None:
        return row.get(col)
    return ""


def find_import_col(normalized_cols, *names):
    for name in names:
        col = normalized_cols.get(normalize_key(name))
        if col is not None:
            return col
    return None


def parse_import_id(value):
    text = clean_text(value)
    if not text:
        return None
    number = parse_brazilian_number(text)
    if pd.notna(number) and number > 0 and abs(number - int(number)) < 1e-9:
        return int(number)
    if text.isdigit():
        return int(text)
    return None


def parse_bulk_action(value, has_action_col):
    if not has_action_col:
        return "I"
    key = normalize_key(value)
    if not key:
        return ""
    if key in {"a", "alterar", "alteracao", "atualizar"}:
        return "A"
    if key in {"e", "excluir", "exclusao", "deletar", "delete"}:
        return "E"
    if key in {"i", "inserir", "inclusao", "novo", "criar"}:
        return "I"
    return "?"


def find_contact_for_import(conn, row, cols_norm):
    contact_id = parse_import_id(get_import_col(row, cols_norm, "ID", "Id", "CRM ID", "Codigo CRM"))
    if contact_id:
        found = conn.execute("SELECT * FROM conversations WHERE id = ?", (contact_id,)).fetchone()
        if found is not None:
            return found

    phone = clean_phone(get_import_col(row, cols_norm, "Telefone", "WhatsApp", "Numero WPP", "Numero", "Celular"))
    if phone:
        return conn.execute(
            "SELECT * FROM conversations WHERE phone = ? ORDER BY id DESC LIMIT 1",
            (phone,),
        ).fetchone()
    return None


def parse_message_direction(value):
    key = normalize_key(value)
    if key in {"recebida", "entrada", "in"}:
        return "in"
    if key in {"enviada", "saida", "out"}:
        return "out"
    return "note"


@app.post("/api/crm/import_contacts")
def api_crm_import_contacts():
    file = request.files.get("file")
    if not file:
        return jsonify({"ok": False, "error": "Nenhum arquivo enviado"}), 400
    path = save_upload(file, "crm")
    df = read_tabular_file(path).fillna("")
    cols_norm = {normalize_key(col): col for col in df.columns}
    action_col = find_import_col(cols_norm, "Acao", "A\u00e7\u00e3o", "Comando", "Operacao", "Opera\u00e7\u00e3o", "Coluna 1")
    has_action_col = action_col is not None
    stats = {"inserted": 0, "updated": 0, "deleted": 0, "ignored": 0, "failed": 0}
    errors = []

    def add_error(row_number, message):
        stats["failed"] += 1
        if len(errors) < 12:
            errors.append(f"Linha {row_number}: {message}")

    def row_text(row, *names):
        return clean_text(get_import_col(row, cols_norm, *names))

    def row_phone(row):
        return clean_phone(get_import_col(row, cols_norm, "Telefone", "WhatsApp", "Numero WPP", "Numero", "Celular"))

    def row_created_at(row, fallback):
        value = get_import_col(row, cols_norm, "Data de criacao", "Data de criação", "Criado em", "Created at", "created_at")
        return parse_contact_datetime(value, "Data de criacao", fallback)

    def insert_contact(conn, row, current_time, default_stage):
        phone = row_phone(row)
        name = row_text(row, "Nome", "Cliente", "Contato")
        if not phone and not name:
            return False
        message = row_text(row, "Mensagem / nota", "Mensagem", "Nota")
        last_message = row_text(row, "Ultima mensagem", "Ultima nota") or message[:240]
        created_at = row_created_at(row, current_time)
        cur = conn.execute(
            """
            INSERT INTO conversations
            (name, phone, stage, kind, source, owner, last_message, notes, property_codes, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                name or phone or "Sem nome",
                phone,
                resolve_stage(row_text(row, "Etapa", "Coluna", "Status"), conn, default_stage),
                row_text(row, "Tipo", "Categoria") or "Cliente",
                row_text(row, "Origem", "Fonte") or "Importacao",
                row_text(row, "Responsavel", "Corretor responsavel"),
                last_message,
                row_text(row, "Observacoes", "Observacao", "Notas gerais"),
                row_text(row, "Codigos imoveis", "Codigos", "Imoveis"),
                created_at,
                current_time,
            ),
        )
        if message:
            conn.execute(
                "INSERT INTO messages (conversation_id, direction, body, created_at) VALUES (?, ?, ?, ?)",
                (
                    cur.lastrowid,
                    parse_message_direction(row_text(row, "Direcao da mensagem", "Direcao", "Tipo da mensagem")),
                    message,
                    current_time,
                ),
            )
        return True

    def update_contact(conn, target, row, current_time):
        updates = []
        values = []

        field_specs = [
            ("name", ("Nome", "Cliente", "Contato"), clean_text, target["name"]),
            ("phone", ("Telefone", "WhatsApp", "Numero WPP", "Numero", "Celular"), clean_phone, target["phone"]),
            ("kind", ("Tipo", "Categoria"), clean_text, target["kind"]),
            ("source", ("Origem", "Fonte"), clean_text, target["source"]),
            ("owner", ("Responsavel", "Corretor responsavel"), clean_text, target["owner"]),
            ("notes", ("Observacoes", "Observacao", "Notas gerais"), clean_text, target["notes"]),
            ("property_codes", ("Codigos imoveis", "Codigos", "Imoveis"), clean_text, target["property_codes"]),
        ]
        for field, aliases, cleaner, fallback in field_specs:
            if find_import_col(cols_norm, *aliases) is None:
                continue
            value = cleaner(get_import_col(row, cols_norm, *aliases))
            if field in {"name", "kind"} and not value:
                value = fallback
            updates.append(f"{field} = ?")
            values.append(value)

        stage_col = find_import_col(cols_norm, "Etapa", "Coluna", "Status")
        if stage_col is not None:
            stage_value = row_text(row, "Etapa", "Coluna", "Status")
            if stage_value:
                updates.append("stage = ?")
                values.append(resolve_stage(stage_value, conn, target["stage"]))

        created_at_col = find_import_col(cols_norm, "Data de criacao", "Data de criação", "Criado em", "Created at", "created_at")
        if created_at_col is not None:
            created_at = row_created_at(row, None)
            if created_at:
                updates.append("created_at = ?")
                values.append(created_at)

        message = row_text(row, "Mensagem / nota", "Mensagem", "Nota")
        last_message_col = find_import_col(cols_norm, "Ultima mensagem", "Ultima nota")
        if last_message_col is not None or message:
            last_message = row_text(row, "Ultima mensagem", "Ultima nota") if last_message_col is not None else ""
            updates.append("last_message = ?")
            values.append(last_message or message[:240])

        if not updates and not message:
            return False

        if updates:
            updates.append("updated_at = ?")
            values.extend([current_time, target["id"]])
            conn.execute(f"UPDATE conversations SET {', '.join(updates)} WHERE id = ?", values)

        if message:
            conn.execute(
                "INSERT INTO messages (conversation_id, direction, body, created_at) VALUES (?, ?, ?, ?)",
                (
                    target["id"],
                    parse_message_direction(row_text(row, "Direcao da mensagem", "Direcao", "Tipo da mensagem")),
                    message,
                    current_time,
                ),
            )
        return True

    now = now_iso()
    with get_db() as conn:
        default_stage = resolve_stage(request.form.get("stage", "ativo"), conn, fallback_stage(conn))
        for index, row in df.iterrows():
            row_number = int(index) + 2
            try:
                action = parse_bulk_action(row.get(action_col) if has_action_col else "", has_action_col)
                if not action:
                    stats["ignored"] += 1
                    continue
                if action == "?":
                    add_error(row_number, "acao invalida. Use A, E, I ou deixe em branco.")
                    continue

                if action == "I":
                    if insert_contact(conn, row, now, default_stage):
                        stats["inserted"] += 1
                    else:
                        stats["ignored"] += 1
                    continue

                target = find_contact_for_import(conn, row, cols_norm)
                if target is None:
                    add_error(row_number, "contato nao encontrado para alterar/excluir. Informe ID ou telefone.")
                    continue

                if action == "E":
                    conn.execute("DELETE FROM messages WHERE conversation_id = ?", (target["id"],))
                    conn.execute("DELETE FROM conversations WHERE id = ?", (target["id"],))
                    stats["deleted"] += 1
                    continue

                if update_contact(conn, target, row, now):
                    stats["updated"] += 1
                else:
                    stats["ignored"] += 1
            except ValueError as exc:
                add_error(row_number, str(exc))

    return jsonify({"ok": True, "imported": stats["inserted"], **stats, "errors": errors})


# Sincronizacao de Leads (Imobzi -> CRM), aba CRM, botao "Sincronizar Leads"
_leads_sync_module = None


def _get_leads_sync_module():
    """Carrega dinamicamente AUTOMAÇÔES/Relatório_Leads/atualizar_crm_v0.py.
    Carregado sob demanda (e nao no topo do arquivo) porque esse script tem
    suas proprias dependencias pesadas (selenium) e pode nao existir em toda
    instalacao do app."""
    global _leads_sync_module
    if _leads_sync_module is not None:
        return _leads_sync_module
    if not LEADS_SYNC_SCRIPT_PATH.exists():
        raise RuntimeError(f"Script nao encontrado: {LEADS_SYNC_SCRIPT_PATH}")
    spec = importlib.util.spec_from_file_location("atualizar_crm_v0", LEADS_SYNC_SCRIPT_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    _leads_sync_module = module
    return _leads_sync_module


leads_sync_lock = threading.Lock()
leads_sync_state = {
    "status": "idle",  # idle | running | done | error
    "logs": [],
    "result": None,
    "error": "",
    "started_at": "",
    "finished_at": "",
}
leads_sync_stop_event = threading.Event()
leads_sync_thread = None


def add_leads_sync_log(message, kind="info"):
    timestamp = datetime.now().strftime("%H:%M:%S")
    with leads_sync_lock:
        leads_sync_state["logs"].append({"t": timestamp, "msg": message, "k": kind})
        leads_sync_state["logs"] = leads_sync_state["logs"][-500:]


def run_leads_sync():
    try:
        module = _get_leads_sync_module()
        module.set_progress_callback(add_leads_sync_log)
        resultado = module.sincronizar_leads(
            pasta_relatorio=str(LEADS_SYNC_SCRIPT_PATH.parent),
            db_path=str(DB_PATH),
            stop_event=leads_sync_stop_event,
        )
        with leads_sync_lock:
            leads_sync_state["status"] = "done"
            leads_sync_state["result"] = resultado
            leads_sync_state["finished_at"] = now_iso()
    except Exception as exc:
        add_leads_sync_log(f"Erro: {exc}", "error")
        with leads_sync_lock:
            leads_sync_state["status"] = "error"
            leads_sync_state["error"] = str(exc)
            leads_sync_state["finished_at"] = now_iso()


@app.post("/api/crm/sync_leads/start")
def api_crm_sync_leads_start():
    global leads_sync_thread
    with leads_sync_lock:
        if leads_sync_state["status"] == "running":
            return jsonify({"ok": False, "error": "Sincronizacao ja esta em andamento"}), 400
        leads_sync_state.update(
            {
                "status": "running",
                "logs": [],
                "result": None,
                "error": "",
                "started_at": now_iso(),
                "finished_at": "",
            }
        )
    leads_sync_stop_event.clear()
    leads_sync_thread = threading.Thread(target=run_leads_sync, daemon=True)
    leads_sync_thread.start()
    return jsonify({"ok": True})


@app.post("/api/crm/sync_leads/stop")
def api_crm_sync_leads_stop():
    with leads_sync_lock:
        if leads_sync_state["status"] != "running":
            return jsonify({"ok": False, "error": "Nenhuma sincronizacao em andamento"}), 400
    leads_sync_stop_event.set()
    add_leads_sync_log("Solicitacao de cancelamento enviada...", "warn")
    return jsonify({"ok": True})


@app.get("/api/crm/sync_leads/status")
def api_crm_sync_leads_status():
    since = int(request.args.get("since", 0) or 0)
    with leads_sync_lock:
        logs = leads_sync_state["logs"][since:]
        return jsonify(
            {
                "ok": True,
                "status": leads_sync_state["status"],
                "result": leads_sync_state["result"],
                "error": leads_sync_state["error"],
                "started_at": leads_sync_state["started_at"],
                "finished_at": leads_sync_state["finished_at"],
                "logs": logs,
                "log_index": len(leads_sync_state["logs"]),
                "selenium_available": SELENIUM_AVAILABLE,
            }
        )


# WhatsApp blaster
blaster_lock = threading.Lock()
blaster_state = {
    "status": "idle",
    "total": 0,
    "sent": 0,
    "failed": 0,
    "current": "",
    "countdown": 0,
    "delay_min": 10,
    "delay_max": 20,
    "logs": [],
}
pause_event = threading.Event()
stop_event = threading.Event()
pause_event.set()
bot_thread = None
driver = None
uploaded_contacts_path = None


def add_blaster_log(message, kind="info"):
    timestamp = datetime.now().strftime("%H:%M:%S")
    with blaster_lock:
        blaster_state["logs"].append({"t": timestamp, "msg": message, "k": kind})
        blaster_state["logs"] = blaster_state["logs"][-500:]


def set_blaster_status(status):
    with blaster_lock:
        blaster_state["status"] = status


def reset_blaster_counts():
    with blaster_lock:
        blaster_state.update(
            {
                "status": "idle",
                "total": 0,
                "sent": 0,
                "failed": 0,
                "current": "",
                "countdown": 0,
                "logs": [],
            }
        )


def setup_driver():
    if not SELENIUM_AVAILABLE:
        detail = f" Detalhe: {SELENIUM_IMPORT_ERROR}" if SELENIUM_IMPORT_ERROR else ""
        raise RuntimeError(
            "Selenium nao esta instalado ou esta incompleto neste Python."
            f"{detail} Feche o app e inicie novamente pelo INICIAR_APP.bat."
        )
    options = Options()
    options.add_argument("--start-maximized")
    options.add_experimental_option("detach", True)
    try:
        return webdriver.Chrome(options=options)
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "Selenium esta incompleto neste Python. Feche o app e rode INICIAR_APP.bat "
            "para reinstalar as dependencias."
        ) from exc


def interruptible_sleep(seconds, countdown=False):
    remaining = float(seconds)
    with blaster_lock:
        if countdown:
            blaster_state["countdown"] = int(remaining)
    while remaining > 0:
        if stop_event.is_set():
            break
        if not pause_event.is_set():
            set_blaster_status("paused")
            time.sleep(0.4)
            continue
        if blaster_state.get("status") == "paused":
            set_blaster_status("running")
        chunk = min(0.5, remaining)
        time.sleep(chunk)
        remaining -= chunk
        if countdown:
            with blaster_lock:
                blaster_state["countdown"] = max(0, int(remaining))
    if countdown:
        with blaster_lock:
            blaster_state["countdown"] = 0


BLASTER_PHONE_COLUMN_ALIASES = {
    "numero wpp",
    "telefone",
    "whatsapp",
    "numero",
    "celular",
    "phone",
}
BLASTER_NAME_COLUMN_ALIASES = {"nome", "name", "contato", "cliente"}
BLASTER_METADATA_COLUMNS = {
    "nome",
    "name",
    "contato",
    "cliente",
    "etapa",
    "stage",
    "coluna",
    "status",
    "tipo",
    "origem",
    "responsavel",
    "lead id",
    "demanda hash",
}


def normalize_blaster_column(column):
    return normalize_key(strip_html_to_text(column) or column)


def is_blaster_metadata_column(column):
    key = normalize_blaster_column(column)
    return (
        key in BLASTER_METADATA_COLUMNS
        or key.startswith("nome ")
        or key.endswith(" nome")
        or key.startswith("etapa ")
        or key.endswith(" etapa")
    )


def find_blaster_column(columns, aliases, fallback=None):
    normalized = [normalize_blaster_column(column) for column in columns]
    for index, column in enumerate(normalized):
        if column in aliases:
            return index
    return fallback


def find_blaster_phone_column(columns):
    return find_blaster_column(columns, BLASTER_PHONE_COLUMN_ALIASES, 0)


def find_blaster_name_column(columns):
    return find_blaster_column(columns, BLASTER_NAME_COLUMN_ALIASES)


def is_blaster_message_column(column):
    key = normalize_blaster_column(column)
    return (
        key == "mensagem"
        or key.startswith("mensagem ")
        or key.startswith("paragrafo")
        or key.startswith("texto")
    )


def blaster_message_column_indexes(columns, phone_index, name_index):
    ignored = {phone_index}
    if name_index is not None:
        ignored.add(name_index)
    explicit = [
        index
        for index, column in enumerate(columns)
        if index not in ignored and is_blaster_message_column(column)
    ]
    if explicit:
        return explicit
    return [
        index
        for index, column in enumerate(columns)
        if index not in ignored and not is_blaster_metadata_column(column)
    ]


def row_value(row, index):
    if index is None:
        return ""
    try:
        return row.iloc[index]
    except Exception:
        return ""


def apply_blaster_variables(message, name="", phone=""):
    replaced = strip_html_to_text(message).replace("{nome}", clean_text(name)).replace("{telefone}", clean_phone(phone))
    return sanitize_outbound_message(replaced, name=name)


def build_contact_message(row):
    columns = list(row.index)
    phone_index = find_blaster_phone_column(columns)
    name_index = find_blaster_name_column(columns)
    message_indexes = blaster_message_column_indexes(columns, phone_index, name_index)
    phone = clean_phone(row_value(row, phone_index))
    name = clean_text(row_value(row, name_index))
    message_parts = [clean_text(row_value(row, index)) for index in message_indexes]
    message = "\n".join([part for part in message_parts if part])
    message = apply_blaster_variables(message, name=name, phone=phone)
    return phone, message


def send_whatsapp_message(web_driver, phone, message):
    try:
        web_driver.get(f"https://web.whatsapp.com/send?phone={phone}&text={quote(message)}")
        add_blaster_log(f"Abrindo conversa com {phone}...", "info")
        WebDriverWait(web_driver, 25).until(
            EC.presence_of_element_located((By.CSS_SELECTOR, "div[contenteditable='true']"))
        )
        button = WebDriverWait(web_driver, 25).until(
            EC.element_to_be_clickable(
                (
                    By.CSS_SELECTOR,
                    "button[aria-label='Enviar'], button[aria-label='Send'], span[data-icon='send']",
                )
            )
        )
        button.click()
        with blaster_lock:
            blaster_state["sent"] += 1
        add_blaster_log(f"Mensagem enviada para {phone}", "success")
        interruptible_sleep(2)
    except TimeoutException:
        with blaster_lock:
            blaster_state["failed"] += 1
        add_blaster_log(f"Timeout para {phone}. Numero invalido ou WhatsApp lento.", "warn")
    except Exception as exc:
        with blaster_lock:
            blaster_state["failed"] += 1
        add_blaster_log(f"Erro com {phone}: {exc}", "error")


def bot_runner(df, qr_wait, delay_min, delay_max):
    global driver
    try:
        set_blaster_status("starting")
        add_blaster_log("Iniciando Chrome...", "info")
        driver = setup_driver()
        driver.get("https://web.whatsapp.com")
        set_blaster_status("waiting_qr")
        add_blaster_log(f"Escaneie o QR Code. Aguardando {qr_wait}s.", "warn")
        interruptible_sleep(qr_wait)
        if stop_event.is_set():
            set_blaster_status("idle")
            add_blaster_log("Parado durante espera do QR.", "warn")
            return
        with blaster_lock:
            blaster_state.update(
                {
                    "status": "running",
                    "total": len(df),
                    "sent": 0,
                    "failed": 0,
                    "delay_min": delay_min,
                    "delay_max": delay_max,
                }
            )
        add_blaster_log(f"Intervalo entre envios: {delay_min}s a {delay_max}s.", "success")
        rows = list(df.iterrows())
        for index, (_, row) in enumerate(rows):
            if stop_event.is_set():
                add_blaster_log("Disparos interrompidos.", "warn")
                break
            while not pause_event.is_set() and not stop_event.is_set():
                set_blaster_status("paused")
                time.sleep(0.4)
            if stop_event.is_set():
                break
            set_blaster_status("running")
            phone, message = build_contact_message(row)
            with blaster_lock:
                blaster_state["current"] = phone
            if not phone or not message:
                with blaster_lock:
                    blaster_state["failed"] += 1
                add_blaster_log("Linha ignorada: telefone ou mensagem vazio.", "warn")
            else:
                send_whatsapp_message(driver, phone, message)
            if index < len(rows) - 1 and not stop_event.is_set():
                interruptible_sleep(random.randint(delay_min, delay_max), countdown=True)
        set_blaster_status("idle" if stop_event.is_set() else "done")
        add_blaster_log("Processo finalizado.", "success")
    except Exception as exc:
        set_blaster_status("error")
        add_blaster_log(f"Erro geral: {exc}", "error")
    finally:
        with blaster_lock:
            blaster_state["current"] = ""
            blaster_state["countdown"] = 0


@app.post("/api/blaster/upload")
def api_blaster_upload():
    global uploaded_contacts_path
    file = request.files.get("file")
    if not file:
        return jsonify({"ok": False, "error": "Nenhum arquivo enviado"}), 400
    path = save_upload(file, "disparos")
    try:
        df = read_tabular_file(path)
        uploaded_contacts_path = path
        reset_blaster_counts()
        preview = df.head(6).fillna("").map(format_preview_cell)
        return jsonify(
            {
                "ok": True,
                "rows": len(df),
                "columns": [str(col) for col in preview.columns],
                "preview": preview.values.tolist(),
            }
        )
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500


@app.post("/api/blaster/from_crm")
def api_blaster_from_crm():
    global uploaded_contacts_path
    payload = request.get_json(silent=True) or {}
    selected_ids = payload.get("contact_ids") or []
    try:
        selected_ids = [int(item) for item in selected_ids]
    except Exception:
        selected_ids = []
    stage = clean_text(payload.get("stage")) or "ativo"
    message_template = clean_text(payload.get("message"))

    with get_db() as conn:
        stage = resolve_stage(stage, conn, "ativo")
        if selected_ids:
            placeholders = ",".join(["?"] * len(selected_ids))
            rows = conn.execute(
                f"""
                SELECT * FROM conversations
                WHERE id IN ({placeholders}) AND stage = ? AND phone <> ''
                ORDER BY name COLLATE NOCASE ASC
                """,
                [*selected_ids, stage],
            ).fetchall()
        else:
            rows = conn.execute(
                """
                SELECT * FROM conversations
                WHERE stage = ? AND phone <> ''
                ORDER BY name COLLATE NOCASE ASC
                """,
                (stage,),
            ).fetchall()

    prepared = []
    skipped = 0
    for row in rows:
        contact = conversation_row_to_dict(row)
        message = message_template or contact["last_message"] or contact["notes"]
        primeiro_nome = clean_text(contact["name"]).split(" ")[0]
        message = apply_blaster_variables(message, name=primeiro_nome, phone=contact["phone"])
        if not contact["phone"] or not clean_text(message):
            skipped += 1
            continue
        prepared.append(
            {
                "Numero WPP": clean_phone(contact["phone"]),
                "Mensagem": message,
            }
        )

    if not prepared:
        return jsonify({"ok": False, "error": "Nenhum contato nesta coluna com telefone e mensagem"}), 400

    df = pd.DataFrame(prepared)
    path = UPLOAD_DIR / f"crm_{stage}_disparos_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Disparos CRM")

    uploaded_contacts_path = path
    reset_blaster_counts()
    preview = df.head(8).fillna("").astype(str)
    return jsonify(
        {
            "ok": True,
            "rows": len(df),
            "skipped": skipped,
            "columns": [str(col) for col in preview.columns],
            "preview": preview.values.tolist(),
        }
    )


@app.get("/api/blaster/template")
def api_blaster_template():
    df = pd.DataFrame(
        [
            {
                "Numero WPP": "5511999999999",
                "Paragrafo 1": "Bom dia {nome}, tudo bem?",
                "Nome": "Maria Silva",
                "Paragrafo 2": "Tenho uma oportunidade alinhada ao seu perfil.",
                "Paragrafo 3": "Fico a disposicao para conversar.",
            }
        ]
    )
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Disparos")
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="modelo_disparos_whatsapp.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.post("/api/blaster/optimize")
def api_blaster_optimize():
    if not uploaded_contacts_path:
        return jsonify({"ok": False, "error": "Nenhuma planilha carregada"}), 400
    try:
        df = read_tabular_file(uploaded_contacts_path).fillna("")
        if len(df.columns) < 2:
            return jsonify({"ok": False, "error": "Planilha precisa de telefone e mensagem"}), 400
        columns = list(df.columns)
        phone_index = find_blaster_phone_column(columns)
        name_index = find_blaster_name_column(columns)
        message_indexes = blaster_message_column_indexes(columns, phone_index, name_index)
        if not message_indexes:
            return jsonify({"ok": False, "error": "Planilha precisa de uma coluna de mensagem"}), 400
        optimized = pd.DataFrame()
        phones = []
        messages = []
        names = []
        for _, row in df.iterrows():
            phone, message = build_contact_message(row)
            phones.append(phone)
            messages.append(message)
            if name_index is not None:
                names.append(clean_text(row_value(row, name_index)))
        optimized["Numero WPP"] = phones
        optimized["Mensagem"] = messages
        if name_index is not None:
            optimized["Nome"] = names
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            optimized.to_excel(writer, index=False, sheet_name="Disparos")
        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name="contatos_otimizados.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500


@app.post("/api/blaster/start")
def api_blaster_start():
    global bot_thread
    if not uploaded_contacts_path:
        return jsonify({"ok": False, "error": "Nenhuma planilha carregada"}), 400
    with blaster_lock:
        if blaster_state["status"] in {"starting", "waiting_qr", "running", "paused"}:
            return jsonify({"ok": False, "error": "Bot ja esta em execucao"}), 400
    payload = request.get_json(silent=True) or {}
    qr_wait = int(payload.get("qr_wait") or 60)
    delay_min = int(payload.get("delay_min") or 10)
    delay_max = int(payload.get("delay_max") or 20)
    if delay_max < delay_min:
        delay_max = delay_min
    df = read_tabular_file(uploaded_contacts_path).fillna("")
    stop_event.clear()
    pause_event.set()
    with blaster_lock:
        blaster_state.update(
            {
                "status": "starting",
                "total": len(df),
                "sent": 0,
                "failed": 0,
                "current": "",
                "countdown": 0,
                "delay_min": delay_min,
                "delay_max": delay_max,
                "logs": [],
            }
        )
    bot_thread = threading.Thread(
        target=bot_runner, args=(df, qr_wait, delay_min, delay_max), daemon=True
    )
    bot_thread.start()
    return jsonify({"ok": True})


@app.post("/api/blaster/pause")
def api_blaster_pause():
    if pause_event.is_set():
        pause_event.clear()
        set_blaster_status("paused")
        add_blaster_log("Disparos pausados.", "warn")
        return jsonify({"ok": True, "action": "paused"})
    pause_event.set()
    set_blaster_status("running")
    add_blaster_log("Disparos retomados.", "success")
    return jsonify({"ok": True, "action": "resumed"})


@app.post("/api/blaster/stop")
def api_blaster_stop():
    stop_event.set()
    pause_event.set()
    add_blaster_log("Solicitacao de parada enviada.", "warn")
    set_blaster_status("idle")
    return jsonify({"ok": True})


@app.get("/api/blaster/status")
def api_blaster_status():
    since = int(request.args.get("since", 0) or 0)
    with blaster_lock:
        logs = blaster_state["logs"][since:]
        return jsonify(
            {
                "ok": True,
                **{key: value for key, value in blaster_state.items() if key != "logs"},
                "logs": logs,
                "log_index": len(blaster_state["logs"]),
                "selenium_available": SELENIUM_AVAILABLE,
            }
        )


# Property selector
# Prefixo do link publico de cada imovel (o codigo e concatenado no final).
# Ex.: PROPERTY_BASE_URL=https://www.seusite.com.br/imovel/
PROPERTY_BASE_URL = os.environ.get("PROPERTY_BASE_URL", "").strip()
property_lock = threading.Lock()
property_df = None
property_source = None
last_property_results = []
last_property_filters = {}

# Seleção Externa (planilha do Vivareal/CODMAX) - fonte de dados separada
# da "Seleção Interna" acima, carregada da pasta "Atualização externa". Os
# dados brutos (endereco, url, ID, m², status, valor) são bem mais simples
# que a planilha do Imobzi, então load_external_property_data() os adapta
# para os mesmos nomes de coluna que o restante do pipeline de busca/PDF/
# apresentação web já entende via COL_ALIASES/get_cell/get_item - assim
# filter_properties, get_property_headline, build_property_summary,
# build_marketing_text etc. funcionam sem precisar de versões separadas.
external_property_lock = threading.Lock()
external_property_df = None
external_property_source = None
external_property_source_mtime = None
last_external_property_results = []
last_external_property_filters = {}

TIPO_IMOVEL_VIVAREAL_RE = re.compile(r"/imovel/([a-z0-9]+)-", re.I)

# As URLs do Vivareal nao tem acento (slug ASCII) - corrige de volta os
# tipos mais comuns pra ficar legivel ("galpao" -> "Galpão").
TIPO_IMOVEL_ACENTOS = {
    "galpao": "Galpão",
    "condominio": "Condomínio",
    "predio": "Prédio",
    "chacara": "Chácara",
    "sitio": "Sítio",
}


def extrair_tipo_do_slug_vivareal(url):
    match = TIPO_IMOVEL_VIVAREAL_RE.search(clean_text(url))
    if not match:
        return ""
    token = match.group(1).lower()
    return TIPO_IMOVEL_ACENTOS.get(token, title_case_if_lower(token))


def _raw_col(raw, *names):
    index = {normalize_key(col): col for col in raw.columns}
    for name in names:
        actual = index.get(normalize_key(name))
        if actual is not None:
            return raw[actual]
    return pd.Series([""] * len(raw), index=raw.index)


def load_external_property_data(path):
    raw = read_tabular_file(path).copy()
    unnamed_cols = [col for col in raw.columns if str(col).startswith("Unnamed:")]
    if unnamed_cols:
        raw = raw.drop(columns=unnamed_cols)

    urls = _raw_col(raw, "url", "URL", "Link", "Link do imovel").apply(clean_text)
    status_norm = _raw_col(raw, "status", "Status").apply(lambda v: normalize_key(clean_text(v)))
    valores = _raw_col(raw, "valor", "Valor")
    enderecos = _raw_col(raw, "endereco", "Endereço", "Endereco").apply(clean_text)

    df = pd.DataFrame(index=raw.index)
    df["Codigo"] = _raw_col(raw, "ID", "Id", "Codigo").apply(clean_text)
    df["Link do imovel"] = urls
    # Sem Bairro/Cidade na planilha do Vivareal - usa o endereço coletado
    # como aproximação, só para dar contexto de localização nos textos
    # (título, resumo, link do mapa). É a informação de local mais
    # granular que essa fonte oferece.
    df["Endereco"] = enderecos
    df["Bairro"] = enderecos
    df["Area"] = _raw_col(raw, "m²", "m2", "Area", "Área")
    df["Valor de Venda"] = valores.where(status_norm == "venda")
    df["Valor de Locacao"] = valores.where(status_norm.isin(["aluguel", "temporada"]))
    df["Tipo do Imovel"] = urls.apply(extrair_tipo_do_slug_vivareal)
    df["Finalidade"] = status_norm.map({"aluguel": "Locação", "venda": "Venda", "temporada": "Temporada"}).fillna("")
    df["Status Vivareal"] = _raw_col(raw, "status", "Status")
    # Tudo que está na planilha acabou de ser coletado do site nesta
    # rodada (era um resultado de busca ativo no Vivareal na hora da
    # coleta) - por isso "Ativo" é sempre "Sim" aqui, ao contrário da
    # planilha do Imobzi, que tem uma coluna própria de status comercial.
    df["Ativo"] = "Sim"

    df["__area_filtro"] = df.apply(choose_area, axis=1)
    prices = df.apply(lambda row: choose_price(row, "Automatico"), axis=1)
    df["__preco_auto"] = [item[0] for item in prices]
    df["__tipo_preco_auto"] = [item[1] for item in prices]
    df["__cep_limpo"] = ""
    # A planilha do Vivareal não tem bairro/cidade em coluna própria - só o
    # link do anúncio carrega essa informação no slug da URL (ex.:
    # ".../galpao-...-vila-sao-pedro-bairros-guarulhos-160-aluguel-.../").
    # O filtro "Região/CEP" busca texto por substring em __regiao_texto, e
    # normalize_key() já troca hífen por espaço - então "guarulhos" ou
    # "vila sao pedro" batem direto contra o link. Endereço entra junto
    # para não perder o que a rua já cobre.
    df["__regiao_texto"] = (df["Endereco"].fillna("") + " " + df["Link do imovel"].fillna("")).astype(str)
    return df


def external_property_type_options(df):
    if df is None:
        return ["Todos"]
    values = [clean_text(get_cell(row, "tipo")) for _, row in df.iterrows()]
    return ["Todos"] + unique_property_types(v for v in values if v)


def latest_external_property_file():
    if not EXTERNAL_SELECTION_DIR.exists():
        return None
    candidatos = [
        path
        for path in EXTERNAL_SELECTION_DIR.glob("*")
        if path.is_file()
        and path.suffix.lower() in {".xlsx", ".xls", ".csv"}
        and not path.name.startswith("~$")
    ]
    if not candidatos:
        return None
    return sorted(candidatos, key=lambda path: path.stat().st_mtime, reverse=True)[0]


def ensure_external_property_data_loaded():
    """Sempre reflete o arquivo mais recente da pasta 'Atualização externa' -
    não confia cegamente no que já está em memória. Cada chamada faz um
    stat() rápido nos arquivos da pasta (barato) e só relê a planilha
    inteira se o arquivo mais recente mudou (nome ou data de modificação)
    desde a última vez. Assim, mesmo que uma planilha nova apareça na pasta
    por fora do botão "Atualizar" (ex.: rodando o script manualmente), a
    próxima vez que a página carregar já mostra ela, sem precisar reiniciar
    o app nem clicar em nada."""
    global external_property_df, external_property_source, external_property_source_mtime
    source = latest_external_property_file()
    if source is None:
        with external_property_lock:
            if external_property_df is not None:
                return external_property_df.copy(), external_property_source
        raise RuntimeError(
            "Nenhuma planilha encontrada em 'Atualização externa'. Clique em 'Atualizar' para gerar uma."
        )

    mtime = source.stat().st_mtime
    with external_property_lock:
        if (
            external_property_df is not None
            and external_property_source == str(source)
            and external_property_source_mtime == mtime
        ):
            return external_property_df.copy(), external_property_source

    try:
        df = load_external_property_data(source)
    except PermissionError:
        # Arquivo aberto no Excel (ou outro programa) - o Windows nega a
        # leitura enquanto estiver aberto. Se já tem algo em memória (de
        # antes de o arquivo ser aberto), mostra isso em vez de quebrar a
        # página com um erro; só falha de vez se nunca conseguiu carregar
        # nada ainda.
        with external_property_lock:
            if external_property_df is not None:
                return external_property_df.copy(), external_property_source
        raise RuntimeError(
            f"O arquivo '{source.name}' está aberto no Excel (ou outro programa) e não pôde ser lido. "
            "Feche o arquivo e clique em 'Atualizar' de novo."
        )
    except Exception as exc:
        with external_property_lock:
            if external_property_df is not None:
                return external_property_df.copy(), external_property_source
        raise RuntimeError(f"Não foi possível ler '{source.name}': {exc}")

    with external_property_lock:
        external_property_df = df
        external_property_source = str(source)
        external_property_source_mtime = mtime
    return df.copy(), external_property_source


def reload_external_property_data():
    global external_property_df, external_property_source, external_property_source_mtime
    with external_property_lock:
        external_property_df = None
        external_property_source = None
        external_property_source_mtime = None
    return ensure_external_property_data_loaded()

# Sincronizacao de Imoveis (Imobzi -> planilha unificada), aba Seletor,
# botao "Atualizar".
_property_sync_module = None


def _get_property_sync_module():
    """Carrega dinamicamente AUTOMAÇÔES/Imoveis/Atualizar_imoveis_do_imobizi.py.
    Carregado sob demanda pelos mesmos motivos do modulo de sync de Leads
    (ver _get_leads_sync_module): dependencias pesadas (selenium) que podem
    nao existir em toda instalacao do app."""
    global _property_sync_module
    if _property_sync_module is not None:
        return _property_sync_module
    if not PROPERTY_SYNC_SCRIPT_PATH.exists():
        raise RuntimeError(f"Script nao encontrado: {PROPERTY_SYNC_SCRIPT_PATH}")
    spec = importlib.util.spec_from_file_location("atualizar_imoveis_do_imobizi", PROPERTY_SYNC_SCRIPT_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    _property_sync_module = module
    return _property_sync_module


property_sync_lock = threading.Lock()
property_sync_state = {
    "status": "idle",  # idle | running | done | error
    "logs": [],
    "result": None,
    "error": "",
    "started_at": "",
    "finished_at": "",
}
property_sync_stop_event = threading.Event()
property_sync_thread = None


def add_property_sync_log(message, kind="info"):
    timestamp = datetime.now().strftime("%H:%M:%S")
    with property_sync_lock:
        property_sync_state["logs"].append({"t": timestamp, "msg": message, "k": kind})
        property_sync_state["logs"] = property_sync_state["logs"][-500:]


def run_property_sync():
    try:
        module = _get_property_sync_module()
        module.set_progress_callback(add_property_sync_log)
        resultado = module.sincronizar_imoveis(
            pasta_relatorio=str(PROPERTY_SYNC_SCRIPT_PATH.parent),
            pasta_destino=str(PROPERTY_DIR) if PROPERTY_DIR else None,
            stop_event=property_sync_stop_event,
        )
        with property_sync_lock:
            property_sync_state["status"] = "done"
            property_sync_state["result"] = resultado
            property_sync_state["finished_at"] = now_iso()
    except Exception as exc:
        add_property_sync_log(f"Erro: {exc}", "error")
        with property_sync_lock:
            property_sync_state["status"] = "error"
            property_sync_state["error"] = str(exc)
            property_sync_state["finished_at"] = now_iso()


@app.post("/api/properties/sync/start")
def api_properties_sync_start():
    global property_sync_thread
    with property_sync_lock:
        if property_sync_state["status"] == "running":
            return jsonify({"ok": False, "error": "Sincronizacao ja esta em andamento"}), 400
        property_sync_state.update(
            {
                "status": "running",
                "logs": [],
                "result": None,
                "error": "",
                "started_at": now_iso(),
                "finished_at": "",
            }
        )
    property_sync_stop_event.clear()
    property_sync_thread = threading.Thread(target=run_property_sync, daemon=True)
    property_sync_thread.start()
    return jsonify({"ok": True})


@app.post("/api/properties/sync/stop")
def api_properties_sync_stop():
    with property_sync_lock:
        if property_sync_state["status"] != "running":
            return jsonify({"ok": False, "error": "Nenhuma sincronizacao em andamento"}), 400
    property_sync_stop_event.set()
    add_property_sync_log("Solicitacao de cancelamento enviada...", "warn")
    return jsonify({"ok": True})


@app.get("/api/properties/sync/status")
def api_properties_sync_status():
    since = int(request.args.get("since", 0) or 0)
    with property_sync_lock:
        logs = property_sync_state["logs"][since:]
        return jsonify(
            {
                "ok": True,
                "status": property_sync_state["status"],
                "result": property_sync_state["result"],
                "error": property_sync_state["error"],
                "started_at": property_sync_state["started_at"],
                "finished_at": property_sync_state["finished_at"],
                "logs": logs,
                "log_index": len(property_sync_state["logs"]),
                "selenium_available": SELENIUM_AVAILABLE,
            }
        )


# Sincronizacao da Seleção Externa (roda "VIVAREAL - CODMAX - V1.py" como
# processo separado - script pronto do usuario, nao um modulo importavel
# com callback de progresso como os outros dois sincronizadores acima, por
# isso aqui o log e capturado lendo a saida (stdout) do processo em vez de
# um callback direto).
external_sync_lock = threading.Lock()
external_sync_state = {
    "status": "idle",  # idle | running | done | error
    "logs": [],
    "result": None,
    "error": "",
    "started_at": "",
    "finished_at": "",
}
external_sync_process = None
external_sync_thread = None


def add_external_sync_log(message, kind="info"):
    timestamp = datetime.now().strftime("%H:%M:%S")
    with external_sync_lock:
        external_sync_state["logs"].append({"t": timestamp, "msg": message, "k": kind})
        external_sync_state["logs"] = external_sync_state["logs"][-1000:]


def _classify_external_log_line(line):
    lowered = line.lower()
    if lowered.startswith("[warn]") or "bloqueado" in lowered:
        return "warn"
    if line.startswith("✅"):
        return "success"
    return "info"


def run_external_sync():
    global external_sync_process
    if not EXTERNAL_SCRIPT_PATH.exists():
        message = f"Script nao encontrado: {EXTERNAL_SCRIPT_PATH}"
        add_external_sync_log(message, "error")
        with external_sync_lock:
            external_sync_state["status"] = "error"
            external_sync_state["error"] = message
            external_sync_state["finished_at"] = now_iso()
        return

    env = dict(os.environ)
    env["VIVAREAL_PASTA_SAIDA"] = str(EXTERNAL_SELECTION_DIR)
    env["PYTHONUNBUFFERED"] = "1"

    try:
        process = subprocess.Popen(
            [sys.executable, str(EXTERNAL_SCRIPT_PATH)],
            cwd=str(EXTERNAL_SELECTION_DIR),
            env=env,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            bufsize=1,
            encoding="utf-8",
            errors="replace",
        )
        with external_sync_lock:
            external_sync_process = process

        for linha in process.stdout:
            linha = linha.rstrip("\n").strip()
            if linha:
                add_external_sync_log(linha, _classify_external_log_line(linha))

        returncode = process.wait()
        with external_sync_lock:
            external_sync_process = None

        if returncode != 0:
            raise RuntimeError(f"O script terminou com erro (codigo {returncode}). Veja o log acima.")

        _, source = reload_external_property_data()
        add_external_sync_log(f"Planilha recarregada: {source}", "success")

        with external_sync_lock:
            external_sync_state["status"] = "done"
            external_sync_state["result"] = {"source": source}
            external_sync_state["finished_at"] = now_iso()
    except Exception as exc:
        add_external_sync_log(f"Erro: {exc}", "error")
        with external_sync_lock:
            external_sync_state["status"] = "error"
            external_sync_state["error"] = str(exc)
            external_sync_state["finished_at"] = now_iso()
    finally:
        with external_sync_lock:
            external_sync_process = None


@app.post("/api/properties/external/sync/start")
def api_properties_external_sync_start():
    global external_sync_thread
    with external_sync_lock:
        if external_sync_state["status"] == "running":
            return jsonify({"ok": False, "error": "Sincronizacao ja esta em andamento"}), 400
        external_sync_state.update(
            {
                "status": "running",
                "logs": [],
                "result": None,
                "error": "",
                "started_at": now_iso(),
                "finished_at": "",
            }
        )
    external_sync_thread = threading.Thread(target=run_external_sync, daemon=True)
    external_sync_thread.start()
    return jsonify({"ok": True})


@app.post("/api/properties/external/sync/stop")
def api_properties_external_sync_stop():
    with external_sync_lock:
        process = external_sync_process
        running = external_sync_state["status"] == "running"
    if not running or process is None:
        return jsonify({"ok": False, "error": "Nenhuma sincronizacao em andamento"}), 400

    add_external_sync_log("Solicitacao de cancelamento enviada...", "warn")
    # O script abre um Chrome NOVO por link (as vezes varios) - so encerrar
    # o processo Python (terminate) pode deixar esses Chromes orfaos no
    # Windows. taskkill /T mata a arvore inteira (processo + filhos).
    try:
        subprocess.run(
            ["taskkill", "/PID", str(process.pid), "/T", "/F"],
            capture_output=True,
            check=False,
        )
    except Exception:
        try:
            process.terminate()
        except Exception:
            pass
    return jsonify({"ok": True})


@app.get("/api/properties/external/sync/status")
def api_properties_external_sync_status():
    since = int(request.args.get("since", 0) or 0)
    with external_sync_lock:
        logs = external_sync_state["logs"][since:]
        return jsonify(
            {
                "ok": True,
                "status": external_sync_state["status"],
                "result": external_sync_state["result"],
                "error": external_sync_state["error"],
                "started_at": external_sync_state["started_at"],
                "finished_at": external_sync_state["finished_at"],
                "logs": logs,
                "log_index": len(external_sync_state["logs"]),
            }
        )


FALLBACK_PROPERTY_TYPES = [
    "Galp\u00e3o",
    "Pr\u00e9dio Inteiro",
    "Sala",
    "Galp\u00e3o em Condom\u00ednio",
    "Terreno",
    "Apartamento",
    "Casa",
    "Casa Comercial",
    "Condom\u00ednio Log\u00edstico",
    "Loja/Sal\u00e3o",
]

COL_ALIASES = {
    "codigo": ["codigo", "cod", "cod imovel", "codigo do imovel", "c\u00f3digo"],
    "area_util": ["area util", "\u00e1rea \u00fatil", "area privativa"],
    "area": ["area", "\u00e1rea"],
    "area_total": ["area total", "\u00e1rea total"],
    "valor_venda": ["valor de venda", "valor venda", "venda"],
    "valor_locacao": ["valor de locacao", "valor de loca\u00e7ao", "valor loca\u00e7\u00e3o", "locacao"],
    "cep": ["cep"],
    "bairro": ["bairro"],
    "cidade": ["cidade"],
    "endereco": ["endereco", "endere\u00e7o"],
    "tipo": ["tipo do imovel", "tipo do im\u00f3vel", "tipo"],
    "finalidade": ["finalidade"],
    "ativo": ["ativo"],
    "docas": ["quantidade de docas", "docas"],
}


def get_cell(row, key):
    aliases = COL_ALIASES.get(key, [key])
    index = {normalize_key(col): col for col in row.index}
    for alias in aliases:
        actual = index.get(normalize_key(alias))
        if actual is not None:
            return row.get(actual)
    return ""


def get_item(item, key):
    aliases = COL_ALIASES.get(key, [key])
    index = {normalize_key(col): col for col in item.keys()}
    for alias in aliases:
        actual = index.get(normalize_key(alias))
        if actual is not None:
            return item.get(actual)
    return item.get(key, "")


def unique_property_types(values):
    options = []
    seen = set()
    for value in values:
        label = clean_text(value)
        key = normalize_key(label)
        if not key or key in seen:
            continue
        seen.add(key)
        options.append(label)
    return options


def find_standard_imobizi_files():
    candidates = []
    for root in (BASE_DIR.parent, BASE_DIR):
        if not root.exists():
            continue
        for path in root.glob("*"):
            name = normalize_key(path.stem)
            if (
                path.is_file()
                and path.suffix.lower() in {".xlsx", ".xls", ".csv"}
                and "imobizi" in name
                and "padrao" in name
            ):
                candidates.append(path)
    return sorted(candidates, key=lambda path: path.stat().st_mtime, reverse=True)


def extract_property_types(df):
    values = []
    for _, row in df.iterrows():
        value = clean_text(get_cell(row, "tipo"))
        if value:
            values.append(value)
    return unique_property_types(values)


def get_property_type_options(extra_df=None):
    values = []
    for path in find_standard_imobizi_files():
        try:
            values.extend(extract_property_types(read_tabular_file(path)))
        except Exception:
            continue
    if extra_df is not None:
        values.extend(extract_property_types(extra_df))
    if not values:
        values = FALLBACK_PROPERTY_TYPES
    return ["Todos"] + unique_property_types(values)


def title_case_if_lower(text):
    raw = clean_text(text)
    if not raw:
        return ""
    if raw == raw.upper() or raw == raw.lower():
        return raw.title()
    return raw


def normalize_cep(value):
    return re.sub(r"\D", "", clean_text(value))


def choose_area(row):
    for key in ("area_util", "area", "area_total"):
        value = parse_brazilian_number(get_cell(row, key))
        if pd.notna(value) and value > 0:
            return value
    return math.nan


def get_valid_prices(row):
    venda = parse_brazilian_number(get_cell(row, "valor_venda"))
    locacao = parse_brazilian_number(get_cell(row, "valor_locacao"))
    return {
        "venda": float(venda) if pd.notna(venda) and venda > 0 else math.nan,
        "locacao": float(locacao) if pd.notna(locacao) and locacao > 0 else math.nan,
    }


def choose_price(row, mode):
    prices = get_valid_prices(row)
    has_venda = pd.notna(prices["venda"])
    has_locacao = pd.notna(prices["locacao"])
    if mode == "Venda":
        return (prices["venda"], "venda") if has_venda else (math.nan, "indefinido")
    if mode == "Locacao":
        return (prices["locacao"], "locacao") if has_locacao else (math.nan, "indefinido")
    if has_venda and has_locacao:
        return min(prices["venda"], prices["locacao"]), "venda+locacao"
    if has_venda:
        return prices["venda"], "venda"
    if has_locacao:
        return prices["locacao"], "locacao"
    return math.nan, "indefinido"


def get_availability_label(row):
    prices = get_valid_prices(row)
    has_venda = pd.notna(prices["venda"])
    has_locacao = pd.notna(prices["locacao"])
    if has_venda and has_locacao:
        return "Venda e locação"
    if has_venda:
        return "Somente venda"
    if has_locacao:
        return "Somente locação"
    return "Sem oferta ativa"


def load_property_data(path):
    df = read_tabular_file(path).copy()
    unnamed_cols = [col for col in df.columns if str(col).startswith("Unnamed:")]
    if unnamed_cols:
        df = df.drop(columns=unnamed_cols)
    df["__area_filtro"] = df.apply(choose_area, axis=1)
    prices = df.apply(lambda row: choose_price(row, "Automatico"), axis=1)
    df["__preco_auto"] = [item[0] for item in prices]
    df["__tipo_preco_auto"] = [item[1] for item in prices]
    df["__cep_limpo"] = df.apply(lambda row: normalize_cep(get_cell(row, "cep")), axis=1)
    df["__regiao_texto"] = df.apply(
        lambda row: " ".join(
            [
                clean_text(get_cell(row, "bairro")),
                clean_text(get_cell(row, "cidade")),
                clean_text(get_cell(row, "endereco")),
            ]
        ).strip(),
        axis=1,
    )
    return df


def filter_properties(
    df,
    area_min,
    area_max,
    valor_min,
    valor_max,
    region_or_cep,
    price_mode,
    only_active=False,
    max_results=None,
    tipo_imovel="Todos",
    selected_codes=None,
):
    filtered = df.copy()
    if price_mode == "Automatico":
        filtered["__preco_filtro"] = filtered["__preco_auto"]
        filtered["__tipo_preco"] = filtered["__tipo_preco_auto"]
    else:
        chosen = filtered.apply(lambda row: choose_price(row, price_mode), axis=1)
        filtered["__preco_filtro"] = [item[0] for item in chosen]
        filtered["__tipo_preco"] = [item[1] for item in chosen]

    if selected_codes:
        code_set = {clean_text(code) for code in selected_codes if clean_text(code)}
        filtered = filtered[filtered.apply(lambda row: clean_text(get_cell(row, "codigo")) in code_set, axis=1)]

    if tipo_imovel and tipo_imovel != "Todos":
        wanted = normalize_key(tipo_imovel)
        filtered = filtered[
            filtered.apply(lambda row: wanted in normalize_key(get_cell(row, "tipo")), axis=1)
        ]

    if only_active:
        active_values = {"sim", "yes", "ativo", "1", "true"}
        filtered = filtered[
            filtered.apply(lambda row: normalize_key(get_cell(row, "ativo")) in active_values, axis=1)
        ]

    if area_min is not None:
        filtered = filtered[filtered["__area_filtro"].fillna(-1) >= area_min]
    if area_max is not None:
        filtered = filtered[filtered["__area_filtro"].fillna(float("inf")) <= area_max]
    if valor_min is not None:
        filtered = filtered[filtered["__preco_filtro"].fillna(-1) >= valor_min]
    if valor_max is not None:
        filtered = filtered[filtered["__preco_filtro"].fillna(float("inf")) <= valor_max]

    search = clean_text(region_or_cep)
    if search:
        digits = re.sub(r"\D", "", search)
        if digits:
            filtered = filtered[filtered["__cep_limpo"].astype(str).str.startswith(digits)]
        else:
            needle = normalize_key(search)
            filtered = filtered[
                filtered["__regiao_texto"].fillna("").astype(str).map(normalize_key).str.contains(needle, na=False)
            ]

    filtered = filtered.sort_values(
        by=["__preco_filtro", "__area_filtro"], ascending=[True, True], na_position="last"
    )
    if max_results is not None and max_results > 0:
        filtered = filtered.head(int(max_results))
    return filtered


def row_to_output_dict(row, link_value=None):
    """link_value: quando informado, usa esse link tal como veio na planilha
    (caso da Seleção Externa, onde a planilha já traz o link real do
    anúncio). Quando None (padrão - Seleção Interna), reconstrói o link a
    partir do Código, como sempre foi."""
    data = {}
    for col, value in row.items():
        if str(col).startswith("__"):
            continue
        data[str(col)] = "" if clean_text(value) == "" else value
    code = clean_text(get_cell(row, "codigo"))
    prices = get_valid_prices(row)
    data["Codigo"] = code
    data["Link do imovel"] = clean_text(link_value) if link_value is not None else (PROPERTY_BASE_URL + code if code and PROPERTY_BASE_URL else "")
    data["Area para filtro"] = None if pd.isna(row.get("__area_filtro")) else float(row.get("__area_filtro"))
    data["Valor para filtro"] = None if pd.isna(row.get("__preco_filtro")) else float(row.get("__preco_filtro"))
    data["Tipo de valor usado no filtro"] = clean_text(row.get("__tipo_preco"))
    data["Disponibilidade comercial"] = get_availability_label(row)
    data["Valor de venda numerico"] = None if pd.isna(prices["venda"]) else float(prices["venda"])
    data["Valor de locacao numerico"] = None if pd.isna(prices["locacao"]) else float(prices["locacao"])
    return data


def get_property_headline(item):
    tipo = title_case_if_lower(get_item(item, "tipo")) or "Imóvel comercial"
    bairro = title_case_if_lower(get_item(item, "bairro"))
    cidade = title_case_if_lower(get_item(item, "cidade"))
    code = clean_text(get_item(item, "codigo") or item.get("Codigo"))
    local = " - ".join([part for part in (bairro, cidade) if part])
    return f"{tipo} | {local} | Cód. {code}" if local else f"{tipo} | Cód. {code}"


def build_property_summary(item):
    finalidade = title_case_if_lower(get_item(item, "finalidade"))
    tipo = title_case_if_lower(get_item(item, "tipo")) or "Imóvel"
    bairro = title_case_if_lower(get_item(item, "bairro"))
    cidade = title_case_if_lower(get_item(item, "cidade"))
    endereco = clean_text(get_item(item, "endereco"))
    area = format_area(item.get("Area para filtro"))
    valor = format_currency(item.get("Valor para filtro"))
    disponibilidade = clean_text(item.get("Disponibilidade comercial")) or "Não informado"
    area_total = format_area(get_item(item, "area_total"))
    area_util = format_area(get_item(item, "area_util"))
    cep = clean_text(get_item(item, "cep")) or "Não informado"
    docas = clean_text(get_item(item, "docas")) or "0"
    code = clean_text(item.get("Codigo") or get_item(item, "codigo"))
    finalidade_text = finalidade.lower() if finalidade else "comercial"
    local = bairro or cidade or "localização não informada"
    endereco_text = endereco or "endereço não informado"
    return (
        f"{tipo} {finalidade_text} com {area} na região de {local}. "
        f"Localizado em {endereco_text}, {local}. Disponibilidade: {disponibilidade}. "
        f"Valor de referência: {valor}. Área útil: {area_util}. Área total: {area_total}. "
        f"CEP: {cep}. Docas: {docas}. Código: {code}."
    )


def build_marketing_text(results, client_name="", report_title="Seleção de imóveis"):
    title = clean_text(report_title) or "Seleção de imóveis"
    client = clean_text(client_name)
    lines = [f"{title} para {client}" if client else title, ""]
    if not results:
        lines.append("Nenhum imóvel encontrado para os filtros informados.")
        return "\n".join(lines)
    lines.append(f"Selecionei {len(results)} opção(ões) alinhadas ao perfil solicitado:")
    lines.append("")
    for index, item in enumerate(results, start=1):
        lines.extend(
            [
                f"{index}. {get_property_headline(item)}",
                build_property_summary(item),
                f"Link: {clean_text(item.get('Link do imovel'))}",
                "",
            ]
        )
    lines.append(
        "Se quiser, posso ajustar a seleção por faixa de área, valor, região ou tipo de imóvel."
    )
    return "\n".join(lines)


PRICE_MODE_LABELS = {"Automatico": "Automático", "Venda": "Venda", "Locacao": "Locação"}


def price_mode_label(value):
    text = clean_text(value) or "Automatico"
    return PRICE_MODE_LABELS.get(text, text)


def build_filter_summary(filters):
    parts = [
        f"Área: {format_area(filters.get('area_min')) if filters.get('area_min') is not None else 'Sem min.'} até {format_area(filters.get('area_max')) if filters.get('area_max') is not None else 'Sem máx.'}",
        f"Valor: {format_currency(filters.get('valor_min')) if filters.get('valor_min') is not None else 'Sem min.'} até {format_currency(filters.get('valor_max')) if filters.get('valor_max') is not None else 'Sem máx.'}",
        f"Região/CEP: {clean_text(filters.get('region_or_cep')) or 'Todas'}",
        f"Tipo de imóvel: {clean_text(filters.get('tipo_imovel')) or 'Todos'}",
        f"Modo de preço: {price_mode_label(filters.get('price_mode'))}",
        f"Somente ativos: {'Sim' if filters.get('only_active') else 'Não'}",
    ]
    if filters.get("selected_codes"):
        parts.append("Códigos: " + ", ".join(filters["selected_codes"]))
    if filters.get("max_results"):
        parts.append(f"Limite: {filters['max_results']}")
    return parts


def slugify(text):
    normalized = unicodedata.normalize("NFKD", text or "")
    ascii_text = normalized.encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", ascii_text).strip("-").lower()
    return slug


def item_value_per_area(item):
    valor = item.get("Valor para filtro")
    area = item.get("Area para filtro")
    if valor is None or area is None or area <= 0:
        return None
    return valor / area


def format_dual_price(item):
    parts = []
    venda = item.get("Valor de venda numerico")
    locacao = item.get("Valor de locacao numerico")
    if venda is not None:
        parts.append(f"Venda: {format_currency(venda)}")
    if locacao is not None:
        parts.append(f"Locação: {format_currency(locacao)}")
    return " · ".join(parts) if parts else format_currency(item.get("Valor para filtro"))


def build_maps_link(item):
    endereco = clean_text(get_item(item, "endereco"))
    bairro = title_case_if_lower(get_item(item, "bairro"))
    cidade = title_case_if_lower(get_item(item, "cidade"))
    parts = [part for part in (endereco, bairro, cidade, "Brasil") if part]
    if len(parts) <= 1:
        return ""
    return "https://www.google.com/maps/search/?api=1&query=" + quote(", ".join(parts))


def build_strategic_tag(item, items):
    tags = []
    vpa = item_value_per_area(item)
    all_vpa = [v for v in (item_value_per_area(i) for i in items) if v is not None]
    if vpa is not None and all_vpa and abs(vpa - min(all_vpa)) < 1e-6:
        tags.append("Melhor custo/m²")
    valores = [i.get("Valor para filtro") for i in items if i.get("Valor para filtro") is not None]
    valor = item.get("Valor para filtro")
    if valor is not None and valores and abs(valor - min(valores)) < 1e-6:
        tags.append("Menor investimento")
    areas = [i.get("Area para filtro") for i in items if i.get("Area para filtro") is not None]
    area = item.get("Area para filtro")
    if area is not None and areas and abs(area - max(areas)) < 1e-6:
        tags.append("Maior área")
    return " · ".join(tags) if tags else "Opção equilibrada"


def build_strategic_note(item, items):
    tipo = title_case_if_lower(get_item(item, "tipo")) or "Imóvel"
    local = " - ".join(
        p for p in (title_case_if_lower(get_item(item, "bairro")), title_case_if_lower(get_item(item, "cidade"))) if p
    ) or "região não informada"
    docas = clean_text(get_item(item, "docas"))
    vpa = item_value_per_area(item)
    all_vpa = [v for v in (item_value_per_area(i) for i in items) if v is not None]
    bits = []
    if vpa is not None and all_vpa:
        media = sum(all_vpa) / len(all_vpa)
        if abs(vpa - min(all_vpa)) < 1e-6:
            bits.append("melhor custo por m² do lote analisado")
        elif vpa <= media:
            bits.append("custo por m² dentro da média das opções")
    if docas and docas not in ("0", "0.0"):
        bits.append(f"{docas} doca(s) de carga")
    disponibilidade = clean_text(item.get("Disponibilidade comercial"))
    if disponibilidade:
        bits.append(disponibilidade.lower())
    tail = "; ".join(bits) if bits else "opção equilibrada dentro do perfil solicitado"
    return f"{tipo} em {local}, com {tail}."


def build_highlights(items):
    bullets = []
    vendas = sum(1 for i in items if i.get("Valor de venda numerico") is not None)
    locacoes = sum(1 for i in items if i.get("Valor de locacao numerico") is not None)
    if vendas or locacoes:
        bits = []
        if vendas:
            bits.append(f"{vendas} para venda")
        if locacoes:
            bits.append(f"{locacoes} para locação")
        bullets.append(" e ".join(bits) + " entre as opções analisadas.")
    seen = set()
    regioes = []
    for i in items:
        local = " - ".join(
            p for p in (title_case_if_lower(get_item(i, "bairro")), title_case_if_lower(get_item(i, "cidade"))) if p
        )
        if local and local not in seen:
            seen.add(local)
            regioes.append(local)
    if regioes:
        shown = regioes[:4]
        extra = len(regioes) - len(shown)
        text = "Presença nas regiões de " + ", ".join(shown)
        if extra > 0:
            text += f" e mais {extra} localidade(s)"
        bullets.append(text + ".")
    docas_total = 0
    for i in items:
        docas_raw = clean_text(get_item(i, "docas"))
        try:
            docas_total += int(float(docas_raw)) if docas_raw else 0
        except ValueError:
            continue
    if docas_total > 0:
        bullets.append(f"{docas_total} doca(s) de carga somadas entre as opções, para operações logísticas.")
    areas = [i.get("Area para filtro") for i in items if i.get("Area para filtro")]
    if areas:
        bullets.append(
            f"Áreas entre {format_area(min(areas))} e {format_area(max(areas))}, atendendo operações de diferentes portes."
        )
    return bullets[:4]


def build_recommendations(items):
    recs = []
    with_vpa = [(i, item_value_per_area(i)) for i in items]
    with_vpa = [pair for pair in with_vpa if pair[1] is not None]
    if with_vpa:
        best_item, best_vpa = min(with_vpa, key=lambda pair: pair[1])
        recs.append(
            {
                "label": "Melhor custo-benefício",
                "item": best_item,
                "reason": f"Menor custo por m² do lote analisado: {format_currency(best_vpa)}/m².",
            }
        )
    with_valor = [i for i in items if i.get("Valor para filtro") is not None]
    if with_valor:
        cheapest = min(with_valor, key=lambda i: i.get("Valor para filtro"))
        if cheapest is not (recs[0]["item"] if recs else None):
            recs.append(
                {
                    "label": "Menor valor de entrada",
                    "item": cheapest,
                    "reason": f"Menor valor entre as opções analisadas: {format_dual_price(cheapest)}.",
                }
            )
    rentable = [i for i in items if "loca" in clean_text(i.get("Disponibilidade comercial")).lower()]
    if rentable:
        quick = min(rentable, key=lambda i: i.get("Area para filtro") or float("inf"))
        recs.append(
            {
                "label": "Melhor para operação imediata",
                "item": quick,
                "reason": "Disponível para locação, permitindo início de operação mais rápido.",
            }
        )
    venda_items = [i for i in items if i.get("Valor de venda numerico") is not None]
    venda_with_vpa = [(i, item_value_per_area(i)) for i in venda_items]
    venda_with_vpa = [pair for pair in venda_with_vpa if pair[1] is not None]
    if venda_with_vpa:
        best_buy, _ = min(venda_with_vpa, key=lambda pair: pair[1])
        recs.append(
            {
                "label": "Melhor para compra e investimento",
                "item": best_buy,
                "reason": "Menor custo por m² entre as opções disponíveis para venda.",
            }
        )
    with_area = [i for i in items if (i.get("Area para filtro") or 0) > 0]
    if with_area:
        largest = max(with_area, key=lambda i: i.get("Area para filtro"))
        recs.append(
            {
                "label": "Melhor para expansão",
                "item": largest,
                "reason": f"Maior área útil do lote analisado: {format_area(largest.get('Area para filtro'))}.",
            }
        )
    return recs[:5]


REPORT_INK = "#151a18"
REPORT_PAPER = "#f4f3ef"
REPORT_WHITE = "#ffffff"
REPORT_MUTED = "#67706b"
REPORT_LINE = "#d7dad5"
REPORT_BLUE = "#0073c3"
REPORT_BLUE_DARK = "#054a7a"
REPORT_SKY = "#bfe3ff"
REPORT_ACCENT = "#4fa8dc"

# Contato exibido nas apresentacoes/PDFs gerados. Configure no .env.
PDF_CONTACT_NAME = os.environ.get("PDF_CONTACT_NAME", "").strip() or "Seu Nome"
PDF_CONTACT_BRAND = os.environ.get("PDF_CONTACT_BRAND", "").strip() or "SP Imóveis | Galpões Industriais"
PDF_CONTACT_WHATSAPP_DISPLAY = os.environ.get("PDF_CONTACT_WHATSAPP_DISPLAY", "").strip() or "(11) 90000-0000"
PDF_CONTACT_WHATSAPP_NUMBER = re.sub(r"\D", "", os.environ.get("PDF_CONTACT_WHATSAPP_NUMBER", "")) or "5511900000000"
PDF_CONTACT_CTA = os.environ.get("PDF_CONTACT_CTA", "").strip() or "Tenho interesse nas opções de imóveis industriais apresentadas."

REPORT_LOGO_PATH = BASE_DIR / "static" / "logo_sp_imoveis.png"
_report_logo_data_uri_cache = None

# Social-share preview image (WhatsApp/Facebook/etc. link cards). These crawlers fetch og:image
# over plain HTTP, so a data: URI does not work here — the file has to be copied next to each
# generated index.html (see next_presentation_path/api_properties_export_web) and referenced by a
# relative filename, so it keeps working no matter which domain the .html ends up hosted on.
PRESENTATION_PREVIEW_IMAGE_PATH = BASE_DIR / "static" / "preview-sp-imoveis.png"
PRESENTATION_PREVIEW_IMAGE_FILENAME = PRESENTATION_PREVIEW_IMAGE_PATH.name
PRESENTATION_PREVIEW_IMAGE_SIZE = (1080, 1350)


def report_logo_data_uri():
    """Reads and base64-encodes the SP Imóveis logo once, so it can be embedded inline
    (as a data URI) in the standalone HTML/PDF reports without depending on an external file."""
    global _report_logo_data_uri_cache
    if _report_logo_data_uri_cache is None:
        try:
            data = REPORT_LOGO_PATH.read_bytes()
            _report_logo_data_uri_cache = "data:image/png;base64," + base64.b64encode(data).decode("ascii")
        except OSError:
            _report_logo_data_uri_cache = ""
    return _report_logo_data_uri_cache


PRESENTATIONS_ROOT = BASE_DIR.parent.parent / "Apresentações"
PRESENTATIONS_INDEX_DIR = PRESENTATIONS_ROOT / "INDEX"
PRESENTATIONS_PDF_DIR = PRESENTATIONS_ROOT / "PDF"


def presentation_filename_slug(text):
    """Slug used for saved presentation filenames: lowercase, underscore-separated (e.g. 'Maria Silva' -> 'maria_silva')."""
    normalized = unicodedata.normalize("NFKD", text or "")
    ascii_text = normalized.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-zA-Z0-9]+", "_", ascii_text).strip("_").lower()


def next_presentation_path(folder, client_name, extension):
    """Builds the next available save path for a generated presentation:

    - Named client: '{cliente}_propostav1{ext}', '{cliente}_propostav2{ext}', ... (one counter per client).
    - No client name: 'cliente(1)v1{ext}', 'cliente(2)v1{ext}', ... (a fresh numbered slot each time).

    Never overwrites an existing file; the folder is created if it does not exist yet.
    """
    folder.mkdir(parents=True, exist_ok=True)
    existing_names = [entry.name for entry in folder.iterdir() if entry.is_file()]
    client_slug = presentation_filename_slug(client_name)
    if client_slug:
        pattern = re.compile(rf"^{re.escape(client_slug)}_propostav(\d+){re.escape(extension)}$", re.IGNORECASE)
        versions = [int(m.group(1)) for name in existing_names if (m := pattern.match(name))]
        version = max(versions, default=0) + 1
        filename = f"{client_slug}_propostav{version}{extension}"
    else:
        pattern = re.compile(rf"^cliente\((\d+)\)v\d+{re.escape(extension)}$", re.IGNORECASE)
        slots = [int(m.group(1)) for name in existing_names if (m := pattern.match(name))]
        next_slot = max(slots, default=0) + 1
        filename = f"cliente({next_slot})v1{extension}"
    return folder / filename


def ensure_preview_image_next_to(html_path):
    """Copies the social-share preview image into the same folder as a generated index.html
    (once, if not already there) so it travels together when the .html is uploaded/hosted
    elsewhere — the og:image tag references it by that relative filename."""
    dest = html_path.parent / PRESENTATION_PREVIEW_IMAGE_FILENAME
    if dest.exists() or not PRESENTATION_PREVIEW_IMAGE_PATH.exists():
        return
    try:
        shutil.copy2(PRESENTATION_PREVIEW_IMAGE_PATH, dest)
    except OSError:
        pass


def report_style_block():
    return f"""
<style>
  :root {{
    --ink: {REPORT_INK}; --paper: {REPORT_PAPER}; --white: {REPORT_WHITE};
    --muted: {REPORT_MUTED}; --line: {REPORT_LINE};
    --blue: {REPORT_BLUE}; --blue-dark: {REPORT_BLUE_DARK};
    --sky: {REPORT_SKY}; --accent: {REPORT_ACCENT};
    --max: 1180px; --nav: 72px;
  }}
  * {{ box-sizing: border-box; }}
  html {{ scroll-behavior: smooth; }}
  body {{
    margin: 0; color: var(--ink); background: var(--paper);
    font-family: Inter, ui-sans-serif, system-ui, -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
    line-height: 1.5;
  }}
  a {{ color: inherit; }}
  .wrap {{ width: min(calc(100% - 40px), var(--max)); margin-inline: auto; }}
  .eyebrow {{ margin: 0 0 10px; color: var(--blue); font-size: .76rem; font-weight: 800; letter-spacing: .12em; text-transform: uppercase; }}
  h1, h2, h3, p {{ margin-top: 0; }}
  h2 {{ margin-bottom: 16px; font-family: Georgia, "Times New Roman", serif; font-size: clamp(2rem, 5vw, 3.6rem); line-height: 1.05; font-weight: 500; }}
  .lead {{ max-width: 720px; color: var(--muted); font-size: 1.06rem; }}
  .btn {{
    min-height: 44px; display: inline-flex; align-items: center; justify-content: center; gap: 8px;
    padding: 11px 16px; border: 1px solid transparent; border-radius: 4px;
    background: var(--blue); color: var(--white); font-size: .88rem; font-weight: 760;
    text-decoration: none; transition: transform .2s ease, background .2s ease, border-color .2s ease;
  }}
  .btn:hover {{ transform: translateY(-2px); background: var(--blue-dark); }}
  .btn.alt {{ border-color: #b7c4cd; background: transparent; color: var(--ink); }}
  .btn.alt:hover {{ border-color: var(--ink); background: var(--white); }}
  .btn.light {{ background: var(--white); color: var(--blue-dark); }}
  .nav {{
    position: fixed; inset: 0 0 auto; z-index: 30; height: var(--nav); color: var(--white);
    background: rgba(21, 26, 24, .96); box-shadow: 0 1px 0 rgba(255,255,255,.12);
  }}
  .nav-inner {{ height: 100%; display: flex; align-items: center; justify-content: space-between; gap: 22px; }}
  .brand {{ display: flex; align-items: center; gap: 12px; line-height: 1.04; text-decoration: none; }}
  .brand-logo {{ height: 42px; width: auto; display: block; }}
  .brand strong {{ font-family: Georgia, "Times New Roman", serif; font-size: 1.22rem; font-weight: 500; }}
  .brand span {{ margin-top: 5px; color: #c9d6e0; font-size: .65rem; font-weight: 700; letter-spacing: .11em; text-transform: uppercase; }}
  .nav-links {{ display: flex; align-items: center; gap: 22px; font-size: .82rem; font-weight: 700; }}
  .nav-links a {{ text-decoration: none; }}
  .hero {{
    position: relative; display: flex; align-items: end; padding: calc(var(--nav) + 60px) 0 64px;
    color: var(--white); background: linear-gradient(160deg, var(--blue-dark) 0%, var(--ink) 100%);
    background-size: cover; background-position: center; overflow: hidden;
  }}
  .hero::after {{ content: ""; position: absolute; width: 440px; height: 440px; right: -180px; top: -220px; border: 1px solid rgba(255,255,255,.15); transform: rotate(35deg); }}
  .hero-content {{ position: relative; max-width: 790px; }}
  .hero .eyebrow {{ color: var(--sky); }}
  .hero h1 {{ margin-bottom: 22px; font-family: Georgia, "Times New Roman", serif; font-size: clamp(2.6rem, 7vw, 5.4rem); line-height: .96; font-weight: 500; }}
  .hero p {{ position: relative; max-width: 620px; color: #dce8f2; font-size: clamp(1rem, 2vw, 1.2rem); }}
  .hero-client {{ position: relative; max-width: 620px; color: #dce8f2; font-size: .95rem; margin-bottom: 8px; }}
  .hero-actions {{ position: relative; display: flex; flex-wrap: wrap; gap: 10px; margin-top: 30px; }}
  section {{ padding: 92px 0; }}
  .summary {{ background: var(--white); }}
  .summary-grid {{ display: grid; grid-template-columns: .85fr 1.15fr; gap: clamp(36px, 8vw, 110px); align-items: start; }}
  .facts {{ border-top: 1px solid var(--line); }}
  .fact {{ display: grid; grid-template-columns: 88px 1fr; gap: 18px; padding: 20px 0; border-bottom: 1px solid var(--line); }}
  .fact strong {{ font-family: Georgia, "Times New Roman", serif; color: var(--blue-dark); font-size: 2rem; font-weight: 500; }}
  .fact span {{ align-self: center; color: var(--muted); }}
  .highlights {{ margin-top: 24px; padding: 0; list-style: none; }}
  .highlights li {{ position: relative; padding-left: 18px; margin-bottom: 8px; color: var(--muted); font-size: .92rem; line-height: 1.55; }}
  .highlights li::before {{ content: ""; position: absolute; left: 0; top: 8px; width: 6px; height: 6px; border-radius: 50%; background: var(--accent); }}
  .kpis {{ background: var(--ink); color: var(--white); }}
  .kpi-grid {{ display: grid; grid-template-columns: repeat(4, 1fr); }}
  .kpi {{ min-height: 190px; padding: 30px; border-left: 1px solid #2f4152; display: flex; flex-direction: column; justify-content: space-between; }}
  .kpi:first-child {{ border-left: 0; }}
  .kpi span {{ color: #a9bbc9; font-size: .77rem; font-weight: 750; letter-spacing: .08em; text-transform: uppercase; }}
  .kpi strong {{ font-family: Georgia, "Times New Roman", serif; font-size: clamp(1.6rem, 3.4vw, 2.6rem); font-weight: 500; line-height: 1.1; }}
  .compare {{ padding-bottom: 52px; }}
  .table-scroll {{ margin-top: 34px; overflow-x: auto; border-top: 2px solid var(--ink); }}
  table {{ width: 100%; min-width: 920px; border-collapse: collapse; background: var(--white); }}
  th, td {{ padding: 16px 14px; border-bottom: 1px solid var(--line); text-align: left; vertical-align: top; }}
  th {{ color: var(--muted); font-size: .7rem; letter-spacing: .07em; text-transform: uppercase; }}
  td {{ font-size: .86rem; }}
  .tag {{ display: inline-block; padding: 3px 7px; border-radius: 3px; background: #e3f1fb; color: var(--blue-dark); font-size: .68rem; font-weight: 780; white-space: nowrap; }}
  .inventory {{ padding-top: 52px; }}
  .property {{ display: grid; grid-template-columns: .34fr .66fr; min-height: 420px; border-top: 1px solid var(--line); background: var(--white); }}
  .property:last-child {{ border-bottom: 1px solid var(--line); }}
  .property-visual {{ position: relative; padding: 34px; display: flex; flex-direction: column; justify-content: space-between; overflow: hidden; color: var(--white); background: var(--blue-dark); background-size: cover; background-position: center; text-decoration: none; cursor: pointer; }}
  .property:nth-child(2n) .property-visual {{ background: #28323d; }}
  .property-visual::after {{ content: ""; position: absolute; right: -65px; bottom: -100px; width: 260px; height: 260px; border: 1px solid rgba(255,255,255,.16); transform: rotate(26deg); }}
  .option {{ position: relative; z-index: 1; font-size: .75rem; font-weight: 800; letter-spacing: .1em; text-transform: uppercase; }}
  .location {{ position: relative; z-index: 1; }}
  .location strong {{ display: block; font-family: Georgia, "Times New Roman", serif; font-size: clamp(1.7rem, 4vw, 2.7rem); line-height: 1; font-weight: 500; }}
  .location span {{ display: block; margin-top: 12px; color: #cfe0ee; }}
  .property-body {{ padding: clamp(32px, 6vw, 60px); }}
  .property-title {{ display: flex; align-items: start; justify-content: space-between; gap: 24px; }}
  .property h3 {{ margin-bottom: 8px; font-family: Georgia, "Times New Roman", serif; font-size: clamp(1.6rem, 3.6vw, 2.4rem); line-height: 1.08; font-weight: 500; }}
  .code {{ flex: 0 0 auto; color: var(--muted); font-size: .74rem; font-weight: 750; white-space: nowrap; }}
  .price {{ margin: 22px 0; color: var(--blue-dark); font-family: Georgia, "Times New Roman", serif; font-size: 1.7rem; }}
  .price small {{ display: block; color: var(--muted); font-family: Inter, sans-serif; font-size: .7rem; font-weight: 750; letter-spacing: .08em; text-transform: uppercase; }}
  .metrics {{ display: grid; grid-template-columns: repeat(3, 1fr); margin: 24px 0; border-block: 1px solid var(--line); }}
  .metric {{ padding: 16px 12px 16px 0; }}
  .metric span {{ display: block; color: var(--muted); font-size: .68rem; font-weight: 750; text-transform: uppercase; }}
  .metric strong {{ display: block; margin-top: 4px; }}
  .analysis {{ display: grid; grid-template-columns: 1fr 1fr; gap: 24px; margin-top: 25px; }}
  .analysis h4 {{ margin: 0 0 6px; color: var(--blue-dark); font-size: .76rem; text-transform: uppercase; }}
  .analysis p {{ margin: 0; color: var(--muted); font-size: .88rem; }}
  .actions {{ display: flex; flex-wrap: wrap; gap: 8px; margin-top: 30px; }}
  .recommend {{ background: #eaf3fa; }}
  .ranking {{ counter-reset: rank; border-top: 2px solid var(--ink); background: var(--white); }}
  .rank {{ counter-increment: rank; display: grid; grid-template-columns: 74px 210px 1fr; gap: 24px; padding: 25px 24px; border-bottom: 1px solid var(--line); align-items: center; }}
  .rank::before {{ content: counter(rank, decimal-leading-zero); color: var(--blue); font-family: Georgia, "Times New Roman", serif; font-size: 2rem; }}
  .rank strong {{ font-size: 1rem; }}
  .rank p {{ margin: 0; color: var(--muted); font-size: .9rem; }}
  .disclaimer {{ margin-top: 22px; color: var(--muted); font-size: .78rem; }}
  .cta {{ position: relative; padding: 100px 0; color: var(--white); background: var(--blue-dark); overflow: hidden; }}
  .cta::after {{ content: ""; position: absolute; width: 440px; height: 440px; right: -180px; top: -220px; border: 1px solid rgba(255,255,255,.15); transform: rotate(35deg); }}
  .cta-inner {{ position: relative; z-index: 1; display: grid; grid-template-columns: 1fr auto; gap: 40px; align-items: end; }}
  .cta h2 {{ max-width: 760px; margin-bottom: 12px; }}
  .cta p {{ margin-bottom: 0; color: #cddceb; }}
  .cta-logo {{ height: 34px; width: auto; display: block; margin-bottom: 14px; }}
  footer.page-footer {{ padding: 22px 0; color: #c5d3de; background: var(--ink); font-size: .74rem; }}
  .footer-inner {{ display: flex; justify-content: space-between; gap: 20px; }}
  .reveal {{ opacity: 0; transform: translateY(16px); transition: opacity .55s ease, transform .55s ease; }}
  .reveal.visible {{ opacity: 1; transform: none; }}
  @media (max-width: 820px) {{
    :root {{ --nav: 64px; }}
    .wrap {{ width: min(calc(100% - 28px), var(--max)); }}
    .nav-links a:not(.nav-cta) {{ display: none; }}
    section {{ padding: 72px 0; }}
    .summary-grid, .cta-inner {{ grid-template-columns: 1fr; }}
    .kpi-grid {{ grid-template-columns: 1fr 1fr; }}
    .kpi {{ min-height: 160px; border-bottom: 1px solid #2f4152; }}
    .property {{ grid-template-columns: 1fr; }}
    .property-visual {{ min-height: 220px; }}
    .metrics {{ grid-template-columns: 1fr 1fr; }}
    .analysis {{ grid-template-columns: 1fr; }}
    .rank {{ grid-template-columns: 54px 1fr; gap: 12px; }}
    .rank p {{ grid-column: 2; }}
    .footer-inner {{ flex-direction: column; }}
  }}
  @media (prefers-reduced-motion: reduce) {{
    html {{ scroll-behavior: auto; }}
    *, *::before, *::after {{ animation-duration: .01ms !important; transition-duration: .01ms !important; }}
    .reveal {{ opacity: 1; transform: none; }}
  }}
  @media print {{
    .nav, .hero-actions {{ display: none; }}
    .hero {{ padding-top: 40px; -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
    .kpis, .cta, .property-visual {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
    .reveal {{ opacity: 1; transform: none; }}
    .summary, .kpis, .compare, .inventory, .recommend, .cta {{ break-inside: avoid-page; }}
    .property, .rank {{ break-inside: avoid; }}
    section {{ padding: 34px 0; }}
  }}
</style>
"""


def build_region_list(items):
    seen = set()
    regioes = []
    for item in items:
        local = " - ".join(
            p for p in (title_case_if_lower(get_item(item, "bairro")), title_case_if_lower(get_item(item, "cidade"))) if p
        )
        if local and local not in seen:
            seen.add(local)
            regioes.append(local)
    return regioes


def build_attention_note(item):
    faltando = []
    if not clean_text(get_item(item, "area_total")):
        faltando.append("área total")
    if not clean_text(item.get("Disponibilidade comercial")):
        faltando.append("disponibilidade comercial")
    docas = clean_text(get_item(item, "docas"))
    if not docas or docas in ("0", "0.0"):
        faltando.append("docas de carga")
    if faltando:
        return f"Confirmar {', '.join(faltando)} e condições comerciais diretamente com o anunciante antes da visita."
    return "Confirmar disponibilidade, valores e condições comerciais diretamente com o anunciante antes da visita."


def report_nav_html():
    wa_link = html.escape(f"https://wa.me/{PDF_CONTACT_WHATSAPP_NUMBER}?text={quote(PDF_CONTACT_CTA)}", quote=True)
    logo = report_logo_data_uri()
    brand_mark = (
        f'<img class="brand-logo" src="{logo}" alt="SP Imóveis">'
        if logo
        else '<strong>SP Imóveis</strong>'
    )
    return f"""
  <header class="nav" id="nav">
    <div class="wrap nav-inner">
      <a class="brand" href="#inicio" aria-label="SP Imóveis, voltar ao início">
        {brand_mark}
        <span>Galpões Industriais</span>
      </a>
      <nav class="nav-links" aria-label="Navegação principal">
        <a href="#comparativo">Comparativo</a>
        <a href="#imoveis">Imóveis</a>
        <a href="#prioridades">Prioridades</a>
        <a class="btn light nav-cta" href="{wa_link}" target="_blank" rel="noreferrer">WhatsApp</a>
      </nav>
    </div>
  </header>
    """


def report_hero_html(items, title, client, filters):
    areas = [i.get("Area para filtro") for i in items if i.get("Area para filtro") is not None]
    area_line = f" entre {format_area(min(areas))} e {format_area(max(areas))}" if areas else ""
    regiao = clean_text(filters.get("region_or_cep")) if filters else ""
    eyebrow = f"Seleção consultiva · {html.escape(regiao)}" if regiao else "Seleção consultiva de imóveis industriais"
    client_line = (
        f'<p class="hero-client">Proposta preparada para <strong>{html.escape(client)}</strong></p>'
        if client
        else ""
    )
    photo = clean_text(items[0].get("_foto")) if items else ""
    hero_style = ""
    if photo:
        hero_style = (
            ' style="background-image: linear-gradient(90deg, rgba(21,26,24,.92) 0%, rgba(21,26,24,.72) 46%, '
            f"rgba(21,26,24,.2) 82%), url('{html.escape(photo, quote=True)}');\""
        )
    return f"""
    <section class="hero" id="inicio"{hero_style}>
      <div class="wrap hero-content">
        <p class="eyebrow">{eyebrow}</p>
        <h1>{html.escape(title)}</h1>
        {client_line}
        <p>{len(items)} oportunidade(s) confirmada(s){html.escape(area_line)}, organizadas para apoiar sua decisão com clareza.</p>
        <div class="hero-actions">
          <a class="btn light" href="#comparativo">Comparar opções</a>
          <a class="btn" href="#prioridades">Ver prioridades</a>
        </div>
      </div>
    </section>
    """


def report_summary_html(items):
    regioes = build_region_list(items)
    vendas = sum(1 for i in items if i.get("Valor de venda numerico") is not None)
    bullets = "".join(f"<li>{html.escape(b)}</li>" for b in build_highlights(items))
    regioes_shown = ", ".join(regioes[:4]) or "não informadas"
    return f"""
    <section class="summary">
      <div class="wrap summary-grid reveal">
        <div>
          <p class="eyebrow">Resumo da seleção</p>
          <h2>Um recorte compacto e operacional.</h2>
        </div>
        <div>
          <p class="lead">A seleção reúne alternativas em {html.escape(regioes_shown)}, comparadas por área, valor e condições comerciais para apoiar a decisão.</p>
          <ul class="highlights">{bullets}</ul>
          <div class="facts">
            <div class="fact"><strong>{len(items):02d}</strong><span>imóvel(is) confirmado(s) e comparado(s)</span></div>
            <div class="fact"><strong>{len(regioes):02d}</strong><span>região(ões)/bairro(s) considerados</span></div>
            <div class="fact"><strong>{vendas:02d}</strong><span>opção(ões) com venda confirmada</span></div>
          </div>
        </div>
      </div>
    </section>
    """


def report_kpis_html(items):
    valores = [i.get("Valor para filtro") for i in items if i.get("Valor para filtro") is not None]
    areas_totais = [
        v for v in (parse_brazilian_number(get_item(i, "area_total")) for i in items) if not math.isnan(v)
    ]
    vpa_list = [v for v in (item_value_per_area(i) for i in items) if v is not None]
    venda_count = sum(1 for i in items if i.get("Valor de venda numerico") is not None)
    valor_min = format_currency(min(valores)) if valores else "Não informado"
    area_max = format_area(max(areas_totais)) if areas_totais else "Não informado"
    vpa_range = (
        f"{format_currency(min(vpa_list))} – {format_currency(max(vpa_list))}" if vpa_list else "Não informado"
    )
    tile4_label = "Opções para venda" if venda_count else "Imóveis analisados"
    tile4_value = f"{venda_count:02d}" if venda_count else f"{len(items):02d}"
    return f"""
    <section class="kpis" aria-label="Indicadores principais">
      <div class="wrap kpi-grid reveal">
        <div class="kpi"><span>Menor valor</span><strong>{valor_min}</strong></div>
        <div class="kpi"><span>Maior área construída</span><strong>{area_max}</strong></div>
        <div class="kpi"><span>Faixa de valor por m²</span><strong>{vpa_range}</strong></div>
        <div class="kpi"><span>{tile4_label}</span><strong>{tile4_value}</strong></div>
      </div>
    </section>
    """


def report_compare_html(items):
    rows = []
    for index, item in enumerate(items, start=1):
        tipo = html.escape(title_case_if_lower(get_item(item, "tipo")) or "Não informado")
        local = html.escape(
            " - ".join(
                p for p in (title_case_if_lower(get_item(item, "bairro")), title_case_if_lower(get_item(item, "cidade"))) if p
            )
            or "Não informado"
        )
        rows.append(
            f"""
              <tr>
                <td><strong>{index:02d}</strong><br><span class="tag">Cód. {html.escape(clean_text(item.get('Codigo')) or 'N/D')}</span></td>
                <td>{local}</td>
                <td>{tipo}</td>
                <td>{html.escape(format_dual_price(item))}</td>
                <td>{html.escape(format_area(get_item(item, 'area_util')))}</td>
                <td>{html.escape(format_area(get_item(item, 'area_total')))}</td>
                <td>{html.escape(build_strategic_tag(item, items))}</td>
              </tr>
            """
        )
    return f"""
    <section class="compare" id="comparativo">
      <div class="wrap reveal">
        <p class="eyebrow">Tabela comparativa</p>
        <h2>O essencial, lado a lado.</h2>
        <p class="lead">Valores e disponibilidade devem ser reconfirmados antes da visita ou proposta.</p>
        <div class="table-scroll" role="region" aria-label="Comparativo dos imóveis" tabindex="0">
          <table>
            <thead><tr><th>Opção</th><th>Localização</th><th>Tipo</th><th>Valor</th><th>Área útil</th><th>Área total</th><th>Leitura estratégica</th></tr></thead>
            <tbody>{''.join(rows)}</tbody>
          </table>
        </div>
      </div>
    </section>
    """


def report_property_html(index, item, items):
    headline = html.escape(get_property_headline(item))
    link = html.escape(clean_text(item.get("Link do imovel")), quote=True)
    maps_link = build_maps_link(item)
    bairro = title_case_if_lower(get_item(item, "bairro"))
    cidade = title_case_if_lower(get_item(item, "cidade"))
    local = " - ".join(p for p in (bairro, cidade) if p) or "Localização não informada"
    disponibilidade = html.escape(clean_text(item.get("Disponibilidade comercial")) or "Não informado")
    endereco = html.escape(clean_text(get_item(item, "endereco")) or local)
    wa_link = html.escape(
        f"https://wa.me/{PDF_CONTACT_WHATSAPP_NUMBER}?text={quote('Tenho interesse na opção ' + f'{index:02d}' + '.')}",
        quote=True,
    )
    codigo = clean_text(item.get("Codigo") or get_item(item, "codigo")) or f"{index:02d}"
    visita_link = html.escape(
        f"https://wa.me/{PDF_CONTACT_WHATSAPP_NUMBER}?text={quote(' Gostaria de visitar o imóvel ' + codigo)}",
        quote=True,
    )
    maps_html = (
        f'<a class="btn alt" href="{html.escape(maps_link, quote=True)}" target="_blank" rel="noreferrer">Ver mapa</a>'
        if maps_link
        else ""
    )
    photo = clean_text(item.get("_foto"))
    visual_style = ""
    if photo:
        visual_style = (
            ' style="background-image: linear-gradient(180deg, rgba(5,74,122,.55) 0%, '
            f"rgba(5,74,122,.2) 45%, rgba(5,74,122,.7) 100%), url('{html.escape(photo, quote=True)}');\""
        )
    # A foto abre o mesmo anúncio do botão "Abrir anúncio" - mesmo link, mesmo destino.
    visual_tag, visual_attrs = (
        ("a", f' href="{link}" target="_blank" rel="noreferrer" aria-label="Abrir anúncio - {html.escape(local)}"')
        if link
        else ("div", "")
    )
    return f"""
        <article class="property reveal">
          <{visual_tag} class="property-visual"{visual_attrs}{visual_style}><span class="option">Opção {index:02d}</span><div class="location"><strong>{html.escape(local)}</strong><span>Cód. {html.escape(clean_text(item.get('Codigo')) or 'N/D')}</span></div></{visual_tag}>
          <div class="property-body">
            <div class="property-title"><div><h3>{headline}</h3><p>{endereco}</p></div><span class="code">{html.escape(format_dual_price(item))}</span></div>
            <div class="price"><small>Disponibilidade</small>{disponibilidade}</div>
            <div class="metrics">
              <div class="metric"><span>Área útil</span><strong>{html.escape(format_area(get_item(item, 'area_util')))}</strong></div>
              <div class="metric"><span>Área total</span><strong>{html.escape(format_area(get_item(item, 'area_total')))}</strong></div>
              <div class="metric"><span>Docas</span><strong>{html.escape(clean_text(get_item(item, 'docas')) or '0')}</strong></div>
            </div>
            <div class="analysis">
              <div><h4>Vocação recomendada</h4><p>{html.escape(build_strategic_note(item, items))}</p></div>
              <div><h4>Ponto de atenção</h4><p>{html.escape(build_attention_note(item))}</p></div>
            </div>
            <div class="actions"><a class="btn" href="{visita_link}" target="_blank" rel="noreferrer">Agendar visita</a><a class="btn alt" href="{link}" target="_blank" rel="noreferrer">Abrir anúncio</a>{maps_html}<a class="btn alt" href="{wa_link}" target="_blank" rel="noreferrer">Falar no WhatsApp</a></div>
          </div>
        </article>
    """


def report_inventory_html(items):
    articles = "".join(report_property_html(index, item, items) for index, item in enumerate(items, start=1))
    return f"""
    <section class="inventory" id="imoveis">
      <div class="wrap">
        <p class="eyebrow">Imóveis selecionados</p>
        <h2>Leitura individual.</h2>
        <p class="lead">Cada imóvel foi analisado sob critérios operacionais e comerciais. Dados não informados dependem de vistoria e validação documental.</p>
      </div>
      <div class="wrap" style="margin-top:42px">{articles}</div>
    </section>
    """


def report_priorities_html(items):
    recs = build_recommendations(items)
    if not recs:
        ranking = "<p style='padding:25px 24px;color:var(--muted)'>Sem dados suficientes para recomendações automáticas.</p>"
    else:
        ranking = "".join(
            f"""
            <div class="rank"><strong>{html.escape(rec['label'])} · {html.escape(get_property_headline(rec['item']))}</strong><p>{html.escape(rec['reason'])}</p></div>
            """
            for rec in recs
        )
    return f"""
    <section class="recommend" id="prioridades">
      <div class="wrap reveal">
        <p class="eyebrow">Recomendação de prioridade</p>
        <h2>Leitura estratégica da seleção.</h2>
        <p class="lead">As prioridades são indicativas e mudam conforme acesso, layout e condição efetiva de cada imóvel.</p>
        <div class="ranking">{ranking}</div>
        <p class="disclaimer">Recomendação comercial preliminar, sem garantia de disponibilidade, liquidez, viabilidade técnica, legal, ambiental ou urbanística. Recomenda-se visita técnica e validação profissional antes de qualquer decisão.</p>
      </div>
    </section>
    """


def report_cta_html(client):
    wa_link = html.escape(f"https://wa.me/{PDF_CONTACT_WHATSAPP_NUMBER}?text={quote(PDF_CONTACT_CTA)}", quote=True)
    heading = f"Vamos transformar a lista em uma agenda objetiva de visitas, {html.escape(client)}?" if client else "Vamos transformar a lista em uma agenda objetiva de visitas."
    logo = report_logo_data_uri()
    logo_html = f'<img class="cta-logo" src="{logo}" alt="SP Imóveis">' if logo else ""
    return f"""
    <section class="cta">
      <div class="wrap cta-inner">
        <div>
          {logo_html}
          <p class="eyebrow" style="color:var(--sky)">Próximo passo</p>
          <h2>{heading}</h2>
          <p>{html.escape(PDF_CONTACT_NAME)} · {html.escape(PDF_CONTACT_BRAND)} · {html.escape(PDF_CONTACT_WHATSAPP_DISPLAY)}</p>
        </div>
        <a class="btn light" href="{wa_link}" target="_blank" rel="noreferrer">Falar com {html.escape(PDF_CONTACT_NAME.split(' ')[0])}</a>
      </div>
    </section>
    """


MESES_PT = {
    1: "janeiro", 2: "fevereiro", 3: "março", 4: "abril", 5: "maio", 6: "junho",
    7: "julho", 8: "agosto", 9: "setembro", 10: "outubro", 11: "novembro", 12: "dezembro",
}


def report_footer_html(items, skipped_count=0):
    skipped_note = f" · {skipped_count} imóvel(is) excedente(s) não incluído(s)" if skipped_count > 0 else ""
    now = datetime.now()
    generated_at = f"{MESES_PT[now.month]} de {now.year}"
    return f"""
  <footer class="page-footer"><div class="wrap footer-inner"><span>SP Imóveis | Galpões Industriais</span><span>{len(items)} imóvel(is) selecionado(s){skipped_note} · {html.escape(generated_at)}</span></div></footer>
    """


REPORT_SCRIPT = """
  <script>
    var nav = document.getElementById('nav');
    if (nav) {
      var setNav = function () { nav.classList.toggle('scrolled', window.scrollY > 24); };
      setNav();
      window.addEventListener('scroll', setNav, { passive: true });
    }
    if (window.IntersectionObserver) {
      var observer = new IntersectionObserver(function (entries) {
        entries.forEach(function (entry) { if (entry.isIntersecting) entry.target.classList.add('visible'); });
      }, { threshold: 0.08 });
      document.querySelectorAll('.reveal').forEach(function (item) { observer.observe(item); });
    } else {
      document.querySelectorAll('.reveal').forEach(function (item) { item.classList.add('visible'); });
    }
  </script>
"""


def build_report_html(items, client_name, report_title, filters, skipped_count=0):
    title = clean_text(report_title) or "Seleção de imóveis"
    client = clean_text(client_name)
    doc_title = html.escape(title) + (f" - {html.escape(client)}" if client else "")
    filters = filters or {}
    description = html.escape(
        f"{len(items)} imóvel(is) selecionado(s) para {client}."
        if client
        else f"Seleção de {len(items)} imóvel(is) industriais preparada por {PDF_CONTACT_BRAND}."
    )
    body = "".join(
        [
            report_nav_html(),
            report_hero_html(items, title, client, filters),
            report_summary_html(items),
            report_kpis_html(items),
            report_compare_html(items),
            report_inventory_html(items),
            report_priorities_html(items),
            report_cta_html(client),
            report_footer_html(items, skipped_count),
        ]
    )
    return f"""<!doctype html>
<html lang="pt-BR">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{doc_title}</title>
<meta name="description" content="{description}">
<meta property="og:type" content="website">
<meta property="og:locale" content="pt_BR">
<meta property="og:site_name" content="{html.escape(PDF_CONTACT_BRAND)}">
<meta property="og:title" content="{doc_title}">
<meta property="og:description" content="{description}">
<meta property="og:image" content="{PRESENTATION_PREVIEW_IMAGE_FILENAME}">
<meta property="og:image:width" content="{PRESENTATION_PREVIEW_IMAGE_SIZE[0]}">
<meta property="og:image:height" content="{PRESENTATION_PREVIEW_IMAGE_SIZE[1]}">
<meta name="twitter:card" content="summary_large_image">
<meta name="twitter:title" content="{doc_title}">
<meta name="twitter:description" content="{description}">
<meta name="twitter:image" content="{PRESENTATION_PREVIEW_IMAGE_FILENAME}">
{report_style_block()}
</head>
<body>
{body}
{REPORT_SCRIPT}
</body>
</html>"""


def build_pdf_html(items, client_name, report_title, filters):
    return build_report_html(items, client_name, report_title, filters)


PDF_FOOTER_TEMPLATE = f"""
<div style="font-size:8px; width:100%; padding:0 16mm; color:#8a94a6; font-family:Arial,Helvetica,sans-serif; display:flex; justify-content:space-between; align-items:center;">
  <span>{html.escape(PDF_CONTACT_BRAND)}</span>
  <span>Página <span class="pageNumber"></span> de <span class="totalPages"></span></span>
</div>
"""


def render_pdf_from_html(html_str, output_path):
    driver = setup_scrape_driver()
    tmp_path = None
    try:
        fd, tmp_path = tempfile.mkstemp(suffix=".html")
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            f.write(html_str)
        driver.get(Path(tmp_path).resolve().as_uri())
        try:
            WebDriverWait(driver, 15).until(lambda d: d.execute_script("return document.readyState") == "complete")
        except Exception:
            pass
        time.sleep(1.5)
        result = driver.execute_cdp_cmd(
            "Page.printToPDF",
            {
                "printBackground": True,
                "paperWidth": 8.27,
                "paperHeight": 11.69,
                "marginTop": 0.2,
                "marginBottom": 0.45,
                "marginLeft": 0,
                "marginRight": 0,
                "displayHeaderFooter": True,
                "headerTemplate": "<span></span>",
                "footerTemplate": PDF_FOOTER_TEMPLATE,
            },
        )
        pdf_bytes = base64.b64decode(result["data"])
        with open(output_path, "wb") as f:
            f.write(pdf_bytes)
    finally:
        try:
            driver.quit()
        except Exception:
            pass
        if tmp_path:
            try:
                os.unlink(tmp_path)
            except Exception:
                pass


PRESENTATION_MAX_ITEMS = 40

LARGEST_IMAGE_JS = """
const imgs = Array.from(document.querySelectorAll('img'));
let best = null;
let bestArea = 0;
for (const img of imgs) {
  const w = img.naturalWidth || img.width || 0;
  const h = img.naturalHeight || img.height || 0;
  const area = w * h;
  if (w >= 200 && h >= 150 && area > bestArea) {
    bestArea = area;
    best = img.currentSrc || img.src;
  }
}
return best;
"""


def setup_scrape_driver():
    if not SELENIUM_AVAILABLE:
        detail = f" Detalhe: {SELENIUM_IMPORT_ERROR}" if SELENIUM_IMPORT_ERROR else ""
        raise RuntimeError(
            "Selenium nao esta instalado ou esta incompleto neste Python."
            f"{detail} Feche o app e inicie novamente pelo INICIAR_APP.bat."
        )
    options = Options()
    options.add_argument("--headless=new")
    options.add_argument("--disable-gpu")
    options.add_argument("--window-size=1366,900")
    options.add_argument("--log-level=3")
    try:
        driver = webdriver.Chrome(options=options)
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "Selenium esta incompleto neste Python. Feche o app e rode INICIAR_APP.bat "
            "para reinstalar as dependencias."
        ) from exc
    driver.set_page_load_timeout(20)
    return driver


UNAVAILABLE_LISTING_PATTERNS = [
    "nao esta mais disponivel",
    "imovel nao esta mais disponivel",
    "imovel indisponivel",
    "anuncio indisponivel",
    "anuncio nao encontrado",
    "anuncio expirado",
    "anuncio foi removido",
    "removido pelo anunciante",
    "pagina nao encontrada",
    "conteudo nao encontrado",
    "este imovel ja foi vendido",
    "este imovel ja foi alugado",
    "erro 404",
    "error 404",
]


def _normalize_for_match(text):
    normalized = unicodedata.normalize("NFKD", text or "")
    return normalized.encode("ascii", "ignore").decode("ascii").lower()


def listing_is_unavailable(driver):
    """Checks the currently loaded ad page for phrases indicating the listing was taken down,
    expired or errored out (e.g. "Infelizmente o imóvel não está mais disponível")."""
    try:
        body_text = driver.execute_script("return document.body ? document.body.innerText : '';") or ""
    except Exception:
        return False
    normalized = _normalize_for_match(body_text)
    return any(pattern in normalized for pattern in UNAVAILABLE_LISTING_PATTERNS)


def scrape_listing_image(driver, url):
    """Opens the ad page with Selenium and returns the main photo URL, or None if the link is
    invalid, the listing is unavailable/errored out, or no photo could be found."""
    if not url:
        return None
    try:
        driver.get(url)
    except Exception:
        return None
    try:
        WebDriverWait(driver, 12).until(lambda d: d.execute_script("return document.readyState") == "complete")
    except Exception:
        pass
    time.sleep(1.2)
    if listing_is_unavailable(driver):
        return None
    try:
        driver.execute_script("window.scrollBy(0, 500);")
        time.sleep(0.6)
    except Exception:
        pass
    if listing_is_unavailable(driver):
        return None
    try:
        og_image = driver.execute_script(
            "const m = document.querySelector('meta[property=\"og:image\"]'); return m ? m.content : null;"
        )
        if og_image:
            return og_image.strip()
    except Exception:
        pass
    try:
        largest = driver.execute_script(LARGEST_IMAGE_JS)
        if largest:
            return largest.strip()
    except Exception:
        pass
    return None


def collect_presentation_items(results, max_items=PRESENTATION_MAX_ITEMS):
    """Visits each result's ad link with Selenium and keeps only the ones that are still available
    and have a working photo. Listings that report themselves as unavailable/removed/errored are
    discarded, same as ones where no link or photo could be found.

    Returns (items_with_photos, skipped_count). Raises RuntimeError if Selenium/Chrome cannot start.
    """
    considered = results[:max_items]
    items_with_photos = []
    driver = None
    try:
        driver = setup_scrape_driver()
        for item in considered:
            link = clean_text(item.get("Link do imovel"))
            image_url = scrape_listing_image(driver, link) if link else None
            if not link or not image_url:
                continue
            items_with_photos.append({**item, "_foto": image_url})
    finally:
        if driver is not None:
            try:
                driver.quit()
            except Exception:
                pass
    return items_with_photos, len(considered) - len(items_with_photos)


def build_presentation_html(items, client_name, report_title, filters, skipped_count=0):
    return build_report_html(items, client_name, report_title, filters, skipped_count=skipped_count)


@app.get("/api/properties/options")
def api_properties_options():
    try:
        df, _ = ensure_property_data_loaded()
    except RuntimeError:
        df = None
    return jsonify({"ok": True, "types": get_property_type_options(df)})


@app.get("/api/properties/external/options")
def api_properties_external_options():
    try:
        df, _ = ensure_external_property_data_loaded()
    except RuntimeError:
        df = None
    return jsonify({"ok": True, "types": external_property_type_options(df)})


@app.get("/api/properties/table")
def api_properties_table():
    try:
        df, source = ensure_property_data_loaded()
    except RuntimeError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    columns = [str(col) for col in df.columns if not str(col).startswith("__")]
    rows = [{col: cell_display(row[col]) for col in columns} for _, row in df.iterrows()]
    return jsonify({"ok": True, "source": source, "columns": columns, "rows": rows})


@app.post("/api/properties/reload")
def api_properties_reload():
    try:
        df, source = reload_property_data()
    except RuntimeError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    return jsonify({"ok": True, "rows": len(df), "source": source, "types": get_property_type_options(df)})


@app.get("/api/properties/external/table")
def api_properties_external_table():
    try:
        df, source = ensure_external_property_data_loaded()
    except RuntimeError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    columns = [str(col) for col in df.columns if not str(col).startswith("__")]
    rows = [{col: cell_display(row[col]) for col in columns} for _, row in df.iterrows()]
    return jsonify({"ok": True, "source": source, "columns": columns, "rows": rows})


@app.post("/api/properties/external/reload")
def api_properties_external_reload():
    try:
        df, source = reload_external_property_data()
    except RuntimeError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    return jsonify({"ok": True, "rows": len(df), "source": source})


@app.post("/api/properties/upload")
def api_properties_upload():
    global property_df, property_source
    file = request.files.get("file")
    if not file:
        return jsonify({"ok": False, "error": "Nenhum arquivo enviado"}), 400
    path = save_upload(file, "imoveis")
    try:
        df = load_property_data(path)
        with property_lock:
            property_df = df
            property_source = str(path)
        return jsonify({"ok": True, "rows": len(df), "types": get_property_type_options(df)})
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500


@app.post("/api/properties/search")
def api_properties_search():
    global last_property_results, last_property_filters
    try:
        df, _ = ensure_property_data_loaded()
    except RuntimeError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    payload = request.get_json(silent=True) or {}
    try:
        selected_codes = [
            token.strip()
            for token in re.split(r"[\s,;]+", clean_text(payload.get("selected_codes")))
            if token.strip()
        ]
        filters = {
            "area_min": parse_optional_number(payload.get("area_min"), "area minima"),
            "area_max": parse_optional_number(payload.get("area_max"), "area maxima"),
            "valor_min": parse_optional_number(payload.get("valor_min"), "valor minimo"),
            "valor_max": parse_optional_number(payload.get("valor_max"), "valor maximo"),
            "region_or_cep": clean_text(payload.get("region_or_cep")),
            "price_mode": clean_text(payload.get("price_mode")) or "Automatico",
            "tipo_imovel": clean_text(payload.get("tipo_imovel")) or "Todos",
            "only_active": bool(payload.get("only_active")),
            "max_results": parse_positive_int(payload.get("max_results"), "limite"),
            "selected_codes": selected_codes,
        }
        filtered = filter_properties(df, **filters)
        results = [row_to_output_dict(row) for _, row in filtered.iterrows()]
        marketing = build_marketing_text(
            results,
            client_name=payload.get("client_name"),
            report_title=payload.get("report_title") or "Seleção de imóveis",
        )
        with property_lock:
            last_property_results = results
            last_property_filters = filters
        cards = [
            {
                "headline": get_property_headline(item),
                "summary": build_property_summary(item),
                "link": clean_text(item.get("Link do imovel")),
                "codigo": clean_text(item.get("Codigo")),
                "area": format_area(item.get("Area para filtro")),
                "valor": format_currency(item.get("Valor para filtro")),
                "bairro": title_case_if_lower(get_item(item, "bairro")),
                "cidade": title_case_if_lower(get_item(item, "cidade")),
            }
            for item in results
        ]
        return jsonify(
            {
                "ok": True,
                "count": len(results),
                "results": results,
                "cards": cards,
                "marketing_text": marketing,
            }
        )
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400


@app.post("/api/properties/external/search")
def api_properties_external_search():
    global last_external_property_results, last_external_property_filters
    try:
        df, _ = ensure_external_property_data_loaded()
    except RuntimeError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    payload = request.get_json(silent=True) or {}
    try:
        selected_codes = [
            token.strip()
            for token in re.split(r"[\s,;]+", clean_text(payload.get("selected_codes")))
            if token.strip()
        ]
        filters = {
            "area_min": parse_optional_number(payload.get("area_min"), "area minima"),
            "area_max": parse_optional_number(payload.get("area_max"), "area maxima"),
            "valor_min": parse_optional_number(payload.get("valor_min"), "valor minimo"),
            "valor_max": parse_optional_number(payload.get("valor_max"), "valor maximo"),
            "region_or_cep": clean_text(payload.get("region_or_cep")),
            "price_mode": clean_text(payload.get("price_mode")) or "Automatico",
            "tipo_imovel": clean_text(payload.get("tipo_imovel")) or "Todos",
            "only_active": bool(payload.get("only_active")),
            "max_results": parse_positive_int(payload.get("max_results"), "limite"),
            "selected_codes": selected_codes,
        }
        filtered = filter_properties(df, **filters)
        # link_value preserva o link real do anuncio (a planilha externa
        # ja traz o link de verdade - nao ha "Codigo" de catalogo proprio
        # para reconstruir um link, como acontece na Seleção Interna).
        results = [
            row_to_output_dict(row, link_value=row.get("Link do imovel")) for _, row in filtered.iterrows()
        ]
        marketing = build_marketing_text(
            results,
            client_name=payload.get("client_name"),
            report_title=payload.get("report_title") or "Seleção de imóveis",
        )
        with external_property_lock:
            last_external_property_results = results
            last_external_property_filters = filters
        cards = [
            {
                "headline": get_property_headline(item),
                "summary": build_property_summary(item),
                "link": clean_text(item.get("Link do imovel")),
                "codigo": clean_text(item.get("Codigo")),
                "area": format_area(item.get("Area para filtro")),
                "valor": format_currency(item.get("Valor para filtro")),
                "bairro": title_case_if_lower(get_item(item, "bairro")),
                "cidade": title_case_if_lower(get_item(item, "cidade")),
            }
            for item in results
        ]
        return jsonify(
            {
                "ok": True,
                "count": len(results),
                "results": results,
                "cards": cards,
                "marketing_text": marketing,
            }
        )
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400


@app.get("/api/properties/export_json")
def api_properties_export_json():
    with property_lock:
        data = list(last_property_results)
    output = io.BytesIO(json.dumps(data, ensure_ascii=False, indent=2, default=str).encode("utf-8"))
    return send_file(output, as_attachment=True, download_name="imoveis_filtrados.json", mimetype="application/json")


@app.get("/api/properties/external/export_json")
def api_properties_external_export_json():
    with external_property_lock:
        data = list(last_external_property_results)
    output = io.BytesIO(json.dumps(data, ensure_ascii=False, indent=2, default=str).encode("utf-8"))
    return send_file(
        output, as_attachment=True, download_name="imoveis_externos_filtrados.json", mimetype="application/json"
    )


@app.post("/api/properties/export_pdf")
def api_properties_export_pdf():
    payload = request.get_json(silent=True) or {}
    with property_lock:
        results = list(last_property_results)
        filters = dict(last_property_filters)
    if not results:
        return jsonify({"ok": False, "error": "Faça uma busca antes de gerar o PDF"}), 400
    if not SELENIUM_AVAILABLE:
        detail = f" Detalhe: {SELENIUM_IMPORT_ERROR}" if SELENIUM_IMPORT_ERROR else ""
        return jsonify(
            {"ok": False, "error": f"Selenium não está disponível para montar o PDF.{detail}"}
        ), 500

    try:
        items, _skipped = collect_presentation_items(results)
    except RuntimeError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500

    if not items:
        return jsonify(
            {"ok": False, "error": "Nenhum anúncio com link válido foi encontrado para montar o PDF."}
        ), 400

    client_name = payload.get("client_name") or ""
    report_title = payload.get("report_title") or "Seleção de imóveis"
    html_doc = build_pdf_html(items, client_name, report_title, filters)

    dest_path = next_presentation_path(PRESENTATIONS_PDF_DIR, client_name, ".pdf")
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    tmp.close()
    try:
        render_pdf_from_html(html_doc, tmp.name)
        shutil.move(tmp.name, dest_path)
        return jsonify(
            {
                "ok": True,
                "filename": dest_path.name,
                "path": str(dest_path),
                "url": f"/apresentacoes/pdf/{quote(dest_path.name)}",
            }
        )
    except RuntimeError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500
    finally:
        try:
            if os.path.exists(tmp.name):
                os.unlink(tmp.name)
        except Exception:
            pass


@app.post("/api/properties/export_web")
def api_properties_export_web():
    payload = request.get_json(silent=True) or {}
    with property_lock:
        results = list(last_property_results)
        filters = dict(last_property_filters)
    if not results:
        return jsonify({"ok": False, "error": "Faça uma busca antes de gerar a apresentação"}), 400
    if not SELENIUM_AVAILABLE:
        detail = f" Detalhe: {SELENIUM_IMPORT_ERROR}" if SELENIUM_IMPORT_ERROR else ""
        return jsonify(
            {"ok": False, "error": f"Selenium não está disponível para buscar as fotos dos anúncios.{detail}"}
        ), 500

    try:
        items, skipped_count = collect_presentation_items(results)
    except RuntimeError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500

    if not items:
        return jsonify(
            {"ok": False, "error": "Nenhum anúncio com link válido foi encontrado para montar a apresentação."}
        ), 400

    client_name = payload.get("client_name") or ""
    html_doc = build_presentation_html(
        items,
        client_name,
        payload.get("report_title") or "Seleção de imóveis",
        filters,
        skipped_count=skipped_count,
    )
    dest_path = next_presentation_path(PRESENTATIONS_INDEX_DIR, client_name, ".html")
    dest_path.write_text(html_doc, encoding="utf-8")
    ensure_preview_image_next_to(dest_path)
    return jsonify(
        {
            "ok": True,
            "filename": dest_path.name,
            "path": str(dest_path),
            "url": f"/apresentacoes/index/{quote(dest_path.name)}",
        }
    )


@app.post("/api/properties/external/export_pdf")
def api_properties_external_export_pdf():
    payload = request.get_json(silent=True) or {}
    with external_property_lock:
        results = list(last_external_property_results)
        filters = dict(last_external_property_filters)
    if not results:
        return jsonify({"ok": False, "error": "Faça uma busca antes de gerar o PDF"}), 400
    if not SELENIUM_AVAILABLE:
        detail = f" Detalhe: {SELENIUM_IMPORT_ERROR}" if SELENIUM_IMPORT_ERROR else ""
        return jsonify(
            {"ok": False, "error": f"Selenium não está disponível para montar o PDF.{detail}"}
        ), 500

    try:
        items, _skipped = collect_presentation_items(results)
    except RuntimeError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500

    if not items:
        return jsonify(
            {"ok": False, "error": "Nenhum anúncio com link válido foi encontrado para montar o PDF."}
        ), 400

    client_name = payload.get("client_name") or ""
    report_title = payload.get("report_title") or "Seleção de imóveis"
    html_doc = build_pdf_html(items, client_name, report_title, filters)

    dest_path = next_presentation_path(PRESENTATIONS_PDF_DIR, client_name, ".pdf")
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    tmp.close()
    try:
        render_pdf_from_html(html_doc, tmp.name)
        shutil.move(tmp.name, dest_path)
        return jsonify(
            {
                "ok": True,
                "filename": dest_path.name,
                "path": str(dest_path),
                "url": f"/apresentacoes/pdf/{quote(dest_path.name)}",
            }
        )
    except RuntimeError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500
    finally:
        try:
            if os.path.exists(tmp.name):
                os.unlink(tmp.name)
        except Exception:
            pass


@app.post("/api/properties/external/export_web")
def api_properties_external_export_web():
    payload = request.get_json(silent=True) or {}
    with external_property_lock:
        results = list(last_external_property_results)
        filters = dict(last_external_property_filters)
    if not results:
        return jsonify({"ok": False, "error": "Faça uma busca antes de gerar a apresentação"}), 400
    if not SELENIUM_AVAILABLE:
        detail = f" Detalhe: {SELENIUM_IMPORT_ERROR}" if SELENIUM_IMPORT_ERROR else ""
        return jsonify(
            {"ok": False, "error": f"Selenium não está disponível para buscar as fotos dos anúncios.{detail}"}
        ), 500

    try:
        items, skipped_count = collect_presentation_items(results)
    except RuntimeError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500

    if not items:
        return jsonify(
            {"ok": False, "error": "Nenhum anúncio com link válido foi encontrado para montar a apresentação."}
        ), 400

    client_name = payload.get("client_name") or ""
    html_doc = build_presentation_html(
        items,
        client_name,
        payload.get("report_title") or "Seleção de imóveis",
        filters,
        skipped_count=skipped_count,
    )
    dest_path = next_presentation_path(PRESENTATIONS_INDEX_DIR, client_name, ".html")
    dest_path.write_text(html_doc, encoding="utf-8")
    ensure_preview_image_next_to(dest_path)
    return jsonify(
        {
            "ok": True,
            "filename": dest_path.name,
            "path": str(dest_path),
            "url": f"/apresentacoes/index/{quote(dest_path.name)}",
        }
    )


def list_presentation_files(folder, url_prefix):
    if not folder.exists():
        return []
    files = []
    for entry in folder.iterdir():
        if not entry.is_file() or entry.name.startswith("~$"):
            continue
        stat = entry.stat()
        files.append(
            {
                "name": entry.name,
                "url": f"{url_prefix}/{quote(entry.name)}",
                "modified_at": stat.st_mtime,
                "modified_label": datetime.fromtimestamp(stat.st_mtime).strftime("%d/%m/%Y %H:%M"),
                "size": stat.st_size,
            }
        )
    files.sort(key=lambda item: item["modified_at"], reverse=True)
    for item in files:
        del item["modified_at"]
    return files


@app.get("/api/presentations/list")
def api_presentations_list():
    return jsonify(
        {
            "ok": True,
            "index": list_presentation_files(PRESENTATIONS_INDEX_DIR, "/apresentacoes/index"),
            "pdf": list_presentation_files(PRESENTATIONS_PDF_DIR, "/apresentacoes/pdf"),
        }
    )


@app.get("/apresentacoes/index/<path:filename>")
def serve_presentation_index(filename):
    return send_from_directory(PRESENTATIONS_INDEX_DIR, filename, as_attachment=False, mimetype="text/html")


@app.get("/apresentacoes/pdf/<path:filename>")
def serve_presentation_pdf(filename):
    return send_from_directory(PRESENTATIONS_PDF_DIR, filename, as_attachment=False, mimetype="application/pdf")


# N8N / CRM integration endpoints
NUMBER_PATTERN = r"(\d{1,3}(?:\.\d{3})+(?:,\d+)?|\d+(?:,\d+)?)"


def demand_empty():
    return {
        "metragem_min": None,
        "metragem_max": None,
        "valor_min": None,
        "valor_max": None,
        "regiao": "",
    }


def demand_score(demand):
    return sum(1 for value in demand.values() if value not in (None, ""))


def extract_number(match):
    number = parse_brazilian_number(match.group(1))
    return None if pd.isna(number) else float(number)


def plausible_area_number(value):
    return value is not None and 1 <= float(value) <= 500000


def plausible_value_number(value):
    return value is not None and 100 <= float(value) <= 1_000_000_000


def has_area_unit_near(text):
    return bool(re.search(r"\b(m2|m\s*2|m²|metros?|area|metragem)\b", text, flags=re.IGNORECASE))


def has_value_context(text):
    return bool(
        re.search(
            r"(r\$|valor|orcamento|orçamento|aluguel|locacao|locação|preco|preço|mensal|ate|até)",
            text,
            flags=re.IGNORECASE,
        )
    )


def clean_region_candidate(line):
    text = strip_html_to_text(line)
    text = re.sub(r"\b(r\$|ate|até|acima de|mais de|a partir de|minimo|mínimo|maximo|máximo)\b", " ", text, flags=re.IGNORECASE)
    text = re.sub(NUMBER_PATTERN, " ", text)
    text = re.sub(r"[+:;|,/]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip(" -")
    return text


def extract_demand_from_text(value):
    text = strip_html_to_text(value)
    demand = demand_empty()
    if not text:
        return demand
    lines = [line.strip() for line in re.split(r"[\n\r]+", text) if line.strip()]
    for line in lines:
        area_max_match = re.search(
            rf"(?:ate|até|max(?:imo)?|máx(?:imo)?)\s*(?:de\s*)?{NUMBER_PATTERN}\s*(?:m2|m\s*2|m²|metros?)",
            line,
            flags=re.IGNORECASE,
        )
        if area_max_match:
            value = extract_number(area_max_match)
            if plausible_area_number(value):
                demand["metragem_max"] = value

        area_min_match = re.search(
            rf"(?:\+|acima\s+de|mais\s+de|a\s+partir\s+de|min(?:imo)?|mín(?:imo)?)\s*(?:de\s*)?{NUMBER_PATTERN}\s*(?:m2|m\s*2|m²|metros?)?",
            line,
            flags=re.IGNORECASE,
        )
        if area_min_match and (has_area_unit_near(line) or "+" in line or not has_value_context(line)):
            value = extract_number(area_min_match)
            digits = re.sub(r"\D", "", area_min_match.group(1))
            if plausible_area_number(value) and len(digits) <= 6:
                demand["metragem_min"] = value

        explicit_area_match = re.search(rf"{NUMBER_PATTERN}\s*(?:m2|m\s*2|m²|metros?)", line, flags=re.IGNORECASE)
        if explicit_area_match and demand["metragem_min"] is None and demand["metragem_max"] is None:
            value = extract_number(explicit_area_match)
            if plausible_area_number(value):
                demand["metragem_min"] = value

        value_max_match = re.search(
            rf"(?:ate|até|max(?:imo)?|máx(?:imo)?|teto|orcamento|orçamento|valor|aluguel|locacao|locação|preco|preço)\D{{0,20}}(?:r\$)?\s*{NUMBER_PATTERN}",
            line,
            flags=re.IGNORECASE,
        )
        if value_max_match and not has_area_unit_near(line):
            value = extract_number(value_max_match)
            if plausible_value_number(value):
                demand["valor_max"] = value

        value_min_match = re.search(
            rf"(?:a\s+partir\s+de|min(?:imo)?|mín(?:imo)?)\D{{0,20}}(?:r\$)?\s*{NUMBER_PATTERN}",
            line,
            flags=re.IGNORECASE,
        )
        if value_min_match and has_value_context(line) and not has_area_unit_near(line):
            value = extract_number(value_min_match)
            if plausible_value_number(value):
                demand["valor_min"] = value

        currency_match = re.search(rf"r\$\s*{NUMBER_PATTERN}", line, flags=re.IGNORECASE)
        if currency_match and demand["valor_max"] is None:
            value = extract_number(currency_match)
            if plausible_value_number(value):
                demand["valor_max"] = value

    region_candidates = []
    for line in lines:
        key = normalize_key(line)
        if not re.search(r"[A-Za-zÀ-ÿ]", line):
            continue
        if any(word in key for word in ("visita", "codigo", "cod", "telefone", "whatsapp", "orcamento", "valor", "metragem")):
            continue
        if has_area_unit_near(line) or has_value_context(line):
            candidate = clean_region_candidate(line)
        else:
            candidate = strip_html_to_text(line)
        candidate_key = normalize_key(candidate)
        if candidate and re.search(r"[A-Za-zÀ-ÿ]", candidate) and candidate_key not in {"nota", "mensagem"}:
            region_candidates.append(candidate)
    if region_candidates:
        demand["regiao"] = region_candidates[-1]
    return demand


def demand_hash_for(conversation_id, demand):
    normalized = {
        "lead_id": str(conversation_id),
        "demanda": {
            "metragem_min": demand.get("metragem_min"),
            "metragem_max": demand.get("metragem_max"),
            "valor_min": demand.get("valor_min"),
            "valor_max": demand.get("valor_max"),
            "regiao": normalize_key(demand.get("regiao")),
        },
    }
    return hashlib.sha256(json_dumps(normalized).encode("utf-8")).hexdigest()[:16]


def get_lead_messages(conn, conversation_id):
    rows = conn.execute(
        "SELECT * FROM messages WHERE conversation_id = ? ORDER BY datetime(created_at) ASC, id ASC",
        (conversation_id,),
    ).fetchall()
    return [
        {
            "id": row["id"],
            "conversation_id": row["conversation_id"],
            "direction": row["direction"],
            "body": row["body"],
            "created_at": row["created_at"],
        }
        for row in rows
    ]


def extract_lead_demand(contact, messages):
    candidates = []
    for message in reversed(messages):
        candidates.append(
            {
                "text": message.get("body", ""),
                "origin": f"message:{message.get('id')}",
                "created_at": message.get("created_at", ""),
            }
        )
    for field in ("last_message", "notes"):
        if clean_text(contact.get(field)):
            candidates.append({"text": contact.get(field), "origin": f"conversation.{field}", "created_at": contact.get("updated_at", "")})
    best = {"demanda": demand_empty(), "score": 0, "mensagem_original": "", "origem": "crm_notes"}
    for candidate in candidates:
        demand = extract_demand_from_text(candidate["text"])
        score = demand_score(demand)
        if score > best["score"]:
            best = {
                "demanda": demand,
                "score": score,
                "mensagem_original": strip_html_to_text(candidate["text"]),
                "origem": candidate["origin"],
            }
    best["lead_id"] = str(contact["id"])
    best["nome"] = contact.get("name", "")
    best["telefone"] = contact.get("phone", "")
    best["demand_hash"] = demand_hash_for(contact["id"], best["demanda"]) if best["score"] else ""
    return best


def lead_payload(contact, messages):
    extraction = extract_lead_demand(contact, messages)
    return {
        "lead_id": str(contact["id"]),
        "nome": contact.get("name", ""),
        "telefone": contact.get("phone", ""),
        "stage": contact.get("stage", ""),
        "kind": contact.get("kind", ""),
        "source": contact.get("source", ""),
        "owner": contact.get("owner", ""),
        "last_message": contact.get("last_message", ""),
        "notes": contact.get("notes", ""),
        "property_codes": contact.get("property_codes", ""),
        "created_at": contact.get("created_at", ""),
        "updated_at": contact.get("updated_at", ""),
        "messages": messages,
        "extraction": extraction,
    }


def latest_property_file():
    if DEFAULT_PROPERTY_FILE:
        configured = Path(os.path.expandvars(DEFAULT_PROPERTY_FILE)).expanduser()
        if configured.exists():
            return configured
    if PROPERTY_DIR and PROPERTY_DIR.exists():
        dir_candidates = [
            path
            for path in PROPERTY_DIR.glob("*")
            if path.is_file()
            and path.suffix.lower() in {".xlsx", ".xls", ".csv"}
            and not path.name.startswith("~$")
        ]
        if dir_candidates:
            return sorted(dir_candidates, key=lambda path: path.stat().st_mtime, reverse=True)[0]
    candidates = []
    for folder in (UPLOAD_DIR, BASE_DIR, BASE_DIR.parent):
        if not folder.exists():
            continue
        candidates.extend(path for path in folder.glob("imoveis_*") if path.suffix.lower() in {".xlsx", ".xls", ".csv"})
    candidates.extend(find_standard_imobizi_files())
    candidates = [path for path in candidates if path.exists()]
    if not candidates:
        return None
    return sorted(candidates, key=lambda path: path.stat().st_mtime, reverse=True)[0]


def ensure_property_data_loaded():
    global property_df, property_source
    with property_lock:
        if property_df is not None:
            return property_df.copy(), property_source
    source = latest_property_file()
    if source is None:
        raise RuntimeError("Nenhuma planilha de imoveis encontrada. Configure CRM_PROPERTY_DIR ou CRM_PROPERTY_FILE.")
    try:
        df = load_property_data(source)
    except PermissionError:
        # Arquivo aberto no Excel (ou outro programa) - o Windows nega a
        # leitura enquanto estiver aberto.
        raise RuntimeError(
            f"O arquivo '{source.name}' está aberto no Excel (ou outro programa) e não pôde ser lido. "
            "Feche o arquivo e tente de novo."
        )
    except Exception as exc:
        raise RuntimeError(f"Não foi possível ler '{source.name}': {exc}")
    with property_lock:
        property_df = df
        property_source = str(source)
    return df.copy(), property_source


def reload_property_data():
    global property_df, property_source
    with property_lock:
        property_df = None
        property_source = None
    return ensure_property_data_loaded()


def filters_from_demand(demand, payload=None):
    payload = payload or {}
    return {
        "area_min": demand.get("metragem_min"),
        "area_max": demand.get("metragem_max"),
        "valor_min": demand.get("valor_min"),
        "valor_max": demand.get("valor_max"),
        "region_or_cep": clean_text(demand.get("regiao")),
        "price_mode": clean_text(payload.get("price_mode")) or "Automatico",
        "tipo_imovel": clean_text(payload.get("tipo_imovel")) or "Todos",
        "only_active": bool(payload.get("only_active", True)),
        "max_results": parse_positive_int(payload.get("max_results"), "limite") if clean_text(payload.get("max_results")) else 5,
        "selected_codes": payload.get("selected_codes") or [],
    }


def option_from_property(item):
    bairro = title_case_if_lower(get_item(item, "bairro"))
    cidade = title_case_if_lower(get_item(item, "cidade"))
    regiao = " - ".join([part for part in (bairro, cidade) if part])
    return {
        "id": clean_text(item.get("Codigo") or get_item(item, "codigo")),
        "codigo": clean_text(item.get("Codigo") or get_item(item, "codigo")),
        "titulo": get_property_headline(item),
        "metragem": item.get("Area para filtro"),
        "valor": item.get("Valor para filtro"),
        "regiao": regiao,
        "link": clean_text(item.get("Link do imovel")),
        "resumo": build_property_summary(item),
    }


def run_selector_for_payload(payload):
    df, source = ensure_property_data_loaded()
    demand = payload.get("demanda") or {}
    filters = filters_from_demand(demand, payload)
    filtered = filter_properties(df, **filters)
    results = [row_to_output_dict(row) for _, row in filtered.iterrows()]
    opcoes = [option_from_property(item) for item in results]
    return {
        "lead_id": str(payload.get("lead_id") or ""),
        "filtros_usados": {
            "metragem_min": filters["area_min"],
            "metragem_max": filters["area_max"],
            "valor_min": filters["valor_min"],
            "valor_max": filters["valor_max"],
            "regiao": filters["region_or_cep"],
            "price_mode": filters["price_mode"],
            "tipo_imovel": filters["tipo_imovel"],
            "only_active": filters["only_active"],
            "max_results": filters["max_results"],
        },
        "count": len(opcoes),
        "opcoes": opcoes,
        "property_source": source,
    }


def greeting_for_now():
    hour = datetime.now().hour
    if hour < 12:
        return "Bom dia"
    if hour < 18:
        return "Boa tarde"
    return "Boa noite"


def build_dispatch_message(lead, selector_result):
    name = clean_text(lead.get("nome") or lead.get("name")).split(" ")[0]
    options = selector_result.get("opcoes") or []
    lines = [f"{greeting_for_now()}, tudo bem{(' ' + name) if name else ''}?", ""]
    if not options:
        lines.append("Ainda nao encontrei uma opcao certeira com esses filtros. Vou ajustar a busca e te retorno.")
        return sanitize_outbound_message("\n".join(lines), name=name)
    lines.append("Encontrei algumas opcoes que parecem se encaixar no que voce procura:")
    lines.append("")
    for index, option in enumerate(options[:5], start=1):
        area = format_area(option.get("metragem"))
        value = format_currency(option.get("valor"))
        region = clean_text(option.get("regiao")) or "regiao a confirmar"
        title = clean_text(option.get("titulo")) or f"Opcao {index}"
        link = clean_text(option.get("link"))
        lines.append(f"{index}. {title}")
        lines.append(f"Area: {area} | Valor: {value} | Regiao: {region}")
        if link:
            lines.append(f"Link: {link}")
        lines.append("")
    lines.append("Me diz qual chamou mais sua atencao que eu te passo os proximos detalhes.")
    return sanitize_outbound_message("\n".join(lines), name=name)


def dispatch_payload_from_selector(payload):
    lead = payload.get("lead") or {}
    selector_result = payload.get("selector") or {}
    message = sanitize_outbound_message(payload.get("mensagem") or build_dispatch_message(lead, selector_result), name=lead.get("nome") or lead.get("name"))
    return {
        "lead_id": str(payload.get("lead_id") or lead.get("lead_id") or lead.get("id") or ""),
        "telefone": clean_phone(payload.get("telefone") or lead.get("telefone") or lead.get("phone")),
        "nome": clean_text(payload.get("nome") or lead.get("nome") or lead.get("name")),
        "mensagem": message,
        "opcoes": selector_result.get("opcoes") or payload.get("opcoes") or [],
        "demanda": payload.get("demanda") or {},
        "demand_hash": clean_text(payload.get("demand_hash") or payload.get("hash")),
    }


def prepare_blaster_dispatch(dispatch_payload):
    global uploaded_contacts_path
    phone = clean_phone(dispatch_payload.get("telefone"))
    message = sanitize_outbound_message(dispatch_payload.get("mensagem"), name=dispatch_payload.get("nome"))
    if not phone:
        raise ValueError("Lead sem telefone para disparo.")
    if not message:
        raise ValueError("Mensagem de disparo vazia.")
    df = pd.DataFrame(
        [
            {
                "Numero WPP": phone,
                "Nome": clean_text(dispatch_payload.get("nome")),
                "Mensagem": message,
                "Lead ID": clean_text(dispatch_payload.get("lead_id")),
                "Demanda Hash": clean_text(dispatch_payload.get("demand_hash")),
            }
        ]
    )
    path = UPLOAD_DIR / f"n8n_lead_{clean_text(dispatch_payload.get('lead_id')) or 'sem_id'}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Disparo N8N")
    uploaded_contacts_path = path
    reset_blaster_counts()
    return path, df


def start_blaster_thread(qr_wait=60, delay_min=10, delay_max=20):
    global bot_thread
    if not uploaded_contacts_path:
        raise RuntimeError("Nenhuma planilha carregada no disparador.")
    with blaster_lock:
        if blaster_state["status"] in {"starting", "waiting_qr", "running", "paused"}:
            raise RuntimeError("Bot ja esta em execucao.")
    delay_min = int(delay_min or 10)
    delay_max = int(delay_max or delay_min)
    if delay_max < delay_min:
        delay_max = delay_min
    df = read_tabular_file(uploaded_contacts_path).fillna("")
    stop_event.clear()
    pause_event.set()
    with blaster_lock:
        blaster_state.update(
            {
                "status": "starting",
                "total": len(df),
                "sent": 0,
                "failed": 0,
                "current": "",
                "countdown": 0,
                "delay_min": delay_min,
                "delay_max": delay_max,
                "logs": [],
            }
        )
    bot_thread = threading.Thread(target=bot_runner, args=(df, int(qr_wait or 60), delay_min, delay_max), daemon=True)
    bot_thread.start()


@app.get("/api/integration/leads/pending")
def api_integration_pending_leads():
    if not integration_is_authorized():
        return integration_auth_error()
    limit = min(parse_positive_int(request.args.get("limit") or "50", "limit") or 50, 200)
    include_processed = str(request.args.get("include_processed", "")).lower() in {"1", "true", "sim", "yes"}
    leads = []
    with get_db() as conn:
        rows = conn.execute(
            "SELECT * FROM conversations ORDER BY datetime(updated_at) DESC, id DESC LIMIT ?",
            (limit,),
        ).fetchall()
        for row in rows:
            contact = conversation_row_to_dict(row)
            messages = get_lead_messages(conn, contact["id"])
            payload = lead_payload(contact, messages)
            extraction = payload["extraction"]
            if extraction["score"] == 0:
                continue
            run = conn.execute(
                "SELECT status FROM automation_runs WHERE conversation_id = ? AND demand_hash = ?",
                (contact["id"], extraction["demand_hash"]),
            ).fetchone()
            processed = bool(run and automation_status_is_duplicate(run["status"]))
            if processed and not include_processed:
                continue
            payload["already_processed"] = processed
            payload["automation_status"] = run["status"] if run else ""
            leads.append(payload)
    return jsonify({"ok": True, "count": len(leads), "leads": leads})


@app.get("/api/integration/leads/<int:conversation_id>")
def api_integration_lead_detail(conversation_id):
    if not integration_is_authorized():
        return integration_auth_error()
    with get_db() as conn:
        row = conn.execute("SELECT * FROM conversations WHERE id = ?", (conversation_id,)).fetchone()
        if row is None:
            return jsonify({"ok": False, "error": "Lead nao encontrado"}), 404
        contact = conversation_row_to_dict(row)
        messages = get_lead_messages(conn, conversation_id)
        payload = lead_payload(contact, messages)
        extraction = payload["extraction"]
        run = None
        if extraction.get("demand_hash"):
            run = conn.execute(
                "SELECT * FROM automation_runs WHERE conversation_id = ? AND demand_hash = ?",
                (conversation_id, extraction["demand_hash"]),
            ).fetchone()
    payload["already_processed"] = bool(run and automation_status_is_duplicate(run["status"]))
    payload["automation_status"] = run["status"] if run else ""
    return jsonify({"ok": True, **payload})


@app.post("/api/integration/leads/<int:conversation_id>/extract")
def api_integration_extract_lead(conversation_id):
    if not integration_is_authorized():
        return integration_auth_error()
    with get_db() as conn:
        row = conn.execute("SELECT * FROM conversations WHERE id = ?", (conversation_id,)).fetchone()
        if row is None:
            return jsonify({"ok": False, "error": "Lead nao encontrado"}), 404
        contact = conversation_row_to_dict(row)
        messages = get_lead_messages(conn, conversation_id)
        extraction = extract_lead_demand(contact, messages)
        if extraction["score"]:
            upsert_automation_run(conn, conversation_id, extraction["demand_hash"], "criterios_extraidos", demand=extraction)
            add_automation_log(conversation_id, "criterios_extraidos", "Demanda extraida das notas do CRM.", extraction, extraction["demand_hash"], conn)
    return jsonify({"ok": True, "lead": contact, "extraction": extraction})


@app.post("/api/integration/selector/search")
def api_integration_selector_search():
    if not integration_is_authorized():
        return integration_auth_error()
    payload = request.get_json(silent=True) or {}
    try:
        selector_result = run_selector_for_payload(payload)
        lead_id = clean_text(payload.get("lead_id"))
        demand_hash = clean_text(payload.get("demand_hash"))
        if lead_id and demand_hash:
            with get_db() as conn:
                upsert_automation_run(
                    conn,
                    int(lead_id),
                    demand_hash,
                    "opcoes_encontradas" if selector_result["count"] else "sem_opcoes",
                    demand={"demanda": payload.get("demanda") or {}},
                    selector=selector_result,
                )
                add_automation_log(
                    int(lead_id),
                    "busca_executada",
                    f"{selector_result['count']} opcoes encontradas.",
                    selector_result,
                    demand_hash,
                    conn,
                )
        return jsonify({"ok": True, **selector_result})
    except Exception as exc:
        lead_id = clean_text(payload.get("lead_id"))
        demand_hash = clean_text(payload.get("demand_hash"))
        if lead_id and demand_hash:
            with get_db() as conn:
                upsert_automation_run(conn, int(lead_id), demand_hash, "erro_busca", error=str(exc))
                add_automation_log(int(lead_id), "erro_busca", str(exc), payload, demand_hash, conn)
        return jsonify({"ok": False, "error": str(exc)}), 400


@app.post("/api/integration/disparo/montar")
def api_integration_compose_dispatch():
    if not integration_is_authorized():
        return integration_auth_error()
    payload = request.get_json(silent=True) or {}
    dispatch = dispatch_payload_from_selector(payload)
    return jsonify({"ok": True, "disparo": dispatch})


@app.post("/api/integration/disparo/preparar")
def api_integration_prepare_dispatch():
    if not integration_is_authorized():
        return integration_auth_error()
    payload = request.get_json(silent=True) or {}
    dispatch = payload.get("disparo") or payload
    lead_id = clean_text(dispatch.get("lead_id"))
    demand_hash = clean_text(dispatch.get("demand_hash"))
    force = bool(payload.get("force") or dispatch.get("force"))
    try:
        if lead_id and demand_hash and not force:
            with get_db() as conn:
                existing = conn.execute(
                    "SELECT status FROM automation_runs WHERE conversation_id = ? AND demand_hash = ?",
                    (int(lead_id), demand_hash),
                ).fetchone()
                if existing and automation_status_is_duplicate(existing["status"]):
                    add_automation_log(int(lead_id), "duplicado_bloqueado", "Demanda ja preparada/enviada.", dispatch, demand_hash, conn)
                    return jsonify({"ok": True, "duplicate": True, "status": existing["status"], "message": "Disparo duplicado bloqueado."})

        path, df = prepare_blaster_dispatch(dispatch)
        auto_start = bool(payload.get("auto_start") or dispatch.get("auto_start"))
        if auto_start:
            start_blaster_thread(
                qr_wait=payload.get("qr_wait") or dispatch.get("qr_wait") or 60,
                delay_min=payload.get("delay_min") or dispatch.get("delay_min") or 10,
                delay_max=payload.get("delay_max") or dispatch.get("delay_max") or 20,
            )
        if lead_id and demand_hash:
            status = "disparo_iniciado" if auto_start else "disparo_preparado"
            with get_db() as conn:
                upsert_automation_run(conn, int(lead_id), demand_hash, status, dispatch=dispatch, message=dispatch.get("mensagem"))
                note = (
                    f"Automacao N8N: {'disparo iniciado' if auto_start else 'disparo preparado'} em {datetime.now().strftime('%d/%m/%Y %H:%M')}.\n"
                    f"Codigos: {', '.join([clean_text(item.get('codigo') or item.get('id')) for item in dispatch.get('opcoes', []) if clean_text(item.get('codigo') or item.get('id'))]) or 'sem codigos'}"
                )
                insert_crm_message(conn, int(lead_id), "note", note)
                add_automation_log(int(lead_id), status, "Disparo carregado no modulo de envios.", {"path": str(path), "rows": len(df)}, demand_hash, conn)
        preview = df.head(8).fillna("").astype(str)
        return jsonify(
            {
                "ok": True,
                "duplicate": False,
                "status": "disparo_iniciado" if auto_start else "disparo_preparado",
                "rows": len(df),
                "file": str(path),
                "columns": [str(col) for col in preview.columns],
                "preview": preview.values.tolist(),
            }
        )
    except Exception as exc:
        if lead_id and demand_hash:
            with get_db() as conn:
                upsert_automation_run(conn, int(lead_id), demand_hash, "erro_disparo", error=str(exc))
                add_automation_log(int(lead_id), "erro_disparo", str(exc), dispatch, demand_hash, conn)
        return jsonify({"ok": False, "error": str(exc)}), 400


@app.post("/api/integration/leads/<int:conversation_id>/automation-result")
def api_integration_register_result(conversation_id):
    if not integration_is_authorized():
        return integration_auth_error()
    payload = request.get_json(silent=True) or {}
    demand_hash = clean_text(payload.get("demand_hash"))
    status = clean_text(payload.get("status")) or ("disparo_feito" if payload.get("sent") else "erro")
    message = sanitize_outbound_message(payload.get("mensagem") or payload.get("message"), name=payload.get("nome"))
    error = clean_text(payload.get("error"))
    if not demand_hash:
        return jsonify({"ok": False, "error": "Informe demand_hash"}), 400
    with get_db() as conn:
        exists = conn.execute("SELECT id FROM conversations WHERE id = ?", (conversation_id,)).fetchone()
        if exists is None:
            return jsonify({"ok": False, "error": "Lead nao encontrado"}), 404
        sent = status in {"disparo_feito", "sent", "done"}
        upsert_automation_run(
            conn,
            conversation_id,
            demand_hash,
            status,
            dispatch=payload.get("disparo"),
            message=message,
            error=error,
            sent=sent,
        )
        if sent and message:
            insert_crm_message(conn, conversation_id, "out", message)
        else:
            insert_crm_message(conn, conversation_id, "note", f"Automacao N8N: {status}. {error}".strip())
        add_automation_log(conversation_id, status, error or "Resultado registrado pelo N8N.", payload, demand_hash, conn)
    return jsonify({"ok": True, "status": status})


@app.post("/api/integration/leads/<int:conversation_id>/automation-log")
def api_integration_register_log(conversation_id):
    if not integration_is_authorized():
        return integration_auth_error()
    payload = request.get_json(silent=True) or {}
    event = clean_text(payload.get("event")) or "n8n_log"
    message = clean_text(payload.get("message") or payload.get("error"))
    demand_hash = clean_text(payload.get("demand_hash"))
    with get_db() as conn:
        exists = conn.execute("SELECT id FROM conversations WHERE id = ?", (conversation_id,)).fetchone()
        if exists is None:
            return jsonify({"ok": False, "error": "Lead nao encontrado"}), 404
        add_automation_log(conversation_id, event, message, payload, demand_hash, conn)
        if payload.get("note"):
            insert_crm_message(conn, conversation_id, "note", payload.get("note"))
    return jsonify({"ok": True, "event": event})


@app.post("/api/integration/leads/<int:conversation_id>/process")
def api_integration_process_lead(conversation_id):
    if not integration_is_authorized():
        return integration_auth_error()
    payload = request.get_json(silent=True) or {}
    with get_db() as conn:
        row = conn.execute("SELECT * FROM conversations WHERE id = ?", (conversation_id,)).fetchone()
        if row is None:
            return jsonify({"ok": False, "error": "Lead nao encontrado"}), 404
        contact = conversation_row_to_dict(row)
        messages = get_lead_messages(conn, conversation_id)
        extraction = extract_lead_demand(contact, messages)
        if extraction["score"] == 0:
            add_automation_log(conversation_id, "sem_demanda", "Nenhum criterio encontrado nas notas.", extraction, "", conn)
            return jsonify({"ok": False, "error": "Nenhum criterio encontrado nas notas.", "extraction": extraction}), 400
        existing = conn.execute(
            "SELECT status FROM automation_runs WHERE conversation_id = ? AND demand_hash = ?",
            (conversation_id, extraction["demand_hash"]),
        ).fetchone()
        if existing and automation_status_is_duplicate(existing["status"]) and not payload.get("force"):
            add_automation_log(conversation_id, "duplicado_bloqueado", "Demanda ja processada.", extraction, extraction["demand_hash"], conn)
            return jsonify({"ok": True, "duplicate": True, "status": existing["status"], "extraction": extraction})
        upsert_automation_run(conn, conversation_id, extraction["demand_hash"], "criterios_extraidos", demand=extraction)
        add_automation_log(conversation_id, "criterios_extraidos", "Demanda extraida das notas do CRM.", extraction, extraction["demand_hash"], conn)
    try:
        selector_result = run_selector_for_payload(
            {
                **payload,
                "lead_id": str(conversation_id),
                "demanda": extraction["demanda"],
                "demand_hash": extraction["demand_hash"],
            }
        )
        dispatch = dispatch_payload_from_selector(
            {
                "lead_id": str(conversation_id),
                "lead": {"lead_id": str(conversation_id), "nome": contact["name"], "telefone": contact["phone"]},
                "demanda": extraction["demanda"],
                "demand_hash": extraction["demand_hash"],
                "selector": selector_result,
            }
        )
        with get_db() as conn:
            upsert_automation_run(conn, conversation_id, extraction["demand_hash"], "opcoes_encontradas", selector=selector_result)
            add_automation_log(conversation_id, "busca_executada", f"{selector_result['count']} opcoes encontradas.", selector_result, extraction["demand_hash"], conn)
        path, df = prepare_blaster_dispatch(dispatch)
        auto_start = bool(payload.get("auto_start"))
        if auto_start:
            start_blaster_thread(payload.get("qr_wait") or 60, payload.get("delay_min") or 10, payload.get("delay_max") or 20)
        with get_db() as conn:
            status = "disparo_iniciado" if auto_start else "disparo_preparado"
            upsert_automation_run(conn, conversation_id, extraction["demand_hash"], status, dispatch=dispatch, message=dispatch["mensagem"])
            add_automation_log(conversation_id, status, "Disparo carregado no modulo de envios.", {"path": str(path), "rows": len(df)}, extraction["demand_hash"], conn)
        return jsonify({"ok": True, "duplicate": False, "extraction": extraction, "selector": selector_result, "disparo": dispatch, "status": status})
    except Exception as exc:
        with get_db() as conn:
            upsert_automation_run(conn, conversation_id, extraction["demand_hash"], "erro", error=str(exc))
            add_automation_log(conversation_id, "erro", str(exc), payload, extraction["demand_hash"], conn)
        return jsonify({"ok": False, "error": str(exc), "extraction": extraction}), 400


@app.get("/api/health")
def api_health():
    return jsonify(
        {
            "ok": True,
            "build": APP_BUILD,
            "selenium_available": SELENIUM_AVAILABLE,
            "selenium_error": SELENIUM_IMPORT_ERROR if not SELENIUM_AVAILABLE else "",
            "integration_token_required": bool(INTEGRATION_TOKEN),
            "n8n_webhook_configured": bool(N8N_CRM_WEBHOOK_URL),
        }
    )


@app.get("/")
def index():
    return render_template("index.html")


if __name__ == "__main__":
    if getattr(sys, "stdout", None) is None:
        sys.stdout = open(DATA_DIR / "server.out.log", "a", encoding="utf-8")
    if getattr(sys, "stderr", None) is None:
        sys.stderr = open(DATA_DIR / "server.err.log", "a", encoding="utf-8")
    port = int(os.environ.get("PORT", "5000"))
    print("=" * 58)
    print("  SP Imoveis CRM / WhatsApp")
    print(f"  Acesse: http://127.0.0.1:{port}")
    print("=" * 58)
    app.run(debug=False, port=port, threaded=True)
