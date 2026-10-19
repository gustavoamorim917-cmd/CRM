"""
Atualizar_imoveis_do_imobizi.py - une "baixar_relatorio_imoveis.py" +
"baixar_relatorio_imoveis - Copia.py" + "baixar_e_unificar_relatorios.py"
num unico fluxo, acionavel pelo botao "Atualizar" na aba Seletor do app
(app/sp_integrado), e tambem executavel sozinho pelo terminal.

Por que dois relatorios: o relatorio "Imoveis - Cadastros" do Imobzi, pedido
com todos os usuarios desde 01/01/2012 (praticamente "desde sempre"), as
vezes nao traz tudo que um pedido mais recente (desde 17/12/2024) traz - por
isso os scripts originais baixavam os DOIS e uniam o resultado, em vez de
confiar so no de periodo mais longo. Este script preserva esse
comportamento, so que numa UNICA sessao do Chrome (loga uma vez só).

O que este script faz, em ordem:
    1. Abre o Chrome, faz login uma vez no Imobzi.
    2. Baixa o relatorio "Imoveis - Cadastros" desde 01/01/2012 (todos os
       usuarios).
    3. Baixa o MESMO relatorio de novo, agora desde 17/12/2024.
    4. Une as duas planilhas: mantem 100% da primeira e cola, logo em
       seguida, o conteudo da segunda a partir da 3a linha (as duas
       primeiras linhas dela - cabecalho/resumo - sao descartadas).
    5. Converte a planilha unificada em .xlsx (UTF-8, acentos corretos).
    6. Salva SO esse .xlsx final (nome com a data do dia, ex.
       "imoveis_270826_unificado.xlsx") na pasta configurada em
       CRM_PROPERTY_DIR - a mesma pasta de onde o app sempre carrega a
       planilha mais recente. Se ja existir um arquivo do mesmo dia la, ele
       e sobrescrito.
    7. Apaga os arquivos intermediarios (os 2 relatorios brutos baixados,
       os xlsx "irmaos" que a conversao gera e o CSV unificado temporario)
       - so sobra o .xlsx final na pasta de relatorios.

Uso direto pelo terminal (sem passar pelo app):
    python "Atualizar_imoveis_do_imobizi.py"

Uso pelo app (app/sp_integrado/app.py):
    O app carrega este arquivo dinamicamente (importlib) quando o botao
    "Atualizar" da aba Seletor e clicado, chama set_progress_callback() para
    receber os logs em tempo real e depois sincronizar_imoveis().

Configuracao:
    As credenciais do Imobzi (IMOBZI_EMAIL / IMOBZI_SENHA) e a pasta de
    destino (CRM_PROPERTY_DIR) sao lidas do ambiente. Este script tenta
    carrega-las, nessa ordem: (1) variaveis ja definidas no processo (quando
    chamado pelo app, que ja leu o proprio .env), (2) o .env desta pasta,
    (3) o .env do app (app/sp_integrado/.env).
"""

import csv
import io
import os
import re
import shutil
import sys
import threading
import time
from datetime import datetime
from pathlib import Path

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import (
    TimeoutException,
    ElementClickInterceptedException,
    StaleElementReferenceException,
)


# --------------------------------------------------------------------------
# Configuracao
# --------------------------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent           # .../AUTOMAÇÔES/Imoveis
PROJECT_ROOT = BASE_DIR.parent.parent                 # raiz do projeto
APP_DIR = PROJECT_ROOT / "app" / "sp_integrado"       # .../app/sp_integrado

LOGIN_URL = "https://my.imobzi.com/#/login"
# Link do relatorio "Imoveis - Cadastros" da SUA conta Imobzi (copie da barra
# de enderecos ao abrir o relatorio). Lido de IMOBZI_REPORT_URL_IMOVEIS no
# .env - ver sincronizar_imoveis().
REPORT_URL = os.environ.get("IMOBZI_REPORT_URL_IMOVEIS", "").strip()

# Os dois periodos herdados dos scripts originais (ver docstring acima).
DATA_INICIAL_1 = "01/01/2012"
DATA_INICIAL_2 = "17/12/2024"

NOME_BASE_ARQUIVO = "imoveis"

# Pasta fixa de fallback, igual ao antigo baixar_e_unificar_relatorios.py -
# so usada se CRM_PROPERTY_DIR nao estiver configurado em nenhum .env.
PASTA_RELATORIOS_FALLBACK = BASE_DIR / "Relatórios de Imóveis"

CONVERTER_TIPOS_NO_XLSX = True   # numeros viram numero e datas viram data no Excel

TIMEOUT_PADRAO = 30            # segundos para esperas normais de UI
TIMEOUT_PROCESSAMENTO = 300    # segundos para o Imobzi gerar o relatorio (pode ser grande)
TIMEOUT_DOWNLOAD = 180         # segundos para o arquivo terminar de baixar
ESPERA_ANTES_DE_RECARREGAR = 60  # segundos de espera fixa antes do reload, depois da notificacao


class SincronizacaoCancelada(Exception):
    """Levantada quando o usuario pede para cancelar a sincronizacao em andamento."""


# --------------------------------------------------------------------------
# Log / progresso (o app injeta um callback para acompanhar em tempo real)
# --------------------------------------------------------------------------

_PROGRESS_CB = None     # assinatura: cb(mensagem: str, kind: str)
_STOP_EVENT = threading.Event()


def set_progress_callback(callback):
    global _PROGRESS_CB
    _PROGRESS_CB = callback


def log(msg, kind="info"):
    agora = datetime.now().strftime("%H:%M:%S")
    print(f"[{agora}] {msg}", flush=True)
    if _PROGRESS_CB:
        try:
            _PROGRESS_CB(msg, kind)
        except Exception:
            pass


def _checar_cancelamento():
    if _STOP_EVENT.is_set():
        raise SincronizacaoCancelada("Sincronizacao cancelada pelo usuario.")


def _sleep_interruptivel(segundos):
    restante = float(segundos)
    while restante > 0:
        _checar_cancelamento()
        pedaco = min(0.5, restante)
        time.sleep(pedaco)
        restante -= pedaco


# --------------------------------------------------------------------------
# Credenciais / pasta de destino
# --------------------------------------------------------------------------

def _carregar_env_de(caminho_env):
    if not caminho_env.exists():
        return
    for linha_bruta in caminho_env.read_text(encoding="utf-8").splitlines():
        linha = linha_bruta.strip()
        if not linha or linha.startswith("#") or "=" not in linha:
            continue
        chave, valor = linha.split("=", 1)
        chave = chave.strip()
        valor = valor.strip().strip('"').strip("'")
        if chave and chave not in os.environ:
            os.environ[chave] = valor


def carregar_env():
    """Carrega variaveis do .env desta pasta e depois do .env do app - sem
    nunca sobrescrever uma variavel que ja esteja definida no processo
    (assim, quando chamado pelo app, o .env do app sempre manda)."""
    _carregar_env_de(BASE_DIR / ".env")
    _carregar_env_de(APP_DIR / ".env")


def resolver_pasta_destino():
    env_dir = os.environ.get("CRM_PROPERTY_DIR", "").strip()
    if env_dir:
        return Path(os.path.expandvars(env_dir)).expanduser()
    return Path(PASTA_RELATORIOS_FALLBACK)


# --------------------------------------------------------------------------
# Selenium: driver, login, navegacao (herdado do baixar_relatorio_imoveis.py)
# --------------------------------------------------------------------------

def montar_driver(pasta_download):
    os.makedirs(pasta_download, exist_ok=True)

    chrome_options = Options()
    chrome_options.add_argument("--start-maximized")
    chrome_options.add_argument("--disable-notifications")

    prefs = {
        "download.default_directory": str(pasta_download),
        "download.prompt_for_download": False,
        "download.directory_upgrade": True,
        "safebrowsing.enabled": True,
    }
    chrome_options.add_experimental_option("prefs", prefs)

    # Selenium 4.18+ resolve o chromedriver sozinho (Selenium Manager),
    # mesma abordagem ja usada pelo app.py para o Disparador de WhatsApp.
    return webdriver.Chrome(options=chrome_options)


def fazer_login(driver, wait, email, senha):
    log("Abrindo pagina de login do Imobzi...")
    driver.get(LOGIN_URL)

    campo_email = wait.until(EC.presence_of_element_located((By.ID, "email")))
    campo_senha = driver.find_element(By.ID, "password")

    campo_email.clear()
    campo_email.send_keys(email)
    campo_senha.clear()
    campo_senha.send_keys(senha)

    log("Enviando login...")
    botao_entrar = wait.until(EC.element_to_be_clickable(
        (By.XPATH, "//button[@type='submit' and normalize-space()='ENTRAR']")
    ))
    clicar_com_retentativa(driver, botao_entrar)

    try:
        wait.until(lambda d: "/login" not in d.current_url)
    except TimeoutException:
        log("Login nao confirmou de primeira, tentando novamente...")
        clicar_com_retentativa(driver, botao_entrar)
        try:
            wait.until(lambda d: "/login" not in d.current_url)
        except TimeoutException:
            campo_senha = driver.find_element(By.ID, "password")
            campo_senha.send_keys(Keys.RETURN)
            wait.until(lambda d: "/login" not in d.current_url)

    log("Login efetuado.")


def abrir_relatorio(driver, wait):
    log("Abrindo relatorio de Imoveis...")
    # Como este script baixa o MESMO relatorio duas vezes (uma por
    # DATA_INICIAL), a 2a chamada navega para a URL EXATA em que o
    # navegador ja esta - e alguns navegadores/versoes do Chromedriver nao
    # fazem um reload de verdade nesse caso (URL "igual" e tratada como
    # no-op). Sem um reload de verdade, qualquer coisa que tenha ficado
    # aberta por cima da tela (ex.: o modal de notificacoes) continuaria
    # aberta. Por isso, se ja estivermos nessa URL, forca um refresh().
    if driver.current_url == REPORT_URL:
        driver.refresh()
    else:
        driver.get(REPORT_URL)
    wait.until(EC.presence_of_element_located(
        (By.XPATH, "//*[contains(text(),'Total de im')]")
    ))
    log("Relatorio carregado.")


def clicar_com_retentativa(driver, elemento, tentativas=3):
    for _ in range(tentativas):
        try:
            elemento.click()
            return
        except (ElementClickInterceptedException, StaleElementReferenceException):
            time.sleep(1)
    driver.execute_script("arguments[0].click();", elemento)


def _definir_data_inicial(driver, wait, data_inicial):
    dia, mes, ano = data_inicial.split("/")

    dia_xpath = (
        "(//input[@placeholder='Data inicial']"
        "/following::div[contains(@class,'calendar-container')])[1]"
        "//div[contains(@class,'cal-in-month')]"
        f"//span[contains(@class,'cal-day-number') and normalize-space()='{int(dia)}']"
    )

    log("  -> procurando o campo 'Data inicial'...")
    campo_data_inicial = wait.until(EC.element_to_be_clickable(
        (By.CSS_SELECTOR, "input[placeholder='Data inicial']")
    ))

    for tentativa in range(1, 4):
        log(f"  -> clicando no campo e digitando a data (tentativa {tentativa})...")
        clicar_com_retentativa(driver, campo_data_inicial)
        campo_data_inicial.send_keys(data_inicial)
        valor_atual = campo_data_inicial.get_attribute("value") or ""
        log("  -> apertando Tab para sair do campo e fechar o calendario...")
        campo_data_inicial.send_keys(Keys.TAB)
        time.sleep(0.5)
        if data_inicial in valor_atual:
            break
        log(f"  -> campo ficou com '{valor_atual}' em vez de '{data_inicial}', tentando de novo...")
    else:
        raise RuntimeError(
            f"Nao consegui digitar a data no campo 'Data inicial' "
            f"(ficou com '{valor_atual}' apos 3 tentativas)."
        )

    log("  -> procurando o dia no calendario (caso ainda esteja aberto)...")
    try:
        dia_1 = WebDriverWait(driver, 3).until(EC.element_to_be_clickable((By.XPATH, dia_xpath)))
        log("  -> dia encontrado, clicando...")
        clicar_com_retentativa(driver, dia_1)
        time.sleep(0.5)
    except TimeoutException:
        log("  -> calendario ja estava fechado, seguindo.")

    log("Data inicial definida, seguindo para o filtro de usuario...")


def _selecionar_usuario_todos(driver, wait):
    log("  -> apertando Tab para fechar o calendario...")
    driver.switch_to.active_element.send_keys(Keys.TAB)
    time.sleep(0.5)
    log("  -> calendario fechado.")

    valor_atual_usuario_xpath = "//h2[normalize-space()='Usuário']/following::div[contains(@class,'combobox')][1]"
    todos_xpath = "//*[normalize-space(text())='Todos']"

    combobox_usuario = wait.until(EC.element_to_be_clickable((By.XPATH, valor_atual_usuario_xpath)))
    log(f"  -> usuario atual no filtro: '{combobox_usuario.text.strip()}'. Clicando para abrir a lista...")

    opcao_todos = None
    for tentativa in range(1, 4):
        _fechar_painel_notificacoes(driver)  # defensivo - ver _fechar_painel_notificacoes
        clicar_com_retentativa(driver, combobox_usuario)
        try:
            opcao_todos = WebDriverWait(driver, 3).until(
                EC.element_to_be_clickable((By.XPATH, todos_xpath))
            )
            log(f"  -> lista de usuarios abriu na tentativa {tentativa}.")
            break
        except TimeoutException:
            log(f"  -> tentativa {tentativa}: lista de usuarios nao abriu, clicando de novo...")
            continue

    if opcao_todos is None:
        raise RuntimeError("Nao consegui abrir a lista de usuarios para selecionar 'Todos'.")

    for tentativa in range(1, 4):
        _fechar_painel_notificacoes(driver)  # defensivo - ver _fechar_painel_notificacoes
        log(f"  -> clicando em 'Todos' (tentativa {tentativa})...")
        clicar_com_retentativa(driver, opcao_todos)
        time.sleep(0.6)
        texto_combobox = combobox_usuario.text.strip()
        if texto_combobox == "Todos":
            log("  -> usuario confirmado como 'Todos'.")
            break
        log(f"  -> combobox de usuario ainda mostra '{texto_combobox}', tentando de novo...")
    else:
        raise RuntimeError("Nao consegui confirmar a selecao de 'Todos' no filtro de usuario.")


def aplicar_filtro_periodo_total(driver, wait, data_inicial):
    # Defensivo: se a notificacao de "relatorio pronto" da etapa anterior
    # ficou aberta (ou uma nova chegou nesse meio-tempo), ela intercepta os
    # cliques no filtro - fecha antes de comecar a mexer em qualquer coisa.
    _fechar_painel_notificacoes(driver)

    log("Abrindo painel de filtro...")
    botao_filtro = wait.until(EC.element_to_be_clickable(
        (By.XPATH, "//page-reports-view//button[.//i[normalize-space()='filter_list']]")
    ))
    clicar_com_retentativa(driver, botao_filtro)

    log(f"Definindo data inicial ({data_inicial})...")
    _definir_data_inicial(driver, wait, data_inicial)

    log("Definindo usuario como 'Todos'...")
    _selecionar_usuario_todos(driver, wait)

    log("Confirmando filtro (icone check)...")

    def _painel_esta_aberto(d):
        try:
            classes = d.find_element(By.CSS_SELECTOR, ".filter-wrapper").get_attribute("class") or ""
            return "show" in classes.split()
        except Exception:
            return False

    aplicado = False
    for tentativa in range(1, 4):
        botao_aplicar = wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, "#button-save")))
        clicar_com_retentativa(driver, botao_aplicar)
        try:
            WebDriverWait(driver, 6).until(lambda d: not _painel_esta_aberto(d))
            aplicado = True
            break
        except TimeoutException:
            log(f"Tentativa {tentativa}: painel de filtro nao fechou, clicando de novo...")
            continue

    if not aplicado:
        raise RuntimeError("Nao foi possivel confirmar o filtro (icone check).")

    _esperar_total_imoveis_estabilizar(driver)
    log(f"Filtro aplicado (desde {data_inicial}, todos os usuarios). Total de imoveis: {_texto_total_imoveis(driver)}")


def _texto_total_imoveis(driver):
    try:
        elemento = driver.find_element(By.XPATH, "//*[contains(text(),'Total de im')]/..")
        return elemento.text.strip()
    except Exception:
        return None


def _esperar_total_imoveis_estabilizar(driver, timeout=60, intervalo=1.5, estavel_por=3):
    inicio = time.time()
    leituras_iguais = 0
    ultimo_valor = None

    while time.time() - inicio < timeout:
        _checar_cancelamento()
        valor_atual = _texto_total_imoveis(driver)
        if valor_atual and valor_atual == ultimo_valor:
            leituras_iguais += 1
            if leituras_iguais >= estavel_por:
                return
        else:
            leituras_iguais = 0
        ultimo_valor = valor_atual
        time.sleep(intervalo)

    log("Aviso: nao deu para confirmar que a contagem de imoveis estabilizou, seguindo mesmo assim.", "warn")


def exportar_relatorio(driver, wait):
    log("Solicitando exportacao do relatorio...")
    botao_menu_xpath = "//page-reports-view//button[.//i[normalize-space()='more_vert']]"
    exportar_xpath = "//*[normalize-space(text())='Exportar']"

    opcao_exportar = None
    for tentativa in range(1, 4):
        botao_menu = wait.until(EC.element_to_be_clickable((By.XPATH, botao_menu_xpath)))
        clicar_com_retentativa(driver, botao_menu)
        try:
            opcao_exportar = WebDriverWait(driver, 3).until(
                EC.element_to_be_clickable((By.XPATH, exportar_xpath))
            )
            break
        except TimeoutException:
            log(f"Tentativa {tentativa}: menu de opcoes nao abriu, clicando de novo...")
            continue

    if opcao_exportar is None:
        raise RuntimeError("Nao consegui abrir o menu de opcoes para exportar o relatorio.")

    for tentativa in range(1, 4):
        clicar_com_retentativa(driver, opcao_exportar)
        try:
            WebDriverWait(driver, 3).until(EC.presence_of_element_located(
                (By.XPATH, "//*[contains(text(),'relatório é processado')]")
            ))
            log("Exportacao solicitada. O Imobzi esta processando o relatorio em segundo plano...")
            return
        except TimeoutException:
            log(f"Tentativa {tentativa}: clique em 'Exportar' nao confirmou, tentando de novo...")
            continue

    raise RuntimeError("Nao consegui confirmar a solicitacao de exportacao do relatorio.")


def _download_comecou(pasta_download, arquivos_antes, timeout=15):
    inicio = time.time()
    while time.time() - inicio < timeout:
        try:
            arquivos_agora = set(os.listdir(pasta_download))
        except FileNotFoundError:
            arquivos_agora = set()
        if arquivos_agora - arquivos_antes:
            return True
        time.sleep(0.5)
    return False


def _notificacoes_estao_abertas(driver):
    """O painel de notificacoes e um MODAL do Ionic (z-index bem alto, com
    backdrop cobrindo a tela inteira) - se ficar aberto por cima da tela,
    intercepta qualquer clique num elemento que esteja "por baixo" dele
    (ex.: o filtro de Usuario). Detectado ja aconteceu na pratica: a
    notificacao de "relatorio pronto" as vezes reaparece/fica aberta bem na
    hora em que a 2a etapa (segundo relatorio) esta mexendo no filtro."""
    try:
        return bool(driver.find_elements(By.CSS_SELECTOR, "ion-modal.notification-side-menu"))
    except Exception:
        return False


def _fechar_painel_notificacoes(driver):
    """Fecha o painel de notificacoes - melhor esforco, mas agora com mais
    de uma tentativa: primeiro o botao "X" do cabecalho, depois a tecla
    Esc, e por fim um clique no proprio backdrop do modal. Nao trava o
    script se nao conseguir; so loga um aviso pra ficar visivel no log em
    vez de falhar silenciosamente."""
    if not _notificacoes_estao_abertas(driver):
        return

    try:
        botao_fechar = driver.find_element(
            By.XPATH,
            "//notification//button[.//i[normalize-space()='close']]"
            " | //button[.//i[normalize-space()='close']]",
        )
        clicar_com_retentativa(driver, botao_fechar)
        time.sleep(0.5)
    except Exception:
        pass

    if not _notificacoes_estao_abertas(driver):
        return

    try:
        driver.switch_to.active_element.send_keys(Keys.ESCAPE)
        time.sleep(0.5)
    except Exception:
        pass

    if not _notificacoes_estao_abertas(driver):
        return

    try:
        backdrop = driver.find_element(By.CSS_SELECTOR, "ion-modal.notification-side-menu ion-backdrop")
        driver.execute_script("arguments[0].click();", backdrop)
        time.sleep(0.5)
    except Exception:
        pass

    if _notificacoes_estao_abertas(driver):
        log("Aviso: nao consegui fechar o painel de notificacoes (pode atrapalhar o proximo clique).", "warn")


def esperar_relatorio_pronto_e_baixar(driver, wait, pasta_download, arquivos_antes):
    log("Aguardando o Imobzi terminar de processar o relatorio (isso pode levar 1 a 3 minutos)...")

    inicio = time.time()

    while time.time() - inicio < TIMEOUT_PROCESSAMENTO:
        _checar_cancelamento()
        try:
            botao_notificacoes = driver.find_element(
                By.XPATH, "//button[@aria-label='Notificações' or .//i[normalize-space()='notifications']]"
            )
            clicar_com_retentativa(driver, botao_notificacoes)
            time.sleep(1)

            tem_notificacao_pronta = bool(
                driver.find_elements(By.XPATH, "//b[normalize-space()='Baixar']")
            )
            _fechar_painel_notificacoes(driver)

            if not tem_notificacao_pronta:
                _sleep_interruptivel(5)
                continue

            log(
                f"Notificacao de relatorio pronto encontrada. Aguardando "
                f"{ESPERA_ANTES_DE_RECARREGAR}s antes de recarregar a pagina..."
            )
            _sleep_interruptivel(ESPERA_ANTES_DE_RECARREGAR)

            log("Recarregando a pagina antes de baixar...")
            driver.refresh()
            wait.until(EC.presence_of_element_located(
                (By.XPATH, "//*[contains(text(),'Total de im')]")
            ))

            botao_notificacoes = wait.until(EC.element_to_be_clickable(
                (By.XPATH, "//button[@aria-label='Notificações' or .//i[normalize-space()='notifications']]")
            ))
            clicar_com_retentativa(driver, botao_notificacoes)
            time.sleep(1)

            links_baixar = driver.find_elements(By.XPATH, "//b[normalize-space()='Baixar']")
            if not links_baixar:
                log("Aviso: depois do reload a notificacao nao apareceu ainda, tentando de novo...")
                _fechar_painel_notificacoes(driver)
                _sleep_interruptivel(3)
                continue

            log("Relatorio pronto! Clicando em 'Baixar'...")
            clicar_com_retentativa(driver, links_baixar[0])

            if not _download_comecou(pasta_download, arquivos_antes):
                log(
                    "Aviso: cliquei em 'Baixar' mas nenhum arquivo novo apareceu na pasta "
                    "de download (provavel notificacao obsoleta). Tentando de novo...",
                    "warn",
                )
                _fechar_painel_notificacoes(driver)
                _sleep_interruptivel(3)
                continue

            log("Download iniciado!", "success")
            _fechar_painel_notificacoes(driver)
            return

        except SincronizacaoCancelada:
            raise
        except Exception:
            pass

        _sleep_interruptivel(5)

    raise TimeoutException(
        "O relatorio nao ficou pronto (ou nao consegui confirmar um download de verdade) "
        "dentro do tempo esperado. Verifique manualmente as notificacoes no Imobzi."
    )


def esperar_download_terminar(pasta_download, arquivos_antes, timeout=TIMEOUT_DOWNLOAD):
    log("Aguardando o arquivo terminar de baixar...")
    inicio = time.time()
    while time.time() - inicio < timeout:
        _checar_cancelamento()
        arquivos_agora = set(os.listdir(pasta_download))
        novos = arquivos_agora - arquivos_antes
        novos_completos = [
            f for f in novos
            if not f.endswith(".crdownload") and not f.endswith(".tmp")
        ]
        if novos_completos:
            caminho = os.path.join(pasta_download, novos_completos[0])
            tamanho1 = os.path.getsize(caminho)
            time.sleep(2)
            tamanho2 = os.path.getsize(caminho)
            if tamanho1 == tamanho2:
                return caminho
        time.sleep(2)

    raise TimeoutException("O download nao terminou dentro do tempo esperado.")


# --------------------------------------------------------------------------
# Normalizacao de arquivo baixado (BOM/mojibake) - herdado sem alteracoes
# --------------------------------------------------------------------------

BOM_UTF8 = b"\xef\xbb\xbf"
_SINAIS_MOJIBAKE = ("Ã", "Â", "â")


def _decodificar(conteudo):
    if conteudo.startswith(BOM_UTF8):
        return conteudo[len(BOM_UTF8):].decode("utf-8", errors="replace")
    if conteudo.startswith(b"\xff\xfe") or conteudo.startswith(b"\xfe\xff"):
        return conteudo.decode("utf-16", errors="replace")
    try:
        return conteudo.decode("utf-8")
    except UnicodeDecodeError:
        return conteudo.decode("cp1252", errors="replace")


def _reparar_mojibake(texto):
    consertou = False
    for _ in range(3):
        if not texto or not any(sinal in texto for sinal in _SINAIS_MOJIBAKE):
            break
        antes = sum(texto.count(sinal) for sinal in _SINAIS_MOJIBAKE)
        melhorou = False
        for codificacao in ("cp1252", "latin-1"):
            try:
                corrigido = texto.encode(codificacao).decode("utf-8")
            except (UnicodeEncodeError, UnicodeDecodeError):
                continue
            depois = sum(corrigido.count(sinal) for sinal in _SINAIS_MOJIBAKE)
            if depois < antes:
                texto = corrigido
                consertou = True
                melhorou = True
                break
        if not melhorou:
            break
    return texto, consertou


def _ler_texto(caminho):
    with open(caminho, "rb") as arquivo:
        conteudo = arquivo.read()
    texto, consertado = _reparar_mojibake(_decodificar(conteudo))
    if consertado:
        log("Texto estava duplamente codificado - acentuacao consertada.")
    return texto


def garantir_bom_utf8(caminho):
    extensao = os.path.splitext(caminho)[1].lower()
    if extensao not in (".csv", ".txt"):
        return caminho

    with open(caminho, "rb") as arquivo:
        original = arquivo.read()

    texto, consertado = _reparar_mojibake(_decodificar(original))
    if consertado:
        log("ATENCAO: o arquivo veio do Imobzi com acentuacao duplamente codificada. Consertei.", "warn")

    novo = BOM_UTF8 + texto.encode("utf-8")
    if novo == original:
        return caminho

    with open(caminho, "wb") as arquivo:
        arquivo.write(novo)
    log("Arquivo normalizado: UTF-8 com BOM.")
    return caminho


def _detectar_separador(amostra):
    separadores = (";", ",", "\t", "|")
    try:
        return csv.Sniffer().sniff(amostra, delimiters="".join(separadores)).delimiter
    except csv.Error:
        contagens = {sep: amostra.count(sep) for sep in separadores}
        melhor = max(contagens, key=contagens.get)
        return melhor if contagens[melhor] else ";"


_RE_DATA = re.compile(
    r"^(\d{1,2})/(\d{1,2})/(\d{4})(?:[\s,]+(\d{1,2}):(\d{2})(?::(\d{2}))?)?$"
)
_RE_NUMERO_BR = re.compile(r"^-?\d{1,3}(?:\.\d{3})+(?:,\d+)?$|^-?\d+(?:,\d+)?$")


def _converter_valor(texto):
    valor = (texto or "").strip()
    if not valor:
        return None
    if not CONVERTER_TIPOS_NO_XLSX:
        return texto

    achou_data = _RE_DATA.match(valor)
    if achou_data:
        dia, mes, ano, hora, minuto, segundo = achou_data.groups()
        try:
            return datetime(
                int(ano), int(mes), int(dia),
                int(hora or 0), int(minuto or 0), int(segundo or 0),
            )
        except ValueError:
            return texto

    bruto = valor.replace("R$", "").replace("\xa0", " ").strip()
    if _RE_NUMERO_BR.match(bruto):
        parte_inteira = bruto.lstrip("-").split(",")[0].replace(".", "")
        inteiro_puro = "," not in bruto and "." not in bruto
        if len(parte_inteira) > 1 and parte_inteira.startswith("0"):
            return texto
        if inteiro_puro and len(parte_inteira) > 9:
            return texto
        if len(parte_inteira) > 15:
            return texto
        try:
            numero = float(bruto.replace(".", "").replace(",", "."))
        except ValueError:
            return texto
        if "," not in bruto and numero.is_integer():
            return int(numero)
        return numero

    return texto


def converter_para_xlsx(caminho_csv):
    extensao = os.path.splitext(caminho_csv)[1].lower()
    if extensao not in (".csv", ".txt"):
        return None

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        log("AVISO: 'openpyxl' nao esta instalado - nao gerei o .xlsx (o CSV ja foi baixado).", "warn")
        return None

    caminho_xlsx = os.path.splitext(caminho_csv)[0] + ".xlsx"
    log(f"Convertendo para Excel: {os.path.basename(caminho_xlsx)} ...")

    texto = _ler_texto(caminho_csv)
    separador = _detectar_separador(texto[:16384])

    planilha = Workbook()
    aba = planilha.active
    aba.title = "Relatório"

    larguras = {}
    total_linhas = 0

    for linha in csv.reader(io.StringIO(texto, newline=""), delimiter=separador):
        total_linhas += 1
        aba.append([_converter_valor(celula) for celula in linha])
        for indice, celula in enumerate(aba[aba.max_row], start=1):
            bruto = linha[indice - 1] if indice - 1 < len(linha) else ""
            if isinstance(celula.value, datetime):
                celula.number_format = (
                    "DD/MM/YYYY HH:MM" if (celula.value.hour or celula.value.minute) else "DD/MM/YYYY"
                )
            elif isinstance(celula.value, (int, float)) and "R$" in (bruto or ""):
                celula.number_format = "R$ #,##0.00"
        if total_linhas <= 300:
            for indice, celula in enumerate(linha, start=1):
                larguras[indice] = min(60, max(larguras.get(indice, 10), len(celula) + 2))

    if total_linhas == 0:
        log("AVISO: o CSV estava vazio - .xlsx nao foi gerado.", "warn")
        return None

    for celula in aba[1]:
        celula.font = Font(bold=True)
    aba.freeze_panes = "A2"
    if aba.max_row > 1:
        aba.auto_filter.ref = aba.dimensions
    for indice, largura in larguras.items():
        aba.column_dimensions[get_column_letter(indice)].width = largura

    planilha.save(caminho_xlsx)
    log(f"Excel gerado com sucesso ({total_linhas} linhas): {caminho_xlsx}")
    return caminho_xlsx


def proximo_nome_disponivel(pasta_destino, data_str, extensao):
    versao = 1
    while True:
        nome = f"{NOME_BASE_ARQUIVO}_{data_str}v{versao}{extensao}"
        caminho = os.path.join(pasta_destino, nome)
        if not os.path.exists(caminho):
            return caminho
        versao += 1


def renomear_arquivo(caminho_baixado, pasta_destino):
    extensao = os.path.splitext(caminho_baixado)[1]
    data_str = datetime.now().strftime("%d%m%Y")
    destino = proximo_nome_disponivel(pasta_destino, data_str, extensao)
    os.makedirs(pasta_destino, exist_ok=True)
    os.replace(caminho_baixado, destino)
    log(f"Arquivo salvo como: {destino}")
    garantir_bom_utf8(destino)
    return destino


def salvar_diagnostico(driver, pasta_download):
    try:
        pasta_diag = os.path.join(pasta_download, "diagnostico")
        os.makedirs(pasta_diag, exist_ok=True)
        carimbo = datetime.now().strftime("%Y%m%d_%H%M%S")
        caminho_print = os.path.join(pasta_diag, f"erro_{carimbo}.png")
        caminho_html = os.path.join(pasta_diag, f"erro_{carimbo}.html")
        driver.save_screenshot(caminho_print)
        with open(caminho_html, "w", encoding="utf-8") as f:
            f.write(driver.page_source)
        log(f"Diagnostico salvo em: {caminho_print}", "warn")
    except Exception as e:
        log(f"Nao foi possivel salvar o diagnostico do erro: {e}", "warn")


# --------------------------------------------------------------------------
# Unificacao das duas planilhas (herdado do baixar_e_unificar_relatorios.py)
# --------------------------------------------------------------------------

def unir_planilhas(caminho1, caminho2, caminho_unificado):
    """Une duas planilhas em uma so: mantem TODO o conteudo da primeira e
    cola, logo em seguida, o conteudo da segunda a partir da 3a linha (pula
    as duas primeiras linhas dela - cabecalho/resumo)."""
    conteudo1 = _ler_texto(caminho1)
    conteudo2 = _ler_texto(caminho2)

    resto2 = "".join(conteudo2.splitlines(keepends=True)[2:])

    if conteudo1 and not conteudo1.endswith(("\n", "\r")):
        conteudo1 += "\r\n"

    with open(caminho_unificado, "wb") as saida:
        saida.write(BOM_UTF8)
        saida.write((conteudo1 + resto2).encode("utf-8"))


# --------------------------------------------------------------------------
# Orquestracao
# --------------------------------------------------------------------------

def baixar_um_relatorio(driver, wait, pasta_download, data_inicial):
    """Abre o relatorio, aplica o filtro de periodo/usuario, exporta,
    espera ficar pronto, baixa e renomeia. Devolve o caminho final do
    arquivo baixado (ja com BOM UTF-8 garantido)."""
    arquivos_antes = set(os.listdir(pasta_download)) if os.path.isdir(pasta_download) else set()

    abrir_relatorio(driver, wait)
    aplicar_filtro_periodo_total(driver, wait, data_inicial)
    exportar_relatorio(driver, wait)
    esperar_relatorio_pronto_e_baixar(driver, wait, str(pasta_download), arquivos_antes)

    caminho_baixado = esperar_download_terminar(str(pasta_download), arquivos_antes)
    return renomear_arquivo(caminho_baixado, str(pasta_download))


def sincronizar_imoveis(pasta_relatorio=None, pasta_destino=None, progress=None, stop_event=None):
    """Ponto de entrada usado pelo app (e pelo `python
    "Atualizar_imoveis_do_imobizi.py"`).

    pasta_relatorio: onde salvar/baixar os relatorios brutos do Imobzi
        (padrao: a propria pasta deste script).
    pasta_destino: pasta fixa onde o .xlsx unificado final e salvo (padrao:
        CRM_PROPERTY_DIR do .env do app, a mesma pasta de onde o app
        carrega a planilha mais recente).
    progress: callback opcional cb(mensagem, kind) para acompanhar em tempo
        real (o app passa uma funcao que guarda os logs para o polling).
    stop_event: threading.Event opcional; se for setado durante a execucao,
        a sincronizacao para de forma limpa (fecha o Chrome) e levanta
        SincronizacaoCancelada.
    """
    if progress:
        set_progress_callback(progress)

    global _STOP_EVENT
    _STOP_EVENT = stop_event or threading.Event()

    carregar_env()
    email = os.environ.get("IMOBZI_EMAIL", "").strip()
    senha = os.environ.get("IMOBZI_SENHA", "").strip()
    if not email or not senha:
        raise RuntimeError(
            "Defina IMOBZI_EMAIL e IMOBZI_SENHA no .env do app (app/sp_integrado/.env) "
            "ou no .env desta pasta."
        )
    global REPORT_URL
    REPORT_URL = os.environ.get("IMOBZI_REPORT_URL_IMOVEIS", "").strip()
    if not REPORT_URL:
        raise RuntimeError(
            "Defina IMOBZI_REPORT_URL_IMOVEIS (link do relatorio 'Imoveis - Cadastros' "
            "do Imobzi) no .env do app (app/sp_integrado/.env) ou no .env desta pasta."
        )

    pasta_download = Path(pasta_relatorio).resolve() if pasta_relatorio else BASE_DIR
    pasta_final = Path(pasta_destino).resolve() if pasta_destino else resolver_pasta_destino()

    driver = montar_driver(pasta_download)
    wait = WebDriverWait(driver, TIMEOUT_PADRAO)
    caminho_1 = None
    caminho_2 = None

    try:
        fazer_login(driver, wait, email, senha)

        log("=== Etapa 1/2: baixando o primeiro relatório (desde 01/01/2012) ===")
        caminho_1 = baixar_um_relatorio(driver, wait, pasta_download, DATA_INICIAL_1)
        log(f"Primeiro relatório baixado: {caminho_1}")

        log("=== Etapa 2/2: baixando o segundo relatório (desde 17/12/2024) ===")
        caminho_2 = baixar_um_relatorio(driver, wait, pasta_download, DATA_INICIAL_2)
        log(f"Segundo relatório baixado: {caminho_2}")
    except SincronizacaoCancelada:
        log("Sincronizacao cancelada pelo usuario.", "warn")
        raise
    except Exception:
        salvar_diagnostico(driver, str(pasta_download))
        raise
    finally:
        driver.quit()

    carimbo_dia = datetime.now().strftime("%d%m%y")

    # Arquivos intermediarios que serao apagados no final.
    arquivos_temporarios = [
        caminho_1,
        caminho_2,
        os.path.splitext(caminho_1)[0] + ".xlsx",  # xlsx "irmao" do 1o relatorio (se houver)
        os.path.splitext(caminho_2)[0] + ".xlsx",  # xlsx "irmao" do 2o relatorio (se houver)
    ]

    extensao_1 = os.path.splitext(caminho_1)[1].lower()
    extensao_2 = os.path.splitext(caminho_2)[1].lower()
    if extensao_1 not in (".csv", ".txt") or extensao_2 not in (".csv", ".txt"):
        raise RuntimeError(
            "Pelo menos um dos arquivos baixados nao e .csv/.txt - nao da "
            "para unificar nem gerar o .xlsx final."
        )

    log("Unificando as duas planilhas...")
    nome_unificado_csv = f"{NOME_BASE_ARQUIVO}_{carimbo_dia}_unificado.csv"
    caminho_unificado_csv = os.path.join(str(pasta_download), nome_unificado_csv)
    unir_planilhas(caminho_1, caminho_2, caminho_unificado_csv)
    arquivos_temporarios.append(caminho_unificado_csv)

    caminho_xlsx_temp = converter_para_xlsx(caminho_unificado_csv)
    if not caminho_xlsx_temp:
        raise RuntimeError(
            "Nao foi possivel gerar o .xlsx unificado (veja os avisos acima - "
            "provavelmente falta instalar 'openpyxl')."
        )

    pasta_final.mkdir(parents=True, exist_ok=True)
    nome_final = f"{NOME_BASE_ARQUIVO}_{carimbo_dia}_unificado.xlsx"
    destino_final = pasta_final / nome_final
    if destino_final.exists():
        destino_final.unlink()  # ja rodou hoje - substitui pelo mais novo
    shutil.move(caminho_xlsx_temp, destino_final)
    log(f"Relatório final salvo em: {destino_final}", "success")

    for arquivo in arquivos_temporarios:
        if arquivo and os.path.isfile(arquivo):
            try:
                os.remove(arquivo)
                log(f"Removido (temporário): {arquivo}")
            except OSError as exc:
                log(f"Aviso: nao consegui remover o temporario '{arquivo}': {exc}", "warn")

    log("Concluído com sucesso!", "success")

    return {
        "destino_final": str(destino_final),
        "carimbo_dia": carimbo_dia,
        "relatorio_1": str(caminho_1),
        "relatorio_2": str(caminho_2),
    }


def main():
    try:
        resultado = sincronizar_imoveis()
    except Exception as exc:
        log(f"ERRO: {exc}", "error")
        sys.exit(1)
    log(f"Resultado: {resultado}")


if __name__ == "__main__":
    main()
